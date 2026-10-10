"""Real models + real upload/chat/policy/report APIs. Only LLM is a spy."""

from contextlib import nullcontext
from io import BytesIO
from pathlib import Path
from urllib.parse import urlsplit
from uuid import UUID, uuid4
import base64
import json

import pytest
from PIL import Image
from docx import Document

from app.llm.client import ProviderDelta
from app.services.engine_adapter import (
    DatabasePolicyStore,
    DatabaseEventSink,
    PipelineAdapter,
    privacy_engine,
)
from app.db.models.submission import Submission
from app.db.models.attachment import Attachment
from app.storage.database import DatabaseFileStorage
from app.storage.files import storage_for_request
from conftest import login_as

FACE = Path(__file__).parents[3] / "packages/engine/tests/fixtures/astronaut.png"


class SpyLLM:
    def __init__(self):
        self.requests = []

    async def stream(self, request):
        self.requests.append(request)
        yield ProviderDelta(text="Sanitized request received.")


class BoundSink:
    def __init__(self, env):
        self.env = env
        self.policy_version = None

    def current(self):
        sub = self.env.db.query(Submission).filter_by(state="processing").one()
        sink = DatabaseEventSink(self.env.db, sub, sub.employee_id)
        sink.policy_version = self.policy_version
        return sink

    def record(self, batch):
        self.current().record(batch)

    def record_failure(self):
        self.current().record_failure()


@pytest.mark.parametrize("kind", ["png", "docx"])
def test_employee_upload_safe_provider_download_and_manager_report(
    env, monkeypatch, kind
):
    spy = SpyLLM()
    for name, value in [
        ("policy_store", DatabasePolicyStore(env.db)),
        ("sanitization_engine", PipelineAdapter()),
        ("event_sink", BoundSink(env)),
        ("llm_client", spy),
    ]:
        monkeypatch.setattr(env.client.app.state, name, value, raising=False)
    # Same durable store instantiated twice to simulate independent cold starts.
    store = lambda: DatabaseFileStorage(
        lambda: nullcontext(env.db), b"fictional-shared-signing-secret"
    )
    monkeypatch.setattr("app.routes.uploads.storage_for_request", store)
    monkeypatch.setattr("app.routes.chat.storage_for_request", store)
    env.client.app.dependency_overrides[storage_for_request] = store
    manager_headers = login_as(env, env.users[0])
    employee = env.client.post(
        "/api/v1/organization/employees",
        headers=manager_headers,
        json={
            "email": "new@employee.test",
            "display_name": "New Employee",
            "password": "",
        },
    ).json()["user"]
    policy = env.client.get("/api/v1/rules").json()
    custom = env.client.post(
        "/api/v1/rules/custom",
        headers=manager_headers,
        json={
            "expected_version": policy["version"],
            "name": "Internal reference",
            "kind": "exact",
            "pattern": "CONFIDENTIAL-777",
            "case_sensitive": True,
            "enabled": True,
        },
    )
    assert custom.status_code == 201
    env.client.cookies.clear()
    signed = env.client.post(
        "/api/v1/auth/login",
        json={"email": employee["email"], "password": ""},
        headers={"Origin": "http://localhost:5173"},
    )
    assert signed.status_code == 200
    headers = {
        "Origin": "http://localhost:5173",
        "X-CSRF-Token": signed.json()["csrf_token"],
    }
    raw = FACE.read_bytes()
    mime = "image/png"
    if kind == "docx":
        doc = Document()
        doc.add_paragraph("Email: dana@example.com CONFIDENTIAL-777")
        doc.add_picture(BytesIO(raw))
        out = BytesIO()
        doc.save(out)
        raw = out.getvalue()
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    init = env.client.post(
        "/api/v1/uploads/init",
        headers=headers,
        json={
            "filename": f"face.{kind}",
            "declared_mime": mime,
            "size_bytes": len(raw),
        },
    )
    assert init.status_code == 201
    aid = init.json()["attachment"]["id"]
    grant = init.json()["upload"]
    url = urlsplit(grant["url"])
    assert (
        env.client.put(
            url.path + "?" + url.query, content=raw, headers=grant["headers"]
        ).status_code
        == 204
    )
    assert (
        env.client.post(
            f"/api/v1/uploads/{aid}/complete", json={}, headers=headers
        ).status_code
        == 200
    )
    conversation = env.client.post(
        "/api/v1/conversations", json={"title": "Guarded test"}, headers=headers
    ).json()["id"]
    body = {
        "text": "Email: dana@example.com CONFIDENTIAL-777",
        "attachment_ids": [aid],
        "client_request_id": str(uuid4()),
    }
    response = env.client.post(
        f"/api/v1/conversations/{conversation}/messages", json=body, headers=headers
    )
    assert "event: done" in response.text, response.text
    assert len(spy.requests) == 1
    payload = spy.requests[0]
    assert "dana@example.com" not in json.dumps(payload.model_dump(mode="json"))
    assert "CONFIDENTIAL-777" not in json.dumps(payload.model_dump(mode="json"))
    download = env.client.get(f"/api/v1/uploads/{aid}/sanitized")
    assert download.status_code == 200
    if kind == "png":
        assert (
            base64.b64decode(payload.artifacts[0].sanitized_image_base64)
            == download.content
        )
        assert not privacy_engine().face_detector.detect(
            download.content, artifact_id=UUID(aid), source_ref="image"
        )
    else:
        assert "[REDACTED]" in payload.artifacts[0].sanitized_text
    assert env.client.get("/api/v1/reports").status_code == 403
    replay = env.client.post(
        f"/api/v1/conversations/{conversation}/messages", json=body, headers=headers
    )
    assert replay.text == response.text
    assert len(spy.requests) == 1
    login_as(env, env.users[0])
    reports = env.client.get("/api/v1/reports").json()["items"]
    assert len(reports) == 1
    assert reports[0]["employee"]["id"] == employee["id"]
    assert reports[0]["status"] == "completed"
    assert any(
        f["rule_id"] == "FACE" and f["action"] == "redacted"
        for f in reports[0]["findings"]
    )
    assert any(f["rule_id"].startswith("custom:") for f in reports[0]["findings"])
    assert "dana@example.com" not in json.dumps(reports)
    assert "CONFIDENTIAL-777" not in json.dumps(reports)
    login_as(env, env.users[2])
    assert env.client.get("/api/v1/reports").json()["items"] == []
    login_as(env, env.users[3])
    assert env.client.get(f"/api/v1/uploads/{aid}/sanitized").status_code == 404


def test_engine_failure_records_safe_failed_report_and_never_calls_provider(
    env, monkeypatch
):
    from app.db.models.manager_assignment import ManagerAssignment
    from datetime import datetime, UTC

    env.db.add(
        ManagerAssignment(
            employee_id=env.users[1].id,
            manager_id=env.users[0].id,
            organization_id=env.companies[0].id,
            updated_at=datetime.now(UTC),
        )
    )
    env.db.commit()

    class Broken:
        def sanitize(self, request):
            raise ValueError("secret payload must never surface")

    spy = SpyLLM()
    for name, value in [
        ("policy_store", DatabasePolicyStore(env.db)),
        ("sanitization_engine", Broken()),
        ("event_sink", BoundSink(env)),
        ("llm_client", spy),
    ]:
        monkeypatch.setattr(env.client.app.state, name, value, raising=False)
    headers = login_as(env, env.users[1])
    cid = env.client.post(
        "/api/v1/conversations", headers=headers, json={"title": "Failure"}
    ).json()["id"]
    response = env.client.post(
        f"/api/v1/conversations/{cid}/messages",
        headers=headers,
        json={
            "text": "secret payload",
            "attachment_ids": [],
            "client_request_id": str(uuid4()),
        },
    )
    assert "event: error" in response.text and "secret payload" not in response.text
    assert not spy.requests
    login_as(env, env.users[0])
    report = env.client.get("/api/v1/reports").json()["items"][0]
    assert report["status"] == "failed"
    assert report["failure_code"] == "PROCESSING_FAILED"
    assert not report["findings"]


def test_database_files_survive_new_storage_instances_and_sessions(env):
    from sqlalchemy.orm import sessionmaker
    from app.schemas.common import ApiError

    employee = env.users[1]
    aid = uuid4()
    key = f"original/{employee.organization_id}/{employee.id}/{aid}"
    env.db.add(
        Attachment(
            id=aid,
            organization_id=employee.organization_id,
            employee_id=employee.id,
            original_key=key,
            filename="private.png",
            mime_type="image/png",
            size_bytes=4,
        )
    )
    env.db.commit()
    factory = sessionmaker(env.db.get_bind(), expire_on_commit=False)
    first = DatabaseFileStorage(factory, b"fictional-shared-signing-secret")
    first.receive_upload(employee.organization_id, employee.id, aid, b"test")
    second = DatabaseFileStorage(factory, b"fictional-shared-signing-secret")
    assert second.read_original(key) == b"test"
    sanitized = second.write_sanitized(
        employee.organization_id, employee.id, aid, b"safe"
    )
    assert (
        DatabaseFileStorage(factory, b"fictional-shared-signing-secret").read_sanitized(
            sanitized
        )
        == b"safe"
    )
    with pytest.raises(ApiError):
        second.write_sanitized(env.companies[1].id, env.users[3].id, aid, b"leak")


def test_document_without_readable_text_can_be_sent_and_retained(env, monkeypatch):
    spy = SpyLLM()
    for name, value in [
        ("policy_store", DatabasePolicyStore(env.db)),
        ("sanitization_engine", PipelineAdapter()),
        ("event_sink", BoundSink(env)),
        ("llm_client", spy),
    ]:
        monkeypatch.setattr(env.client.app.state, name, value, raising=False)
    headers = login_as(env, env.users[1])
    out = BytesIO()
    Document().save(out)
    raw = out.getvalue()
    init = env.client.post(
        "/api/v1/uploads/init",
        headers=headers,
        json={
            "filename": "empty.docx",
            "declared_mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "size_bytes": len(raw),
        },
    )
    assert init.status_code == 201
    grant = init.json()["upload"]
    url = urlsplit(grant["url"])
    assert (
        env.client.put(
            url.path + "?" + url.query, content=raw, headers=grant["headers"]
        ).status_code
        == 204
    )
    aid = init.json()["attachment"]["id"]
    assert (
        env.client.post(
            f"/api/v1/uploads/{aid}/complete", json={}, headers=headers
        ).status_code
        == 200
    )
    cid = env.client.post(
        "/api/v1/conversations", headers=headers, json={"title": "Empty document"}
    ).json()["id"]
    for text, attachments in [("Summarize", [aid]), ("Follow up", [])]:
        response = env.client.post(
            f"/api/v1/conversations/{cid}/messages",
            headers=headers,
            json={
                "text": text,
                "attachment_ids": attachments,
                "client_request_id": str(uuid4()),
            },
        )
        assert "event: done" in response.text, response.text
    assert len(spy.requests) == 2
    assert (
        spy.requests[0].artifacts[0].sanitized_text
        == "[Sanitized document contains no readable text.]"
    )
    assert (
        spy.requests[1].artifacts[0].sanitized_text
        == "[Sanitized document contains no readable text.]"
    )
