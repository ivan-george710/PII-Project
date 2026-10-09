from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .models import ExtractedPage, ExtractionResult


def _shape_text(shape) -> list[str]:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return [
            text
            for child in shape.shapes
            for text in _shape_text(child)
        ]

    if shape.has_table:
        return [
            "\t".join(cell.text.strip() for cell in row.cells)
            for row in shape.table.rows
            if any(cell.text.strip() for cell in row.cells)
        ]

    if shape.has_text_frame and shape.text.strip():
        return [shape.text.strip()]

    return []


def extract_pptx(path: Path) -> ExtractionResult:
    """Extract text from every PowerPoint slide."""
    presentation = Presentation(path)
    pages = []

    for page_number, slide in enumerate(presentation.slides, start=1):
        slide_text = [
            text
            for shape in slide.shapes
            for text in _shape_text(shape)
        ]
        pages.append(
            ExtractedPage(
                page_number=page_number,
                text="\n".join(slide_text),
                extraction_method="text",
            )
        )

    return ExtractionResult(
        source_path=path,
        filename=path.name,
        file_type="pptx",
        pages=pages,
    )