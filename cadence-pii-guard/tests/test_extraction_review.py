from pathlib import Path
from types import SimpleNamespace

import pytest

from app.engines.extraction import (
    ExtractedPage,
    ExtractionEngine,
    ExtractionResult,
    OCRTextBlock,
)
from app.engines.extraction import engine as extraction_module
from app.engines.extraction import ocr as ocr_module
from app.engines.extraction.ocr import OCREngine

from docx import Document
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter


FIXTURES = Path(__file__).parent / "fixtures"
EXPECTED_BLOCK = OCRTextBlock(
    text="OCR fixture text",
    confidence=0.91,
    bounding_box=[
        [1.0, 2.0],
        [3.0, 2.0],
        [3.0, 4.0],
        [1.0, 4.0],
    ],
)


def _engine_with_strategy(
    monkeypatch,
    file_type,
    strategy,
    ocr,
    image_mime_type="image/png",
):
    mime_types = {
        "image": image_mime_type,
        "pdf": "application/pdf",
        "docx": "application/vnd.openxmlformats-officedocument."
        "wordprocessingml.document",
        "pptx": "application/vnd.openxmlformats-officedocument."
        "presentationml.presentation",
    }
    metadata = SimpleNamespace(
        file_type=file_type,
        mime_type=mime_types[file_type],
        extraction_strategy=strategy,
    )

    class FakeFileIntelligence:
        def analyze(self, _path):
            return metadata

    monkeypatch.setattr(
        extraction_module,
        "FileIntelligenceEngine",
        FakeFileIntelligence,
    )
    monkeypatch.setattr(extraction_module, "OCREngine", lambda: ocr)
    return ExtractionEngine()


def test_ocr_image_blocks_keep_text_confidence_and_box(monkeypatch):
    class FakeRapidOCR:
        def __call__(self, _image_path):
            return (
                [
                    (
                        [
                            [1.0, 2.0],
                            [3.0, 2.0],
                            [3.0, 4.0],
                            [1.0, 4.0],
                        ],
                        "OCR fixture text",
                        0.91,
                    )
                ],
                None,
            )

    monkeypatch.setattr(ocr_module, "RapidOCR", FakeRapidOCR)
    blocks = OCREngine().extract_image(FIXTURES / "hybrid_page.png")

    assert blocks == [EXPECTED_BLOCK]
    assert blocks[0].text == "OCR fixture text"
    assert blocks[0].confidence == 0.91
    assert blocks[0].bounding_box == EXPECTED_BLOCK.bounding_box


def test_scanned_pdf_strategy_returns_ocr_page_contract(monkeypatch):
    class FakeOCR:
        def extract_pdf(self, path):
            return ExtractionResult(
                source_path=path,
                filename=path.name,
                file_type="pdf",
                pages=[
                    ExtractedPage(
                        page_number=1,
                        text=EXPECTED_BLOCK.text,
                        extraction_method="ocr",
                        blocks=[EXPECTED_BLOCK],
                    )
                ],
            )

    engine = _engine_with_strategy(
        monkeypatch,
        "pdf",
        "ocr",
        FakeOCR(),
    )
    result = engine.extract(FIXTURES / "text.pdf")

    assert result.pages[0].page_number == 1
    assert result.pages[0].extraction_method == "ocr"
    assert result.pages[0].text == EXPECTED_BLOCK.text
    assert result.pages[0].blocks == [EXPECTED_BLOCK]


def test_hybrid_pdf_uses_text_then_ocr_with_correct_page_numbers(
    monkeypatch,
):
    class FakeOCR:
        def extract_pdf_page(self, _path, page_number):
            return ExtractedPage(
                page_number=page_number,
                text=EXPECTED_BLOCK.text,
                extraction_method="ocr",
                blocks=[EXPECTED_BLOCK],
            )

    engine = _engine_with_strategy(
        monkeypatch,
        "pdf",
        "hybrid",
        FakeOCR(),
    )
    result = engine.extract(FIXTURES / "hybrid.pdf")

    assert [page.page_number for page in result.pages] == [1, 2]
    assert [page.extraction_method for page in result.pages] == [
        "text",
        "ocr",
    ]
    assert result.pages[1].text == EXPECTED_BLOCK.text
    assert result.pages[1].blocks == [EXPECTED_BLOCK]


@pytest.mark.parametrize(
    ("mime_type", "file_type"),
    [("image/png", "png"), ("image/jpeg", "jpeg")],
)
def test_image_routes_to_ocr_and_returns_normalized_page(
    monkeypatch,
    tmp_path,
    mime_type,
    file_type,
):
    class FakeOCR:
        def extract_image_file(self, image_path, detected_type):
            return ExtractionResult(
                source_path=image_path,
                filename=image_path.name,
                file_type=detected_type,
                pages=[
                    ExtractedPage(
                        page_number=1,
                        text=EXPECTED_BLOCK.text,
                        extraction_method="ocr",
                        blocks=[EXPECTED_BLOCK],
                    )
                ],
            )

    engine = _engine_with_strategy(
        monkeypatch,
        "image",
        "ocr",
        FakeOCR(),
        image_mime_type=mime_type,
    )
    image_path = FIXTURES / "hybrid_page.png"
    if mime_type == "image/jpeg":
        image_path = tmp_path / "hybrid_page.jpg"
        with Image.open(FIXTURES / "hybrid_page.png") as image:
            image.convert("RGB").save(image_path, format="JPEG")

    result = engine.extract(image_path)

    assert result.file_type == file_type
    assert result.pages[0].page_number == 1
    assert result.pages[0].text == EXPECTED_BLOCK.text
    assert result.pages[0].extraction_method == "ocr"
    assert result.pages[0].blocks == [EXPECTED_BLOCK]


def test_docx_extracts_paragraph_and_table_content(
    monkeypatch,
    tmp_path,
):
    path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("Word paragraph")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Account"
    table.cell(0, 1).text = "1234"
    document.save(path)

    class UnusedOCR:
        pass

    engine = _engine_with_strategy(
        monkeypatch,
        "docx",
        "docx",
        UnusedOCR(),
    )
    result = engine.extract(path)

    assert result.file_type == "docx"
    assert len(result.pages) == 1
    assert result.pages[0].page_number == 1
    assert result.pages[0].extraction_method == "text"
    assert "Word paragraph" in result.pages[0].text
    assert "Account\t1234" in result.pages[0].text
    assert result.pages[0].blocks is None


def test_pptx_extracts_slide_text_and_numbers(
    monkeypatch,
    tmp_path,
):
    path = tmp_path / "sample.pptx"
    presentation = Presentation()
    for text in ("First slide", "Second slide"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        textbox = slide.shapes.add_textbox(
            Inches(1),
            Inches(1),
            Inches(5),
            Inches(1),
        )
        textbox.text = text
    presentation.save(path)

    class UnusedOCR:
        pass

    engine = _engine_with_strategy(
        monkeypatch,
        "pptx",
        "pptx",
        UnusedOCR(),
    )
    result = engine.extract(path)

    assert result.file_type == "pptx"
    assert [page.page_number for page in result.pages] == [1, 2]
    assert [page.extraction_method for page in result.pages] == [
        "text",
        "text",
    ]
    assert [page.text for page in result.pages] == [
        "First slide",
        "Second slide",
    ]


def _text_pdf_engine(monkeypatch):
    class TextPDFIntelligence:
        def analyze(self, _path):
            return SimpleNamespace(
                file_type="pdf",
                mime_type="application/pdf",
                extraction_strategy="text",
            )

    monkeypatch.setattr(
        extraction_module,
        "FileIntelligenceEngine",
        TextPDFIntelligence,
    )
    monkeypatch.setattr(
        extraction_module,
        "OCREngine",
        lambda: SimpleNamespace(),
    )
    return ExtractionEngine()


def test_corrupt_pdf_raises_clear_error(monkeypatch, tmp_path):
    path = tmp_path / "corrupt.pdf"
    path.write_bytes(b"%PDF-1.4\nthis is not a valid PDF")

    with pytest.raises(ValueError, match="Could not read PDF"):
        _text_pdf_engine(monkeypatch).extract(path)


def test_password_protected_pdf_raises_clear_error(
    monkeypatch,
    tmp_path,
):
    path = tmp_path / "encrypted.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt("secret")
    with path.open("wb") as output:
        writer.write(output)

    with pytest.raises(ValueError, match="password-protected"):
        _text_pdf_engine(monkeypatch).extract(path)


def test_pdf_without_pages_raises_clear_error(monkeypatch, tmp_path):
    path = tmp_path / "no-pages.pdf"
    with path.open("wb") as output:
        PdfWriter().write(output)

    with pytest.raises(ValueError, match="contains no pages"):
        _text_pdf_engine(monkeypatch).extract(path)


def test_ocr_pdf_retains_page_number_when_render_returns_no_image(
    monkeypatch,
    tmp_path,
):
    pdf_path = tmp_path / "two-pages.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.add_blank_page(width=100, height=100)
    with pdf_path.open("wb") as output:
        writer.write(output)

    class FakeRapidOCR:
        def __call__(self, _image_path):
            return None, None

    monkeypatch.setattr(ocr_module, "RapidOCR", FakeRapidOCR)
    monkeypatch.setattr(
        ocr_module,
        "convert_from_path",
        lambda *args, **kwargs: [],
    )

    result = OCREngine().extract_pdf(pdf_path)

    assert [page.page_number for page in result.pages] == [1, 2]
    assert [page.extraction_method for page in result.pages] == [
        "ocr",
        "ocr",
    ]
    assert [page.text for page in result.pages] == ["", ""]
    assert [page.blocks for page in result.pages] == [[], []]
