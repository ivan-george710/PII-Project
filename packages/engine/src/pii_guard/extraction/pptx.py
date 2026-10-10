from __future__ import annotations

from hashlib import sha1
from pathlib import Path
from uuid import UUID

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from pii_guard.errors import ProcessingError, ProcessingErrorCode
from pii_guard.extraction.embedded_images import (
    EmbeddedImage, ocr_embedded_images, part_images, part_paragraphs,
)
from pii_guard.inspection.file_type import detect_file_type
from pii_guard.models import ExtractedArtifact, ExtractedTextBlock, SourceLocation
from pii_guard.ocr.engine import OCRProvider


def extract_pptx(path: Path, *, artifact_id: UUID, filename: str,
                 ocr_provider: OCRProvider, max_slides: int = 100) -> ExtractedArtifact:
    detected = detect_file_type(path, artifact_id=artifact_id)
    if detected.file_format != "pptx":
        raise ProcessingError(ProcessingErrorCode.UNSUPPORTED_FILE, "Artifact is not a PPTX presentation.", artifact_id=artifact_id)
    try:
        presentation = Presentation(str(path))
        if len(presentation.slides) > max_slides:
            raise ProcessingError(ProcessingErrorCode.UNSUPPORTED_FILE, "PPTX exceeds the supported slide limit.", artifact_id=artifact_id)
        blocks = []
        images = []
        visited = set()
        slide_parts = set()
        for number, slide in enumerate(presentation.slides, 1):
            # XML traversal includes grouped shapes and table-cell paragraphs.
            blocks.extend(_text_blocks(slide.part, artifact_id, number))
            visited.add(str(slide.part.partname))
            slide_parts.add(str(slide.part.partname))
            if slide.has_notes_slide:
                blocks.extend(_text_blocks(slide.notes_slide.part, artifact_id, number))
                visited.add(str(slide.notes_slide.part.partname))
            placed_targets = set()
            for shape in slide.shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    if any(node.tag.endswith("}blip") for node in shape.element.iter()):
                        raise ProcessingError(ProcessingErrorCode.UNSUPPORTED_FILE, "Grouped image placement is not supported.")
                if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                    continue
                if shape.element.xpath(".//a:xfrm[@flipH='1' or @flipV='1' or @flipH='true' or @flipV='true']"):
                    raise ProcessingError(ProcessingErrorCode.UNSUPPORTED_FILE, "Flipped slide images are not supported.")
                crop = (shape.crop_left, shape.crop_top, shape.crop_right, shape.crop_bottom)
                if any(v < 0 or v >= 1 for v in crop) or crop[0] + crop[2] >= 1 or crop[1] + crop[3] >= 1:
                    raise ProcessingError(ProcessingErrorCode.UNSUPPORTED_FILE, "Unsupported image crop.")
                bbox = (shape.left / presentation.slide_width, shape.top / presentation.slide_height,
                        shape.width / presentation.slide_width, shape.height / presentation.slide_height)
                x, y, width, height = bbox
                if x < 0 or y < 0 or x + width > 1 or y + height > 1:
                    raise ProcessingError(ProcessingErrorCode.UNSUPPORTED_FILE, "Off-slide image placement is not supported.")
                images.append(EmbeddedImage(shape.image.blob, f".{shape.image.ext}",
                              f"pptx:slide:{number}:shape:{shape.shape_id}:image", bbox, slide_number=number,
                              rotation=shape.rotation, crop=crop, slide_size=(presentation.slide_width, presentation.slide_height)))
                placed_targets.add(shape.image.sha1)
            related = part_images(slide.part, prefix="pptx")
            placed_count = sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes)
            if len(slide.part._element.xpath(".//a:blip")) != placed_count or any(sha1(image.data).hexdigest() not in placed_targets for image in related):
                raise ProcessingError(ProcessingErrorCode.UNSUPPORTED_FILE, "Unsupported slide image placement.")
        # Ancillary parts retain part-local references, never invented slide positions.
        for part in presentation.part.package.iter_parts():
            name = str(part.partname)
            if name not in visited and name.startswith("/ppt/") and name.endswith(".xml"):
                blocks.extend(_text_blocks(part, artifact_id, None))
            if name not in slide_parts:
                images.extend(part_images(part, prefix="pptx"))
        blocks.extend(ocr_embedded_images(images, artifact_id=artifact_id, ocr_provider=ocr_provider))
    except ProcessingError as exc:
        exc.artifact_id = artifact_id
        raise
    except Exception as exc:
        raise ProcessingError(ProcessingErrorCode.EXTRACTION_FAILED, "PPTX content could not be extracted.", artifact_id=artifact_id) from exc
    return ExtractedArtifact(artifact_id=artifact_id, filename=filename, mime_type=detected.mime_type,
                             file_type="pptx", blocks=blocks, warnings=[])


def _text_blocks(part: object, artifact_id: UUID, slide_number: int | None) -> list[ExtractedTextBlock]:
    return [ExtractedTextBlock(text=text, extraction_method="native_text", location=SourceLocation(
        artifact_id=artifact_id, kind="native_text", page_number=None, slide_number=slide_number,
        source_ref=f"pptx:part:{part.partname}:paragraph:{index}", start=0, end=len(text), bbox=None))
        for index, text in enumerate(part_paragraphs(part, "http://schemas.openxmlformats.org/drawingml/2006/main"))]
