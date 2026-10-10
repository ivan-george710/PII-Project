from __future__ import annotations

from collections.abc import Iterable
from io import BytesIO
from pathlib import Path
from uuid import UUID
import re

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from lxml import etree as ET

from pii_guard.models import ExtractedTextBlock, SanitizedArtifact
from pii_guard.rules.policy import RedactionTarget
from pii_guard.redaction.metadata import sanitize_office_metadata
from pii_guard.redaction.office import (
    A, P, R, failure, inspect_editable, inspect_package, pack, paragraph_units, publish,
    relationships, rewrite_media, rewrite_unit, serialize, template, unpack, validate_inputs, xml,
)


def _presentation_structure(parts):
    root = xml(parts['ppt/presentation.xml'])
    default = xml(template('pptx')['ppt/presentation.xml'])
    for value in (root, default):
        for node in list(value):
            if node.tag == f'{{{P}}}sldIdLst':
                if any(n.tag != f'{{{P}}}sldId' or set(n.attrib) != {'id', f'{{{R}}}id'}
                       or len(n) or not re.fullmatch(r'[0-9]+', n.get('id', ''))
                       or not re.fullmatch(r'rId[0-9]+', n.get(f'{{{R}}}id', '')) for n in node):
                    raise ValueError('Invalid slide inventory')
                value.remove(node)
            elif node.tag == f'{{{P}}}sldSz':
                if set(node.attrib) != {'cx', 'cy', 'type'} or node.get('type') not in {'screen4x3', 'screen16x9', 'custom'} or len(node):
                    raise ValueError('Invalid presentation dimensions')
                if any(not re.fullmatch(r'[0-9]+', node.get(k, '')) or int(node.get(k)) <= 0 for k in ('cx', 'cy')):
                    raise ValueError('Invalid presentation dimensions')
                value.remove(node)
    if ET.tostring(root, method='c14n') != ET.tostring(default, method='c14n'):
        raise ValueError('Unsupported presentation properties')


def redact_pptx(data: bytes, targets: Iterable[RedactionTarget] = (), *, artifact_id: UUID,
                blocks: Iterable[ExtractedTextBlock] = (), region_sources=(), submission_id: UUID | None = None) -> bytes:
    """Rewrite ordinary shape text, notes and unrotated, uncropped pictures.

    Shared media is replaced once using the union of enabled image-local regions.
    Grouped/transformed shapes, hidden slides, non-stock masters/metadata and other
    unsupported parts are errors, even with no targets.
    """
    try:
        parts = unpack(data)
        sanitize_office_metadata(parts, 'pptx')
        presentation = Presentation(BytesIO(pack(parts)))
        if not 0 < len(presentation.slides) <= 100:
            raise ValueError('Unsupported slide count')
        editable, slide_numbers = set(), {}
        for number, slide in enumerate(presentation.slides, 1):
            name = str(slide.part.partname).lstrip('/')
            if name in editable:
                raise ValueError('Repeated slide part')
            editable.add(name)
            slide_numbers[name] = number
            if slide.has_notes_slide:
                notes = str(slide.notes_slide.part.partname).lstrip('/')
                if notes in editable:
                    raise ValueError('Shared speaker notes')
                editable.add(notes)
                slide_numbers[notes] = number
        rels = relationships(parts)
        inspect_package(parts, 'pptx', editable, rels)
        _presentation_structure(parts)
        roots, native, images = {}, {}, {}
        for name in editable:
            root = xml(parts[name])
            if root.tag != f'{{{P}}}{"notes" if "/notesSlides/" in name else "sld"}':
                raise ValueError('Invalid editable slide part')
            inspect_editable(root, 'pptx')
            ids = [n.get('id') for n in root.iter(f'{{{P}}}cNvPr')]
            if len(ids) != len(set(ids)):
                raise ValueError('Ambiguous slide shape identity')
            roots[name] = root
        for name, content in parts.items():
            if name.startswith('ppt/') and name.endswith('.xml'):
                root = roots.get(name)
                if root is None:
                    root = xml(content)
                for index, (p, nodes, text) in enumerate(paragraph_units(root, A)):
                    native[f'pptx:part:/{name}:paragraph:{index}'] = (p, nodes, text, slide_numbers.get(name))
        for number, slide in enumerate(presentation.slides, 1):
            name = str(slide.part.partname).lstrip('/')
            used = set()
            for shape in slide.shapes:
                if shape.shape_type not in {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.TEXT_BOX, MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.PLACEHOLDER}:
                    raise ValueError('Unsupported slide shape')
                if (shape.rotation != 0 or shape.width <= 0 or shape.height <= 0
                    or shape.left < 0 or shape.top < 0
                    or shape.left + shape.width > presentation.slide_width
                    or shape.top + shape.height > presentation.slide_height):
                    raise ValueError('Unsupported hidden or off-slide shape')
                if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                    continue
                if shape.rotation != 0 or any((shape.crop_left, shape.crop_top, shape.crop_right, shape.crop_bottom)):
                    raise ValueError('Unsupported cropped or rotated picture')
                bbox = (shape.left/presentation.slide_width, shape.top/presentation.slide_height,
                    shape.width/presentation.slide_width, shape.height/presentation.slide_height)
                x, y, w, h = bbox
                if w <= 0 or h <= 0 or x < 0 or y < 0 or x+w > 1 or y+h > 1:
                    raise ValueError('Unmappable picture placement')
                blips = list(shape.element.iter(f'{{{A}}}blip'))
                if len(blips) != 1:
                    raise ValueError('Ambiguous picture payload')
                rid = blips[0].get(f'{{{R}}}embed')
                media, reltype = rels[name][rid]
                if reltype != R + '/image' or '/media/' not in media:
                    raise ValueError('Invalid picture relationship')
                used.add(rid)
                images[f'pptx:slide:{number}:shape:{shape.shape_id}:image'] = (media, number, bbox)
            if used != {rid for rid, (_, typ) in rels[name].items() if typ == R + '/image'}:
                raise ValueError('Unplaced slide media')
            if len(list(roots[name].iter(f'{{{A}}}blip'))) != sum(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in slide.shapes):
                raise ValueError('Unsupported slide image placement')
        if any(list(root.iter(f'{{{A}}}blip')) for name, root in roots.items() if '/notesSlides/' in name):
            raise ValueError('Unsupported notes image placement')
        known, grouped = validate_inputs(artifact_id, blocks, targets, native, images, region_sources)
        for ref, unit_targets in grouped.items():
            if known[ref].location.kind != 'native_text':
                continue
            part_name = ref.split(':part:/', 1)[1].rsplit(':paragraph:', 1)[0]
            if part_name not in editable:
                raise ValueError('Unsupported ancillary text target')
            p, nodes, text, _ = native[ref]
            expected = rewrite_unit(nodes, text, unit_targets, known[ref].location, A)
            actual = paragraph_units(p, A)
            if len(actual) != 1 or actual[0][2] != expected:
                raise ValueError('Slide run rewrite did not complete')
        rewrite_media(parts, images, known, grouped, artifact_id)
        for name, root in roots.items():
            parts[name] = serialize(root)
        output = pack(parts)
        Presentation(BytesIO(output))
        return output
    except Exception:
        raise failure('pptx', artifact_id, submission_id) from None


def redact_pptx_file(source_path: Path, output_path: Path, targets: Iterable[RedactionTarget] = (), *,
                     artifact_id: UUID, blocks: Iterable[ExtractedTextBlock] = (),
                     submission_id: UUID | None = None) -> SanitizedArtifact:
    """Publish a fresh PPTX exclusively after the entire rewrite succeeds."""
    return publish(redact_pptx, 'pptx', source_path, output_path, targets,
        artifact_id=artifact_id, blocks=blocks, submission_id=submission_id)
