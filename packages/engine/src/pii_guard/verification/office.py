"""Serialized OOXML inventories, using a separate paragraph/placement reader."""
from io import BytesIO
import re
from xml.etree import ElementTree as ET

from docx import Document
from pptx import Presentation

from pii_guard.redaction.office import (
    unpack, relationships, inspect_package, inspect_editable, xml, W, A, P, R,
)
from .raster import image_evidence, verify_image
from .evidence import digest, expected_text


def paragraphs(root, ns):
    values=[]
    for paragraph in root.iter(f'{{{ns}}}p'):
        tokens=[]
        def visit(node):
            if node.tag==f'{{{ns}}}t': tokens.append(node.text or '')
            elif node.tag==f'{{{ns}}}tab': tokens.append('\t')
            elif node.tag in {f'{{{ns}}}br',f'{{{ns}}}cr'}: tokens.append('\n')
            for child in node:
                if child.tag!=f'{{{ns}}}p': visit(child)
        visit(paragraph)
        if tokens: values.append(''.join(tokens))
    return values


def inventory(parts,kind):
    rels=relationships(parts)  # Shared strict security policy, not writer state.
    images,units,editable,slides={}, {},set(),{}
    if kind=='docx':
        Document(BytesIO(_zip_bytes(parts)))
        editable.add('word/document.xml')
        root=ET.fromstring(parts['word/document.xml'])
        if root.tag!=f'{{{W}}}document':raise ValueError('Invalid main Word part')
        used=set()
        for node in root.iter(f'{{{A}}}blip'):
            rid=node.attrib[f'{{{R}}}embed'];media,typ=rels['word/document.xml'][rid]
            if typ!=R+'/image': raise ValueError('Invalid picture')
            used.add(rid)
            images[f'docx:part:/word/document.xml:image:{rid}']=(media,None,None)
        if used!={rid for rid,(_,typ) in rels['word/document.xml'].items() if typ==R+'/image'}:
            raise ValueError('Unplaced Word relationship')
        geometry=()
    else:
        # Standard parser verifies relationships/reopening. Source mapping and
        # coordinates are recomputed from serialized XML, not writer image maps.
        obj=Presentation(BytesIO(_zip_bytes(parts)))
        from pii_guard.redaction.pptx import _presentation_structure
        _presentation_structure(parts)  # Shared supported presentation policy.
        root=ET.fromstring(parts['ppt/presentation.xml'])
        size=root.find(f'{{{P}}}sldSz');width=int(size.attrib['cx']);height=int(size.attrib['cy'])
        slide_ids=root.find(f'{{{P}}}sldIdLst')
        names=[]
        if slide_ids is None: raise ValueError('Missing slides')
        for number,node in enumerate(slide_ids,1):
            name,typ=rels['ppt/presentation.xml'][node.attrib[f'{{{R}}}id']]
            if typ!=R+'/slide' or name in editable: raise ValueError('Invalid slide identity')
            names.append(name);editable.add(name);slides[name]=number
            for _,(target,typ) in rels.get(name,{}).items():
                if typ==R+'/notesSlide':
                    if target in editable: raise ValueError('Shared notes')
                    editable.add(target);slides[target]=number
            used=set()
            for pic in ET.fromstring(parts[name]).iter(f'{{{P}}}pic'):
                sid=pic.find(f'{{{P}}}nvPicPr/{{{P}}}cNvPr').attrib['id']
                blip=pic.find(f'{{{P}}}blipFill/{{{A}}}blip')
                rid=blip.attrib[f'{{{R}}}embed'];media,typ=rels[name][rid]
                used.add(rid)
                transform=pic.find(f'{{{P}}}spPr/{{{A}}}xfrm')
                if transform is None or set(transform.attrib)-{'rot','flipH','flipV'} or any(transform.attrib.values()):
                    raise ValueError('Unsupported picture transform')
                off=transform.find(f'{{{A}}}off');ext=transform.find(f'{{{A}}}ext')
                box=(int(off.attrib['x'])/width,int(off.attrib['y'])/height,
                    int(ext.attrib['cx'])/width,int(ext.attrib['cy'])/height)
                if typ!=R+'/image': raise ValueError('Invalid picture relationship')
                images[f'pptx:slide:{number}:shape:{sid}:image']=(media,number,box)
            if used!={rid for rid,(_,typ) in rels[name].items() if typ==R+'/image'}:
                raise ValueError('Unplaced slide relationship')
        if len(obj.slides)!=len(names) or not 0<len(names)<=100: raise ValueError('Invalid slide count')
        geometry=(width,height,tuple(names))
    inspect_package(parts,kind,editable,rels)
    for name in editable:
        parsed=xml(parts[name])
        if kind=='pptx':
            if parsed.tag!=f'{{{P}}}{"notes" if "/notesSlides/" in name else "sld"}':raise ValueError('Wrong slide/notes root')
            ids=[node.get('id') for node in parsed.iter(f'{{{P}}}cNvPr')]
            if len(ids)!=len(set(ids)):raise ValueError('Duplicate shape identity')
            if '/notesSlides/' in name and any(True for _ in parsed.iter(f'{{{A}}}blip')):raise ValueError('Unsupported note images')
        inspect_editable(parsed,kind)  # Reuse supported-feature policy only.
    ns=W if kind=='docx' else A
    for name,data in parts.items():
        if name.startswith('word/' if kind=='docx' else 'ppt/') and name.endswith('.xml'):
            for index,text in enumerate(paragraphs(ET.fromstring(data),ns)):
                units[f'{kind}:part:/{name}:paragraph:{index}']=(text,slides.get(name))
    if {name for name in parts if '/media/' in name}!={entry[0] for entry in images.values()}:
        raise ValueError('Unplaced media')
    return units,images,geometry,editable


def _zip_bytes(parts):
    # Reopening standard libraries from the already bounded package; not a writer.
    from zipfile import ZipFile, ZIP_STORED
    out=BytesIO()
    with ZipFile(out,'w',ZIP_STORED) as archive:
        for name,data in parts.items(): archive.writestr(name,data)
    return out.getvalue()


def prepare(parts,kind,artifact_id,blocks,targets):
    units,images,geometry,editable=inventory(parts,kind)
    grouped={}
    for target in targets: grouped.setdefault(target.location.source_ref,[]).append(target.location)
    expected={}
    for ref,(text,slide) in units.items():
        key=(artifact_id,'native_text',None,slide,ref,None)
        spans=[(loc.start,loc.end) for loc in grouped.get(ref,())]
        expected[ref]=(key,digest(key,expected_text(text,spans)))
    media_boxes={name:[] for name,_,_ in images.values()}
    for ref,block in blocks.items():
        loc=block.location
        if loc.kind=='native_text':
            if ref not in units or units[ref]!=(block.text,loc.slide_number) or loc.page_number is not None or loc.bbox is not None:
                raise ValueError('Wrong native original')
        else:
            base,sep,index=ref.rpartition(':ocr:')
            if not sep or not index.isascii() or not index.isdigit() or base not in images:
                raise ValueError('Wrong picture source')
            media,slide,placement=images[base]
            if loc.slide_number!=slide or loc.page_number is not None: raise ValueError('Wrong picture identity')
            if ref in grouped:
                box=loc.bbox
                if placement is not None:
                    x,y,w,h=placement;bx,by,bw,bh=box
                    box=((bx-x)/w,(by-y)/h,bw/w,bh/h)
                media_boxes[media].append(box)
    plans=tuple(image_evidence(parts[name],(artifact_id,name),boxes)
        for name,boxes in sorted(media_boxes.items()))
    # The original's metadata is already removed in the private preparation copy.
    # Every non-editable/non-media part must remain byte-identical afterward.
    fixed=tuple((name,digest((artifact_id,name),data)) for name,data in sorted(parts.items())
        if name not in editable and '/media/' not in name)
    return (tuple(sorted(expected.items())),plans,geometry,tuple(sorted(parts)),
            tuple(sorted(images.items())),fixed)


def verify(data,kind,artifact_id,plan):
    parts=unpack(data)
    # Explicit serialized-state assertions precede shared feature guards. Never
    # invoke the metadata remover on output: doing so could conceal a regression.
    if any(name.startswith('docProps/') or '/printerSettings/' in name for name in parts):
        raise ValueError('Ancillary metadata survived')
    for name,content in parts.items():
        if name.endswith('.xml') and (name=='word/document.xml' or re.fullmatch(r'ppt/(slides/slide|notesSlides/notesSlide)[0-9]+\.xml',name)):
            for node in ET.fromstring(content).iter():
                if node.tag in {f'{{{P}}}cNvPr','{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr',
                        '{http://schemas.openxmlformats.org/drawingml/2006/picture}cNvPr'}:
                    if node.attrib.get('name') or 'descr' in node.attrib or 'title' in node.attrib:
                        raise ValueError('Shape metadata survived')
    units,images,geometry,_=inventory(parts,kind)
    expected,media_plans,old_geometry,names,old_images,fixed=plan
    if tuple(sorted(parts))!=names or tuple(sorted(images.items()))!=old_images or geometry!=old_geometry:
        raise ValueError('Changed package or source geometry')
    if set(units)!={ref for ref,_ in expected}: raise ValueError('Missing native source')
    for ref,(key,expected_hash) in expected:
        text,slide=units[ref]
        if slide!=key[3] or digest(key,text)!=expected_hash: raise ValueError('Native redaction failed')
    for name,fingerprint in fixed:
        if digest((artifact_id,name),parts[name])!=fingerprint: raise ValueError('Changed ancillary or relationship data')
    for media_plan in media_plans:
        name=media_plan[0][1]
        verify_image(parts[name],media_plan,output_format='PNG' if name.endswith('.png') else 'JPEG')
    return tuple(parts.values())
