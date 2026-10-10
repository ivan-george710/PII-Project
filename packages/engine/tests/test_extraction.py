from __future__ import annotations

from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document
from PIL import Image, ImageDraw
from pptx import Presentation
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from pii_guard.errors import ProcessingError, ProcessingErrorCode
from pii_guard.extraction import extract_docx, extract_image, extract_pdf, extract_pptx
from pii_guard.ocr import OCRBlock
from pii_guard.inspection import analyze_pdf
from pypdf import PdfReader, PdfWriter
from PIL import ImageChops
from pydantic import ValidationError
from zipfile import ZipFile


class FakeOCR:
    def __init__(self, text: str = "OCR TEXT") -> None:
        self.text = text
        self.seen_paths: list[Path] = []
        self.existed_during_call: list[bool] = []

    def recognize_image(self, image_path: Path, **_: object) -> list[OCRBlock]:
        self.seen_paths.append(Path(image_path))
        self.existed_during_call.append(Path(image_path).exists())
        return [OCRBlock(self.text, 0.9, (0.1, 0.2, 0.3, 0.4))]


class FailingOCR:
    def recognize_image(self, image_path: Path, **_: object) -> list[OCRBlock]:
        raise ProcessingError(ProcessingErrorCode.OCR_FAILED, "OCR exploded")


def _write_text_pdf(path: Path, text: str = "Native PDF text") -> None:
    pdf = canvas.Canvas(str(path), pagesize=letter)
    pdf.drawString(72, 720, text)
    pdf.showPage()
    pdf.save()


def _write_image(path: Path, text: str = "scan") -> None:
    image = Image.new("RGB", (300, 160), "white")
    draw = ImageDraw.Draw(image)
    draw.text((20, 40), text, fill="black")
    image.save(path)


def _write_image_pdf(path: Path, image_path: Path) -> None:
    pdf = canvas.Canvas(str(path), pagesize=letter)
    pdf.drawImage(str(image_path), 72, 500, width=200, height=100)
    pdf.showPage()
    pdf.save()


def test_extract_native_text_pdf(tmp_path: Path) -> None:
    path = tmp_path / "native.pdf"
    _write_text_pdf(path, "Policy owner alpha")

    result = extract_pdf(
        path,
        artifact_id=uuid4(),
        filename="native.pdf",
        ocr_provider=FakeOCR(),
    )

    assert result.file_type == "pdf"
    assert result.full_text.strip() == "Policy owner alpha"
    assert result.blocks[0].location.kind == "native_text"
    assert result.blocks[0].location.page_number == 1
    assert result.blocks[0].location.start == 0
    assert result.blocks[0].location.end == len(result.blocks[0].text)


def test_extract_scanned_pdf_uses_ocr_and_cleans_temp_files(tmp_path: Path) -> None:
    image_path = tmp_path / "scan.jpg"
    _write_image(image_path)
    pdf_path = tmp_path / "scan.pdf"
    _write_image_pdf(pdf_path, image_path)
    ocr = FakeOCR("Scanned page text")

    result = extract_pdf(
        pdf_path,
        artifact_id=uuid4(),
        filename="scan.pdf",
        ocr_provider=ocr,
    )

    assert result.full_text == "Scanned page text"
    assert result.blocks[0].location.kind == "ocr"
    assert result.blocks[0].location.page_number == 1
    assert result.blocks[0].location.bbox == (0.1, 0.2, 0.3, 0.4)
    assert ocr.existed_during_call == [True]
    assert all(not path.exists() for path in ocr.seen_paths)


def test_extract_mixed_pdf_preserves_order(tmp_path: Path) -> None:
    image_path = tmp_path / "scan.jpg"
    _write_image(image_path)
    pdf_path = tmp_path / "mixed.pdf"
    pdf = canvas.Canvas(str(pdf_path), pagesize=letter)
    pdf.drawString(72, 720, "First page native")
    pdf.showPage()
    pdf.drawImage(str(image_path), 72, 500, width=200, height=100)
    pdf.showPage()
    pdf.save()

    result = extract_pdf(
        pdf_path,
        artifact_id=uuid4(),
        filename="mixed.pdf",
        ocr_provider=FakeOCR("Second page OCR"),
    )

    assert [block.text.strip() for block in result.blocks] == [
        "First page native",
        "Second page OCR",
    ]
    assert result.blocks[1].location.page_number == 2


def test_extract_pdf_ocr_images_on_native_text_page(tmp_path: Path) -> None:
    image_path = tmp_path / "scan.jpg"
    _write_image(image_path)
    pdf_path = tmp_path / "mixed-page.pdf"
    pdf = canvas.Canvas(str(pdf_path), pagesize=letter)
    pdf.drawString(72, 720, "Native page text")
    pdf.drawImage(str(image_path), 72, 500, width=200, height=100)
    pdf.showPage()
    pdf.save()

    result = extract_pdf(
        pdf_path,
        artifact_id=uuid4(),
        filename="mixed-page.pdf",
        ocr_provider=FakeOCR("Image text on same page"),
    )

    assert [block.text.strip() for block in result.blocks] == [
        "Native page text",
        "Image text on same page",
    ]
    assert result.blocks[0].location.kind == "native_text"
    assert result.blocks[1].location.kind == "ocr"
    assert result.blocks[1].location.page_number == 1


def test_extract_pdf_does_not_hide_ocr_failure(tmp_path: Path) -> None:
    image_path = tmp_path / "scan.jpg"
    _write_image(image_path)
    pdf_path = tmp_path / "scan.pdf"
    _write_image_pdf(pdf_path, image_path)

    with pytest.raises(ProcessingError) as error:
        extract_pdf(
            pdf_path,
            artifact_id=uuid4(),
            filename="scan.pdf",
            ocr_provider=FailingOCR(),
        )

    assert error.value.code is ProcessingErrorCode.OCR_FAILED


def test_extract_docx_native_text_and_embedded_image_ocr(tmp_path: Path) -> None:
    image_path = tmp_path / "scan.png"
    _write_image(image_path)
    path = tmp_path / "document.docx"
    document = Document()
    document.add_paragraph("DOCX native paragraph")
    document.add_picture(str(image_path))
    document.save(path)

    result = extract_docx(
        path,
        artifact_id=uuid4(),
        filename="document.docx",
        ocr_provider=FakeOCR("DOCX image text"),
    )

    assert "DOCX native paragraph" in result.full_text
    assert "DOCX image text" in result.full_text
    assert {block.extraction_method for block in result.blocks} == {"native_text", "ocr"}


def test_extract_pptx_native_text_notes_and_embedded_image_ocr(tmp_path: Path) -> None:
    image_path = tmp_path / "scan.png"
    _write_image(image_path)
    path = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "PPTX native title"
    slide.shapes.add_picture(str(image_path), 0, 0)
    slide.notes_slide.notes_text_frame.text = "PPTX speaker note"
    presentation.save(path)

    result = extract_pptx(
        path,
        artifact_id=uuid4(),
        filename="slides.pptx",
        ocr_provider=FakeOCR("PPTX image text"),
    )

    assert "PPTX native title" in result.full_text
    assert "PPTX speaker note" in result.full_text
    assert "PPTX image text" in result.full_text
    assert all(block.location.slide_number == 1 for block in result.blocks if block.text in {"PPTX native title", "PPTX speaker note", "PPTX image text"})
    ocr_block = next(block for block in result.blocks if block.extraction_method == "ocr")
    assert ocr_block.location.bbox != (0.1, 0.2, 0.3, 0.4)


def test_extract_standalone_image_uses_ocr(tmp_path: Path) -> None:
    path = tmp_path / "scan.png"
    _write_image(path)

    result = extract_image(
        path,
        artifact_id=uuid4(),
        filename="scan.png",
        ocr_provider=FakeOCR("Standalone image text"),
    )

    assert result.file_type == "png"
    assert result.full_text == "Standalone image text"
    assert result.blocks[0].location.bbox == (0.1, 0.2, 0.3, 0.4)


def test_extract_empty_docx_returns_no_blocks(tmp_path: Path) -> None:
    path = tmp_path / "empty.docx"
    Document().save(path)

    result = extract_docx(
        path,
        artifact_id=uuid4(),
        filename="empty.docx",
        ocr_provider=FakeOCR(),
    )

    assert result.blocks == []


def test_vector_scan_fallback_page_routing_and_blank_page(tmp_path: Path) -> None:
    path = tmp_path / "vector.pdf"
    pdf = canvas.Canvas(str(path), pagesize=(200, 100))
    pdf.drawString(10, 50, "Native")
    pdf.showPage()
    pdf.rect(20, 40, 40, 60, fill=1, stroke=0)
    pdf.showPage()
    pdf.showPage()
    pdf.save()
    assert not PdfReader(path).pages[1].images
    assert [page.requires_ocr for page in analyze_pdf(path).pages] == [False, True, False]
    ocr = FakeOCR()
    result = extract_pdf(path, artifact_id=uuid4(), filename=path.name, ocr_provider=ocr)
    assert [block.location.page_number for block in result.blocks] == [1, 2]
    assert result.blocks[1].location.source_ref == "pdf:page:2:rendered:ocr:0"
    assert result.blocks[1].location.bbox == (0.1, 0.2, 0.3, 0.4)
    assert len(ocr.seen_paths) == 1
    assert not ocr.seen_paths[0].parent.exists()


def test_inline_scan_without_xobjects_is_rendered(tmp_path: Path) -> None:
    path = tmp_path / "inline.pdf"
    pdf = canvas.Canvas(str(path), pagesize=(200, 100))
    pdf.drawInlineImage(Image.new("RGB", (20, 20), "black"), 20, 20)
    pdf.save()
    assert not PdfReader(path).pages[0]["/Resources"].get("/XObject")
    result = extract_pdf(path, artifact_id=uuid4(), filename=path.name, ocr_provider=FakeOCR())
    assert result.blocks[0].location.kind == "ocr"


@pytest.mark.parametrize("rotation,expected_size,expected_box", [(0, (400, 200), (0.1, 0, 0.2, 0.6)), (90, (200, 400), (0.4, 0.1, 0.6, 0.2))])
def test_rendered_pdf_pixel_mapping_rotation_and_scale(tmp_path: Path, rotation: int, expected_size: tuple, expected_box: tuple) -> None:
    path = tmp_path / "rotated.pdf"
    pdf = canvas.Canvas(str(path), pagesize=(200, 100))
    pdf.rect(20, 40, 40, 60, fill=1, stroke=0)
    pdf.save()
    writer = PdfWriter(clone_from=path)
    writer.pages[0].rotate(rotation)
    writer.write(path)

    class PixelOCR(FakeOCR):
        def recognize_image(self, image_path: Path, **kwargs: object) -> list[OCRBlock]:
            super().recognize_image(image_path, **kwargs)
            with Image.open(image_path) as image:
                assert image.size == expected_size
                bounds = ImageChops.invert(image.convert("RGB")).getbbox()
                x1, y1, x2, y2 = bounds
                from pii_guard.ocr import block_from_pixel_bbox
                return [block_from_pixel_bbox("REGION", 1, (x1, y1, x2-x1, y2-y1), width=image.width, height=image.height)]

    ocr = PixelOCR()
    result = extract_pdf(path, artifact_id=uuid4(), filename=path.name, ocr_provider=ocr)
    assert result.blocks[0].location.bbox == pytest.approx(expected_box, abs=0.01)
    assert not ocr.seen_paths[0].parent.exists()


def test_rendered_pdf_cropbox_coordinates(tmp_path: Path) -> None:
    path = tmp_path / "crop.pdf"
    pdf = canvas.Canvas(str(path), pagesize=(200, 100))
    pdf.rect(30, 30, 40, 40, fill=1, stroke=0)
    pdf.save()
    writer = PdfWriter(clone_from=path)
    writer.pages[0].cropbox.lower_left = (20, 20)
    writer.pages[0].cropbox.upper_right = (120, 80)
    writer.write(path)
    class SizeOCR(FakeOCR):
        def recognize_image(self, image_path: Path, **kwargs: object) -> list[OCRBlock]:
            with Image.open(image_path) as image:
                assert image.size == (200, 120)
            return super().recognize_image(image_path, **kwargs)
    result = extract_pdf(path, artifact_id=uuid4(), filename=path.name, ocr_provider=SizeOCR())
    assert result.blocks[0].location.bbox == (0.1, 0.2, 0.3, 0.4)


@pytest.mark.parametrize("extractor,suffix", [(extract_pdf,"pdf"), (extract_docx,"docx"), (extract_pptx,"pptx")])
def test_ocr_failure_cleanup_and_safe_errors(tmp_path: Path, extractor: object, suffix: str) -> None:
    image = tmp_path / "image.png"
    _write_image(image)
    path = tmp_path / f"document.{suffix}"
    if suffix == "pdf":
        _write_image_pdf(path, image)
    elif suffix == "docx":
        document = Document()
        document.add_picture(str(image))
        document.save(path)
    else:
        document = Presentation()
        document.slides.add_slide(document.slide_layouts[6]).shapes.add_picture(str(image), 0, 0)
        document.save(path)
    class BrokenOCR(FakeOCR):
        def recognize_image(self, image_path: Path, **kwargs: object) -> list[OCRBlock]:
            super().recognize_image(image_path, **kwargs)
            raise RuntimeError("SECRET DOCUMENT CONTENT")
    ocr = BrokenOCR()
    identity = uuid4()
    with pytest.raises(ProcessingError) as error:
        extractor(path, artifact_id=identity, filename=path.name, ocr_provider=ocr)
    assert error.value.code == ProcessingErrorCode.OCR_FAILED
    assert error.value.artifact_id == identity
    assert "SECRET" not in str(error.value)
    assert not ocr.seen_paths[0].parent.exists()


def test_renderer_failure_is_explicit(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "vector.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.rect(10, 10, 20, 20, fill=1)
    pdf.save()
    import pii_guard.extraction.pdf as module
    def broken(*args: object, **kwargs: object) -> None:
        raise RuntimeError("SECRET")
    monkeypatch.setattr(module.pdfium, "PdfDocument", broken)
    with pytest.raises(ProcessingError) as error:
        extract_pdf(path, artifact_id=uuid4(), filename=path.name, ocr_provider=FakeOCR())
    assert error.value.code == ProcessingErrorCode.OCR_FAILED
    assert "SECRET" not in str(error.value)


def test_docx_nested_text_headers_footers_and_images(tmp_path: Path) -> None:
    path = tmp_path / "nested.docx"
    image = tmp_path / "image.png"
    _write_image(image)
    document = Document()
    document.sections[0].header.paragraphs[0].text = "HEADER"
    document.sections[0].header.paragraphs[0].add_run().add_picture(str(image))
    document.sections[0].footer.paragraphs[0].text = "FOOTER"
    document.add_table(1, 1).cell(0, 0).add_table(1, 1).cell(0, 0).text = "NESTED"
    from docx.oxml import parse_xml
    document._element.body.append(parse_xml('<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:r><w:pict><w:txbxContent><w:p><w:r><w:t>TEXTBOX</w:t></w:r></w:p></w:txbxContent></w:pict></w:r></w:p>'))
    document.add_paragraph("caf\u00e9 \U0001f600\nsecond\tline")
    document.save(path)
    result = extract_docx(path, artifact_id=uuid4(), filename=path.name, ocr_provider=FakeOCR("HEADER IMAGE"))
    for text in ("HEADER", "FOOTER", "NESTED", "TEXTBOX", "HEADER IMAGE", "caf\u00e9 \U0001f600\nsecond\tline"):
        assert text in result.full_text
    assert result.full_text.count("TEXTBOX") == 1
    for block in result.blocks:
        assert block.location.end == len(block.text)
        assert block.location.page_number is None


def test_pptx_table_group_layout_and_master_text(tmp_path: Path) -> None:
    path = tmp_path / "nested.pptx"
    document = Presentation()
    slide = document.slides.add_slide(document.slide_layouts[6])
    slide.shapes.add_table(1, 1, 0, 0, 1000000, 1000000).table.cell(0, 0).text = "CELL"
    slide.shapes.add_group_shape().shapes.add_textbox(0, 0, 1000000, 1000000).text = "GROUP"
    slide.slide_layout.shapes[0].text = "LAYOUT"
    slide.slide_layout.slide_master.shapes[0].text = "MASTER"
    document.save(path)
    result = extract_pptx(path, artifact_id=uuid4(), filename=path.name, ocr_provider=FakeOCR())
    for text in ("CELL", "GROUP", "LAYOUT", "MASTER"):
        assert text in result.full_text
    assert next(b for b in result.blocks if b.text == "CELL").location.slide_number == 1
    assert next(b for b in result.blocks if b.text == "MASTER").location.slide_number is None


@pytest.mark.parametrize("suffix,extractor", [("docx",extract_docx), ("pptx",extract_pptx)])
def test_corrupt_office_package_safe_failure(tmp_path: Path, suffix: str, extractor: object) -> None:
    path = tmp_path / f"bad.{suffix}"
    with ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml" if suffix == "docx" else "ppt/presentation.xml", "SECRET broken XML")
    with pytest.raises(ProcessingError) as error:
        extractor(path, artifact_id=uuid4(), filename=path.name, ocr_provider=FakeOCR())
    assert error.value.code == ProcessingErrorCode.EXTRACTION_FAILED
    assert "SECRET" not in str(error.value)


def test_extraction_model_identity_method_and_spans(tmp_path: Path) -> None:
    path = tmp_path / "image.png"
    _write_image(path)
    from pii_guard.models import ExtractedArtifact, ExtractedTextBlock
    result = extract_image(path, artifact_id=uuid4(), filename=path.name, ocr_provider=FakeOCR("\U0001f600\ncaf\u00e9"))
    assert result.blocks[0].location.end == 6
    data = result.model_dump()
    data["artifact_id"] = uuid4()
    with pytest.raises(ValidationError):
        ExtractedArtifact.model_validate(data)
    for field, value in (("extraction_method", "native_text"),):
        block = result.blocks[0].model_dump()
        block[field] = value
        with pytest.raises(ValidationError):
            ExtractedTextBlock.model_validate(block)
    block = result.blocks[0].model_dump()
    block["location"]["end"] = 7
    with pytest.raises(ValidationError):
        ExtractedTextBlock.model_validate(block)


@pytest.mark.parametrize("rotation,crop,expected", [(90, 0, (0.4, 0.1, 0.4, 0.3)), (0, 0.2, (0.0, 0.2, 0.25, 0.4))])
def test_pptx_rotated_cropped_image_mapping(tmp_path: Path, rotation: int, crop: float, expected: tuple) -> None:
    path = tmp_path / "transform.pptx"
    image = tmp_path / "scan.png"
    _write_image(image)
    document = Presentation()
    document.slide_width = document.slide_height = 1000000
    picture = document.slides.add_slide(document.slide_layouts[6]).shapes.add_picture(str(image), 0, 0, width=1000000, height=1000000)
    picture.rotation = rotation
    picture.crop_left = crop
    document.save(path)
    result = extract_pptx(path, artifact_id=uuid4(), filename=path.name, ocr_provider=FakeOCR())
    block = next(b for b in result.blocks if b.extraction_method == "ocr")
    assert block.location.bbox == pytest.approx(expected)


def test_pptx_grouped_image_fails_explicitly(tmp_path: Path) -> None:
    path = tmp_path / "group.pptx"
    image = tmp_path / "scan.png"
    _write_image(image)
    document = Presentation()
    document.slides.add_slide(document.slide_layouts[6]).shapes.add_group_shape().shapes.add_picture(str(image), 0, 0)
    document.save(path)
    with pytest.raises(ProcessingError) as error:
        extract_pptx(path, artifact_id=uuid4(), filename=path.name, ocr_provider=FakeOCR())
    assert error.value.code == ProcessingErrorCode.UNSUPPORTED_FILE


def test_pptx_slide_limit(tmp_path: Path) -> None:
    path = tmp_path / "slides.pptx"
    document = Presentation()
    for _ in range(2):
        document.slides.add_slide(document.slide_layouts[6])
    document.save(path)
    with pytest.raises(ProcessingError) as error:
        extract_pptx(path, artifact_id=uuid4(), filename=path.name, ocr_provider=FakeOCR(), max_slides=1)
    assert error.value.code == ProcessingErrorCode.UNSUPPORTED_FILE


def test_standalone_ocr_failure_is_safe(tmp_path: Path) -> None:
    path = tmp_path / "scan.png"
    _write_image(path)
    class Broken:
        def recognize_image(self, *args: object, **kwargs: object) -> list[OCRBlock]:
            raise RuntimeError("SECRET")
    with pytest.raises(ProcessingError) as error:
        extract_image(path, artifact_id=uuid4(), filename=path.name, ocr_provider=Broken())
    assert error.value.code == ProcessingErrorCode.OCR_FAILED
    assert "SECRET" not in str(error.value)


def test_corrupt_image_rejected_before_mock_ocr(tmp_path: Path) -> None:
    path = tmp_path / "scan.png"
    path.write_bytes(b"\x89PNG\r\n\x1a\ninvalid")
    ocr = FakeOCR()
    with pytest.raises(ProcessingError) as error:
        extract_image(path, artifact_id=uuid4(), filename=path.name, ocr_provider=ocr)
    assert error.value.code == ProcessingErrorCode.EXTRACTION_FAILED
    assert not ocr.seen_paths


def test_render_limit_is_explicit_and_configurable(tmp_path: Path) -> None:
    path = tmp_path / "large.pdf"
    pdf = canvas.Canvas(str(path), pagesize=(200, 100))
    pdf.rect(20, 20, 20, 20, fill=1)
    pdf.save()
    with pytest.raises(ProcessingError) as error:
        extract_pdf(path, artifact_id=uuid4(), filename=path.name, ocr_provider=FakeOCR(), max_render_edge=399)
    assert error.value.code == ProcessingErrorCode.UNSUPPORTED_FILE
    assert extract_pdf(path, artifact_id=uuid4(), filename=path.name, ocr_provider=FakeOCR(), max_render_edge=400).blocks


def test_pdf_renderer_handles_closed_on_image_encoding_failure(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "vector.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.rect(10, 10, 20, 20, fill=1)
    pdf.save()
    import pii_guard.extraction.pdf as module
    closed = []
    original_page_close = module.pdfium.PdfPage.close
    original_bitmap_close = module.pdfium.PdfBitmap.close
    def page_close(instance: object) -> None:
        closed.append("page")
        original_page_close(instance)
    def bitmap_close(instance: object) -> None:
        closed.append("bitmap")
        original_bitmap_close(instance)
    def broken_save(*args: object, **kwargs: object) -> None:
        raise RuntimeError("SECRET")
    monkeypatch.setattr(module.pdfium.PdfPage, "close", page_close)
    monkeypatch.setattr(module.pdfium.PdfBitmap, "close", bitmap_close)
    monkeypatch.setattr(Image.Image, "save", broken_save)
    with pytest.raises(ProcessingError) as error:
        extract_pdf(path, artifact_id=uuid4(), filename=path.name, ocr_provider=FakeOCR())
    assert error.value.code == ProcessingErrorCode.OCR_FAILED
    assert {"page", "bitmap"} <= set(closed)
    path.unlink()  # The document handle must also be closed on Windows.


def test_docx_comments_footnotes_endnotes_and_external_images(tmp_path: Path) -> None:
    from docx.opc.part import Part
    from docx.opc.packuri import PackURI
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    path = tmp_path / "ancillary.docx"
    document = Document()
    for name in ("comments", "footnotes", "endnotes"):
        data = f'<w:{name} xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:p><w:r><w:t>{name.upper()}</w:t></w:r></w:p></w:{name}>'.encode()
        part = Part(PackURI(f"/word/{name}.xml"), f"application/vnd.openxmlformats-officedocument.wordprocessingml.{name}+xml", data, document.part.package)
        document.part.relate_to(part, getattr(RT, name.upper()))
    document.save(path)
    result = extract_docx(path, artifact_id=uuid4(), filename=path.name, ocr_provider=FakeOCR())
    assert all(name in result.full_text for name in ("COMMENTS", "FOOTNOTES", "ENDNOTES"))
    document.part.relate_to("https://example.invalid/private.png", RT.IMAGE, is_external=True)
    document.save(path)
    with pytest.raises(ProcessingError) as error:
        extract_docx(path, artifact_id=uuid4(), filename=path.name, ocr_provider=FakeOCR())
    assert error.value.code == ProcessingErrorCode.UNSUPPORTED_FILE


def test_pptx_notes_image_has_image_local_coordinates(tmp_path: Path) -> None:
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    path = tmp_path / "notes.pptx"
    image = tmp_path / "image.png"
    _write_image(image)
    document = Presentation()
    slide = document.slides.add_slide(document.slide_layouts[6])
    image_part, relationship_id = slide.part.get_or_add_image_part(str(image))
    slide.notes_slide.part.relate_to(image_part, RT.IMAGE)
    slide.part.drop_rel(relationship_id)
    document.save(path)
    result = extract_pptx(path, artifact_id=uuid4(), filename=path.name, ocr_provider=FakeOCR())
    block = next(b for b in result.blocks if b.extraction_method == "ocr")
    assert "notesSlides" in block.location.source_ref
    assert block.location.slide_number is None
    assert block.location.bbox == (0.1, 0.2, 0.3, 0.4)
