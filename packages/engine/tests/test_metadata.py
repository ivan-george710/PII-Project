"""Actual ancillary-content fixtures: inspect every member/object after rewriting."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile
import posixpath
import traceback

import pymupdf
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree as ET
from PIL import Image, PngImagePlugin, ImageDraw
from pptx import Presentation
from pypdf import PdfReader, PdfWriter
from pypdf.generic import (
    ArrayObject, ByteStringObject, DecodedStreamObject, DictionaryObject,
    NameObject, NullObject, TextStringObject,
)

from pii_guard.errors import ProcessingError, ProcessingErrorCode
from pii_guard.redaction import redact_image, redact_pdf, redact_pdf_file
from pii_guard.redaction.metadata import (
    AncillaryClass, AncillaryDecision, AncillaryReport, classify_office_part,
    sanitize_office_metadata, validate_pdf_metadata, verify_encoded_image_metadata,
)
from pii_guard.redaction.office import pack, template, unpack
from test_redaction import _office_case, _office_renderer, _pdf_case, _raster, _zip_change


MARKER = 'META_PRIVATE_710'
_R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
_CT = 'http://schemas.openxmlformats.org/package/2006/content-types'
_PKGREL = 'http://schemas.openxmlformats.org/package/2006/relationships'
_CUSTOM = 'http://schemas.openxmlformats.org/officeDocument/2006/custom-properties'
_VT = 'http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes'


def _assert_safe_failure(call, caplog):
    with pytest.raises(ProcessingError) as error:
        call()
    assert error.value.code == ProcessingErrorCode.REDACTION_FAILED
    assert MARKER not in ''.join(traceback.format_exception(error.value))
    assert MARKER not in caplog.text


def _add_part(parts, name, data, content_type, *, owner='', reltype=None):
    parts[name] = data
    ct = ET.fromstring(parts['[Content_Types].xml'])
    ET.SubElement(ct, f'{{{_CT}}}Override', PartName='/' + name, ContentType=content_type)
    parts['[Content_Types].xml'] = ET.tostring(ct)
    if reltype:
        relname = '_rels/.rels' if not owner else str(Path(owner).parent).replace('\\','/') + '/_rels/' + Path(owner).name + '.rels'
        rels = ET.fromstring(parts[relname])
        target = name if not owner else posixpath.relpath(name, posixpath.dirname(owner))
        ET.SubElement(rels, f'{{{_PKGREL}}}Relationship', Id='rId999', Type=reltype, Target=target)
        parts[relname] = ET.tostring(rels)


@pytest.mark.parametrize('format', ['PNG', 'JPEG'])
def test_image_metadata_all_removed_and_orientation_baked(format, caplog):
    image = Image.new('RGB', (24, 12), 'white')
    ImageDraw.Draw(image).rectangle((0,0,8,5), fill='red')
    exif = Image.Exif(); exif[270] = MARKER; exif[274] = 6; exif[315] = MARKER
    b = BytesIO()
    if format == 'PNG':
        meta = PngImagePlugin.PngInfo()
        meta.add_text('Description', MARKER)
        meta.add_text('Comment', MARKER, zip=True)
        meta.add_itxt('XML:com.adobe.xmp', MARKER)
        image.save(b, format=format, exif=exif, pnginfo=meta, icc_profile=MARKER.encode())
    else:
        image.save(b, format=format, exif=exif, icc_profile=MARKER.encode(), comment=MARKER.encode())
        # A genuine JPEG APP1 XMP record independent of EXIF/ICC and a COM record.
        original = b.getvalue(); payload=b'http://ns.adobe.com/xap/1.0/\x00'+MARKER.encode()
        app=b'\xff\xe1'+(len(payload)+2).to_bytes(2,'big')+payload
        comment=b'\xff\xfe'+(len(MARKER)+2).to_bytes(2,'big')+MARKER.encode()
        b=BytesIO(original[:2]+app+comment+original[2:])
    result = redact_image(b.getvalue(), artifact_id=uuid4())
    verify_encoded_image_metadata(result.data, format)
    assert MARKER.encode() not in result.data
    with Image.open(BytesIO(result.data)) as out:
        assert out.size == (12,24) and not out.getexif()
        assert not set(out.info)-{'jfif','jfif_version','jfif_unit','jfif_density'}
        assert out.convert('RGB').getpixel((9,3))[0] > 180  # Rotated red region.
    assert result.warnings == ('Image metadata was removed.',)
    assert MARKER not in caplog.text and MARKER not in '\n'.join(result.warnings)


def test_transparent_hidden_rgb_is_cleared_without_visible_pixel_changes():
    image = Image.new('RGBA',(3,1))
    image.putdata([(100,22,45,0),(12,34,56,128),(4,5,6,255)])
    b=BytesIO();image.save(b,format='PNG')
    result=redact_image(b.getvalue(),artifact_id=uuid4())
    with Image.open(BytesIO(result.data)) as out:
        assert out.getpixel((0,0))==(0,0,0,0)
        assert out.getpixel((1,0))==(12,34,56,128)
        assert out.getpixel((2,0))==(4,5,6,255)


@pytest.mark.parametrize('format', ['PNG','JPEG'])
def test_encoded_metadata_guard_rejects_actual_metadata_bearing_files(format):
    image=Image.new('RGB',(10,10),'white');b=BytesIO()
    if format=='PNG':
        info=PngImagePlugin.PngInfo();info.add_text('Comment',MARKER);image.save(b,format=format,pnginfo=info)
    else:
        exif=Image.Exif();exif[270]=MARKER;image.save(b,format=format,exif=exif)
    with pytest.raises(ValueError) as caught: verify_encoded_image_metadata(b.getvalue(),format)
    assert MARKER not in str(caught.value)


def _pdf_bytes(writer):
    b=BytesIO();writer.write(b);return b.getvalue()


def _xmp(data=None):
    stream=DecodedStreamObject()
    stream.set_data(data or (f'<?xpacket begin=""?><x:xmpmeta xmlns:x="adobe:ns:meta/">'
        f'<rdf:RDF xmlns:rdf="http://www.w3.org/1999/02/22-rdf-syntax-ns#"><rdf:Description '
        f'xmlns:dc="http://purl.org/dc/elements/1.1/" dc:title="{MARKER}"/></rdf:RDF>'
        '</x:xmpmeta><?xpacket end="w"?>').encode())
    stream[NameObject('/Type')]=NameObject('/Metadata');stream[NameObject('/Subtype')]=NameObject('/XML')
    return stream


def _assert_pdf_no_marker(data):
    reader=PdfReader(BytesIO(data),strict=True)
    assert not reader.metadata
    owners=[reader.trailer['/Root'],*reader.pages]
    assert all('/Metadata' not in owner or isinstance(owner['/Metadata'],NullObject) for owner in owners)
    assert MARKER not in str(reader.trailer)
    assert MARKER.encode() not in data
    with pymupdf.open(stream=data,filetype='pdf') as doc:
        assert not doc.is_repaired and not doc.get_xml_metadata()
        for xref in range(1,doc.xref_length()):
            assert MARKER not in doc.xref_object(xref)
            if doc.xref_is_stream(xref): assert MARKER.encode() not in doc.xref_stream(xref)


@pytest.mark.parametrize('where', ['document','page','both'])
@pytest.mark.parametrize('compressed', [False, True])
def test_pdf_info_xmp_ids_and_obsolete_objects_permanently_removed(tmp_path,where,compressed,caplog):
    path,aid,blocks,targets=_pdf_case(tmp_path,image=False)
    original=path.read_bytes();writer=PdfWriter(clone_from=PdfReader(path))
    writer.add_metadata({'/Title':MARKER,'/Author':MARKER,'/PrivatePayload':MARKER})
    stream=_xmp();stream=stream.flate_encode() if compressed else stream
    ref=writer._add_object(stream)
    if where in {'document','both'}: writer._root_object[NameObject('/Metadata')]=ref
    if where in {'page','both'}: writer.pages[0][NameObject('/Metadata')]=ref
    writer._ID=ArrayObject([ByteStringObject(MARKER.encode()),ByteStringObject(MARKER.encode())])
    orphan=DecodedStreamObject();orphan.set_data(MARKER.encode());writer._add_object(orphan)
    data=_pdf_bytes(writer);assert MARKER in str(PdfReader(BytesIO(data)).metadata)
    report=validate_pdf_metadata(PdfReader(BytesIO(data)))
    assert {d.structure for d in report.decisions}=={'pdf_info','pdf_xmp'}
    assert MARKER not in report.model_dump_json() and MARKER not in '\n'.join(report.warnings)
    output=redact_pdf(data,targets,artifact_id=aid,blocks=blocks)
    _assert_pdf_no_marker(output)
    assert 'keep' in PdfReader(BytesIO(output)).pages[0].extract_text()
    assert path.read_bytes()==original and MARKER not in caplog.text


@pytest.mark.parametrize('feature',['attachment','annotation','javascript','layers','form','page_action','alternate_image','hidden_stream'])
def test_pdf_unsupported_ancillary_content_rejected(tmp_path,feature,caplog):
    path,aid,blocks,targets=_pdf_case(tmp_path,image=feature=='alternate_image')
    writer=PdfWriter(clone_from=PdfReader(path))
    if feature=='attachment': writer.add_attachment('payload.txt',MARKER.encode())
    elif feature=='annotation':
        annot=DictionaryObject({NameObject('/Type'):NameObject('/Annot'),NameObject('/Subtype'):NameObject('/Text'),NameObject('/Contents'):TextStringObject(MARKER)})
        writer.pages[0][NameObject('/Annots')]=ArrayObject([writer._add_object(annot)])
    elif feature=='alternate_image':
        image=next(iter(writer.pages[0]['/Resources']['/XObject'].values())).get_object()
        image[NameObject('/Alternates')]=ArrayObject([TextStringObject(MARKER)])
    elif feature=='hidden_stream':
        stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 3 Tr 30 100 Td ('+MARKER.encode()+b') Tj ET')
        writer.pages[0][NameObject('/Contents')]=writer._add_object(stream)
    else:
        key={'javascript':'/OpenAction','layers':'/OCProperties','form':'/AcroForm','page_action':'/AA'}[feature]
        owner=writer.pages[0] if feature=='page_action' else writer._root_object
        owner[NameObject(key)]=DictionaryObject({NameObject('/PrivatePayload'):TextStringObject(MARKER)})
    data=_pdf_bytes(writer)
    _assert_safe_failure(lambda:redact_pdf(data,targets,artifact_id=aid,blocks=blocks),caplog)


@pytest.mark.parametrize('bad',['xml','root','dtd','subtype','payload','size','compressed_size','filter','nested_info'])
def test_pdf_malformed_or_unbounded_metadata_fails_closed(tmp_path,bad,caplog):
    path,aid,blocks,targets=_pdf_case(tmp_path,image=False)
    writer=PdfWriter(clone_from=PdfReader(path));meta=_xmp()
    if bad=='xml': meta.set_data(b'<broken>'+MARKER.encode())
    elif bad=='root': meta.set_data(b'<other>'+MARKER.encode()+b'</other>')
    elif bad=='dtd': meta.set_data(b'<!DOCTYPE x [<!ENTITY private "'+MARKER.encode()+b'">]><x/>')
    elif bad=='subtype': meta[NameObject('/Subtype')]=NameObject('/PrivatePayload')
    elif bad=='payload': meta[NameObject('/PrivatePayload')]=TextStringObject(MARKER)
    elif bad in {'size','compressed_size'}:
        meta.set_data(b'<x:xmpmeta xmlns:x="adobe:ns:meta/">'+b'X'*(1024*1024+1)+b'</x:xmpmeta>')
        if bad=='compressed_size': meta=meta.flate_encode()
    elif bad=='filter': meta[NameObject('/Filter')]=NameObject('/Unsupported')
    else: writer._info=writer._add_object(DictionaryObject({NameObject('/PrivatePayload'):ArrayObject([TextStringObject(MARKER)])}))
    writer._root_object[NameObject('/Metadata')]=writer._add_object(meta)
    _assert_safe_failure(lambda:redact_pdf(_pdf_bytes(writer),targets,artifact_id=aid,blocks=blocks),caplog)


def _office_metadata_fixture(parts,kind):
    core=ET.fromstring(parts['docProps/core.xml'])
    for node in core:
        if ET.QName(node).localname in {'title','creator','lastModifiedBy','description','keywords'}: node.text=MARKER
    parts['docProps/core.xml']=ET.tostring(core)
    app=ET.fromstring(parts['docProps/app.xml'])
    ns='http://schemas.openxmlformats.org/officeDocument/2006/extended-properties'
    ET.SubElement(app,f'{{{ns}}}Company').text=MARKER
    parts['docProps/app.xml']=ET.tostring(app)
    custom=(f'<Properties xmlns="{_CUSTOM}" xmlns:vt="{_VT}"><property '
        f'fmtid="{{D5CDD505-2E9C-101B-9397-08002B2CF9AE}}" pid="2" name="{MARKER}">'
        f'<vt:lpwstr>{MARKER}</vt:lpwstr></property></Properties>').encode()
    _add_part(parts,'docProps/custom.xml',custom,'application/vnd.openxmlformats-officedocument.custom-properties+xml',reltype=_R+'/custom-properties')
    parts['docProps/thumbnail.jpeg']=MARKER.encode()  # Non-rendered preview is discarded, never decoded.
    if kind=='pptx': parts['ppt/printerSettings/printerSettings1.bin']=MARKER.encode()
    main='word/document.xml' if kind=='docx' else 'ppt/slides/slide1.xml'
    root=ET.fromstring(parts[main])
    for node in root.iter():
        if ET.QName(node).localname in {'docPr','cNvPr'}:
            node.set('name',MARKER);node.set('descr',MARKER);node.set('title',MARKER)
    parts[main]=ET.tostring(root)
    for name in list(parts):
        if '/media/' in name:
            with Image.open(BytesIO(parts[name])) as im:
                info=PngImagePlugin.PngInfo();info.add_text('PrivatePayload',MARKER)
                b=BytesIO();im.save(b,format='PNG',pnginfo=info);parts[name]=b.getvalue()


@pytest.mark.parametrize('kind',['docx','pptx'])
def test_office_all_supported_metadata_removed_every_zip_member_scanned(tmp_path,kind,caplog):
    path,aid,blocks,targets=_office_case(tmp_path,kind)
    data=_zip_change(path.read_bytes(),lambda parts:_office_metadata_fixture(parts,kind))
    path.write_bytes(data);original=path.read_bytes()
    parts=unpack(data);report=sanitize_office_metadata(parts,kind)
    assert AncillaryClass.STRIPPED in {d.disposition for d in report.decisions}
    assert AncillaryClass.SANITIZED in {d.disposition for d in report.decisions}
    assert MARKER not in report.model_dump_json() and MARKER not in '\n'.join(report.warnings)
    output=tmp_path/f'output.{kind}'
    _office_renderer(kind,True)(path,output,targets,artifact_id=aid,blocks=blocks)
    assert path.read_bytes()==original
    with ZipFile(output) as z:
        for name in z.namelist():
            assert MARKER not in name
            assert MARKER.encode() not in z.read(name)
        assert not any(n.startswith('docProps/') or '/printerSettings/' in n for n in z.namelist())
    if kind=='docx': assert '[REDACTED]' in Document(output).paragraphs[0].text
    else:
        obj=Presentation(output)
        assert '[REDACTED]' in obj.slides[0].shapes[0].text
        assert obj.slides[0].notes_slide.notes_text_frame.text=='[REDACTED] notes'
    assert MARKER not in caplog.text
    assert not list(tmp_path.glob('.pii-redaction-*'))


@pytest.mark.parametrize('kind',['docx','pptx'])
@pytest.mark.parametrize('feature',['comments','custom_xml','embedded','macros','external','alternate','fields','hidden'])
def test_office_unsupported_hidden_structures_rejected_even_no_targets(tmp_path,kind,feature,caplog):
    path,aid,blocks,targets=_office_case(tmp_path,kind,pictures=False)
    prefix='word' if kind=='docx' else 'ppt'
    main='word/document.xml' if kind=='docx' else 'ppt/slides/slide1.xml'
    def change(parts):
        if feature=='comments':
            _add_part(parts,prefix+'/comments/comment1.xml',f'<comments>{MARKER}</comments>'.encode(),
                'application/vnd.openxmlformats-officedocument.'+('wordprocessingml.comments+xml' if kind=='docx' else 'presentationml.comments+xml'),owner=main,reltype=_R+'/comments')
        elif feature=='custom_xml':
            # Modify the stock empty bibliography or add arbitrary custom XML.
            if kind=='docx': parts['customXml/item1.xml']=f'<private>{MARKER}</private>'.encode()
            else: _add_part(parts,'customXml/item1.xml',f'<private>{MARKER}</private>'.encode(),'application/xml',owner=main,reltype=_R+'/customXml')
        elif feature in {'embedded','macros'}:
            name=prefix+('/embeddings/oleObject1.bin' if feature=='embedded' else '/vbaProject.bin')
            _add_part(parts,name,MARKER.encode(),'application/vnd.ms-office.'+('oleObject' if feature=='embedded' else 'vbaProject'),owner=main,reltype=_R+('/oleObject' if feature=='embedded' else '/vbaProject'))
        elif feature=='external':
            relname=prefix+'/_rels/document.xml.rels' if kind=='docx' else 'ppt/slides/_rels/slide1.xml.rels'
            root=ET.fromstring(parts[relname]);ET.SubElement(root,f'{{{_PKGREL}}}Relationship',Id='rId999',Type=_R+'/hyperlink',Target='https://example.invalid/'+MARKER,TargetMode='External');parts[relname]=ET.tostring(root)
        else:
            root=ET.fromstring(parts[main])
            if feature=='hidden':
                if kind=='docx': root.find('.//'+qn('w:r')).insert(0,ET.fromstring('<w:rPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:vanish/></w:rPr>'))
                else: root.set('show','0')
            elif feature=='fields':
                if kind=='docx': ET.SubElement(root.find('.//'+qn('w:p')),qn('w:fldSimple'),{qn('w:instr'):MARKER})
                else: ET.SubElement(root.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}p'),'{http://schemas.openxmlformats.org/drawingml/2006/main}fld',id=MARKER,type='datetime')
            else:
                mc='http://schemas.openxmlformats.org/markup-compatibility/2006'
                alternate=ET.SubElement(root,f'{{{mc}}}AlternateContent');ET.SubElement(alternate,f'{{{mc}}}Fallback').text=MARKER
            parts[main]=ET.tostring(root)
    data=_zip_change(path.read_bytes(),change)
    _assert_safe_failure(lambda:_office_renderer(kind)(data,artifact_id=aid),caplog)


@pytest.mark.parametrize('feature',['header','footer','footnote','endnote','revision','variable'])
def test_docx_ancillary_text_and_revisions_are_explicitly_rejected(tmp_path,feature,caplog):
    path,aid,_,_=_office_case(tmp_path,'docx',pictures=False)
    doc=Document(path)
    if feature=='header': doc.sections[0].header.paragraphs[0].text=MARKER
    elif feature=='footer': doc.sections[0].footer.paragraphs[0].text=MARKER
    elif feature=='revision':
        p=doc.paragraphs[0]._p
        revision=OxmlElement('w:ins');revision.set(qn('w:author'),MARKER)
        revision.append(p[0]);p.append(revision)
    elif feature=='variable':
        variables=OxmlElement('w:docVars');variable=OxmlElement('w:docVar');variable.set(qn('w:name'),MARKER);variable.set(qn('w:val'),MARKER);variables.append(variable);doc.settings.element.append(variables)
    if feature not in {'footnote','endnote'}:
        b=BytesIO();doc.save(b);data=b.getvalue()
    else:
        name='word/'+feature+'s.xml'
        def add(parts): _add_part(parts,name,f'<private>{MARKER}</private>'.encode(),
            'application/vnd.openxmlformats-officedocument.wordprocessingml.'+feature+'s+xml',owner='word/document.xml',reltype=_R+'/'+feature+'s')
        data=_zip_change(path.read_bytes(),add)
    _assert_safe_failure(lambda:_office_renderer('docx')(data,artifact_id=aid),caplog)


@pytest.mark.parametrize('feature',['transition','action','notes_metadata'])
def test_pptx_unknown_links_transitions_notes_payloads_rejected(tmp_path,feature,caplog):
    path,aid,_,_=_office_case(tmp_path,'pptx',pictures=False)
    def change(parts):
        name='ppt/notesSlides/notesSlide1.xml' if feature=='notes_metadata' else 'ppt/slides/slide1.xml'
        root=ET.fromstring(parts[name]);ns='http://schemas.openxmlformats.org/presentationml/2006/main'
        if feature=='transition': ET.SubElement(root,f'{{{ns}}}transition',advTm=MARKER)
        elif feature=='action': ET.SubElement(root.find('.//'+f'{{{ns}}}cNvPr'),'{http://schemas.openxmlformats.org/drawingml/2006/main}hlinkClick',action=MARKER)
        else: ET.SubElement(root,f'{{{ns}}}extLst').text=MARKER
        parts[name]=ET.tostring(root)
    _assert_safe_failure(lambda:_office_renderer('pptx')(_zip_change(path.read_bytes(),change),artifact_id=aid),caplog)


@pytest.mark.parametrize('kind',['docx','pptx'])
@pytest.mark.parametrize('bad',['broken','wrong_root','bad_custom','bad_reltype','owned_payload','bad_content_type'])
def test_malformed_office_ancillary_structures_fail_closed(tmp_path,kind,bad,caplog):
    path,aid,_,_=_office_case(tmp_path,kind,pictures=False)
    def change(parts):
        if bad=='broken': parts['docProps/core.xml']=b'<broken>'+MARKER.encode()
        elif bad=='wrong_root': parts['docProps/core.xml']=f'<private>{MARKER}</private>'.encode()
        elif bad=='bad_custom': _add_part(parts,'docProps/custom.xml',f'<private>{MARKER}</private>'.encode(),'application/vnd.openxmlformats-officedocument.custom-properties+xml',reltype=_R+'/custom-properties')
        elif bad=='bad_reltype':
            root=ET.fromstring(parts['_rels/.rels'])
            next(n for n in root if n.get('Target')=='docProps/core.xml').set('Type',_R+'/oleObject')
            parts['_rels/.rels']=ET.tostring(root)
        elif bad=='owned_payload':
            root=ET.Element(f'{{{_PKGREL}}}Relationships',nsmap={None:_PKGREL})
            ET.SubElement(root,f'{{{_PKGREL}}}Relationship',Id='rId1',Type=_R+'/image',Target='thumbnail.jpeg')
            parts['docProps/_rels/core.xml.rels']=ET.tostring(root)
        else:
            root=ET.fromstring(parts['[Content_Types].xml'])
            next(n for n in root if n.get('PartName')=='/docProps/core.xml').set('ContentType','application/xml')
            parts['[Content_Types].xml']=ET.tostring(root)
    _assert_safe_failure(lambda:_office_renderer(kind)(_zip_change(path.read_bytes(),change),artifact_id=aid),caplog)


@pytest.mark.parametrize('kind',['docx','pptx'])
def test_stock_structural_classification_has_content_evidence(kind):
    stock=template(kind)
    editable={'word/document.xml'} if kind=='docx' else {'ppt/slides/slide1.xml','ppt/notesSlides/notesSlide1.xml'}
    parts=dict(stock);report=sanitize_office_metadata(parts,kind)
    assert report.warnings
    for name,data in parts.items():
        decision=classify_office_part(name,data,stock,editable,kind)
        assert decision.disposition!=AncillaryClass.UNSUPPORTED
        if name not in editable and name!='ppt/presentation.xml' and not name.endswith('.rels') and name!='[Content_Types].xml':
            assert data==stock[name]
            assert classify_office_part(name,data+b' '+MARKER.encode(),stock,editable,kind).disposition==AncillaryClass.UNSUPPORTED
    assert classify_office_part('unknown.bin',MARKER.encode(),stock,editable,kind).disposition==AncillaryClass.UNSUPPORTED


@pytest.mark.parametrize('kind',['docx','pptx','pdf'])
@pytest.mark.parametrize('failure',['metadata','publish'])
def test_metadata_failure_and_publication_failure_leave_original_and_no_partial_output(tmp_path,monkeypatch,kind,failure,caplog):
    if kind=='pdf':
        path,aid,blocks,targets=_pdf_case(tmp_path,image=False)
        if failure=='metadata':
            writer=PdfWriter(clone_from=PdfReader(path));writer._root_object[NameObject('/Metadata')]=writer._add_object(_xmp(b'<broken>'))
            path.write_bytes(_pdf_bytes(writer))
        renderer=redact_pdf_file
    else:
        path,aid,blocks,targets=_office_case(tmp_path,kind,pictures=False)
        if failure=='metadata': path.write_bytes(_zip_change(path.read_bytes(),lambda parts:parts.update({'docProps/core.xml':b'<broken>'})))
        renderer=_office_renderer(kind,True)
    original=path.read_bytes();output=tmp_path/f'output.{kind}'
    if failure=='publish':
        def reject(*args): raise OSError(MARKER)
        monkeypatch.setattr('os.link',reject)
    _assert_safe_failure(lambda:renderer(path,output,targets,artifact_id=aid,blocks=blocks),caplog)
    assert path.read_bytes()==original and not output.exists()
    assert not list(tmp_path.glob('.pii-redaction-*'))


def test_classification_and_warning_contract_contains_only_fixed_categories():
    report=AncillaryReport(decisions=tuple(AncillaryDecision(structure=structure,disposition=AncillaryClass.STRIPPED)
        for structure in ['office_properties','office_preview','office_printer_data','pdf_info','pdf_xmp','image_metadata']))
    assert len(report.warnings)==6 and MARKER not in report.model_dump_json()
    with pytest.raises(ValueError): AncillaryDecision(structure=MARKER,disposition=AncillaryClass.STRIPPED)


@pytest.mark.parametrize('key',['/PageMode','/PageLayout','/Version'])
@pytest.mark.parametrize('payload',['text','name','dictionary'])
def test_pdf_catalog_preferences_cannot_carry_hidden_payloads(tmp_path,key,payload,caplog):
    path,aid,_,_=_pdf_case(tmp_path,image=False)
    writer=PdfWriter(clone_from=PdfReader(path))
    value={'text':TextStringObject(MARKER),'name':NameObject('/'+MARKER),
        'dictionary':DictionaryObject({NameObject('/PrivatePayload'):TextStringObject(MARKER)})}[payload]
    writer._root_object[NameObject(key)]=value
    _assert_safe_failure(lambda:redact_pdf(_pdf_bytes(writer),artifact_id=aid),caplog)


@pytest.mark.parametrize('key,value',[('/PageMode','/UseNone'),('/PageLayout','/SinglePage'),('/Version','/1.7')])
def test_pdf_known_structural_catalog_preferences_remain_supported(tmp_path,key,value):
    path,aid,_,_=_pdf_case(tmp_path,image=False)
    writer=PdfWriter(clone_from=PdfReader(path))
    writer._root_object[NameObject(key)]=NameObject(value)
    output=redact_pdf(_pdf_bytes(writer),artifact_id=aid)
    assert PdfReader(BytesIO(output)).trailer['/Root'][key]==value


@pytest.mark.parametrize('feature',['catalog_type','font_name'])
@pytest.mark.parametrize('payload',['text','name'])
def test_pdf_other_structural_name_fields_reject_private_payloads(tmp_path,feature,payload,caplog):
    path,aid,_,_=_pdf_case(tmp_path,image=False)
    writer=PdfWriter(clone_from=PdfReader(path))
    value=TextStringObject(MARKER) if payload=='text' else NameObject('/'+MARKER)
    if feature=='catalog_type': writer._root_object[NameObject('/Type')]=value
    else:
        font=next(iter(writer.pages[0]['/Resources']['/Font'].values())).get_object()
        font[NameObject('/Name')]=value
    _assert_safe_failure(lambda:redact_pdf(_pdf_bytes(writer),artifact_id=aid),caplog)


def test_actual_word_comment_package_reopens_but_redaction_rejects_it(tmp_path,caplog):
    path,aid,_,_=_office_case(tmp_path,'docx',pictures=False)
    doc=Document(path)
    doc.add_comment(doc.paragraphs[0].runs[0],text=MARKER,author=MARKER,initials='MP')
    b=BytesIO();doc.save(b);data=b.getvalue()
    assert MARKER in next(iter(Document(BytesIO(data)).comments)).text
    _assert_safe_failure(lambda:_office_renderer('docx')(data,artifact_id=aid),caplog)


def test_actual_powerpoint_comment_package_reopens_but_is_rejected(tmp_path,caplog):
    path,aid,_,_=_office_case(tmp_path,'pptx',pictures=False)
    ns='http://schemas.openxmlformats.org/presentationml/2006/main'
    def add(parts):
        root=ET.Element(f'{{{ns}}}cmLst',nsmap={'p':ns})
        cm=ET.SubElement(root,f'{{{ns}}}cm',authorId='0',dt='2026-01-01T00:00:00Z',idx='1')
        ET.SubElement(cm,f'{{{ns}}}pos',x='0',y='0');ET.SubElement(cm,f'{{{ns}}}text').text=MARKER
        _add_part(parts,'ppt/comments/comment1.xml',ET.tostring(root),
            'application/vnd.openxmlformats-officedocument.presentationml.comments+xml',owner='ppt/slides/slide1.xml',reltype=_R+'/comments')
        authors=ET.Element(f'{{{ns}}}cmAuthorLst',nsmap={'p':ns})
        ET.SubElement(authors,f'{{{ns}}}cmAuthor',id='0',name=MARKER,initials='MP',lastIdx='1',clrIdx='0')
        _add_part(parts,'ppt/commentAuthors.xml',ET.tostring(authors),
            'application/vnd.openxmlformats-officedocument.presentationml.commentAuthors+xml',owner='ppt/presentation.xml',reltype=_R+'/commentAuthors')
    data=_zip_change(path.read_bytes(),add)
    assert len(Presentation(BytesIO(data)).slides)==1
    _assert_safe_failure(lambda:_office_renderer('pptx')(data,artifact_id=aid),caplog)


@pytest.mark.parametrize('feature',['annotation','widget'])
def test_actual_pdf_annotations_and_form_widgets_rejected(tmp_path,feature,caplog):
    path,aid,_,_=_pdf_case(tmp_path,image=False)
    with pymupdf.open(path) as doc:
        page=doc[0]
        if feature=='annotation': page.add_text_annot((10,10),MARKER)
        else:
            widget=pymupdf.Widget();widget.field_name='private';widget.field_value=MARKER
            widget.field_type=pymupdf.PDF_WIDGET_TYPE_TEXT;widget.rect=pymupdf.Rect(20,20,150,60)
            page.add_widget(widget)
        data=doc.tobytes()
    assert PdfReader(BytesIO(data)).pages[0]['/Annots']
    _assert_safe_failure(lambda:redact_pdf(data,artifact_id=aid),caplog)


@pytest.mark.parametrize('kind',['docx','pptx'])
def test_office_metadata_removed_with_no_enabled_targets_and_no_event_mutation(tmp_path,kind):
    from pii_guard import BUILTIN_RULE_IDS, PolicySnapshot, CustomRule, PolicyEvaluator
    path,aid,blocks,_=_office_case(tmp_path,kind,pictures=False)
    data=_zip_change(path.read_bytes(),lambda parts:_office_metadata_fixture(parts,kind))
    policy=PolicySnapshot(organization_id=uuid4(),version=8,builtins={r:False for r in BUILTIN_RULE_IDS},
        custom_rules=[CustomRule(id=uuid4(),name='disabled',kind='exact',pattern='SECRET',case_sensitive=True,enabled=False)])
    evaluated=PolicyEvaluator(policy).evaluate(blocks=blocks)
    before=[f.model_dump_json() for f in evaluated.findings]
    assert before and not evaluated.redaction_targets
    output=_office_renderer(kind)(data,evaluated.redaction_targets,artifact_id=aid,blocks=blocks)
    assert [f.model_dump_json() for f in evaluated.findings]==before
    with ZipFile(BytesIO(output)) as z:
        assert all(MARKER.encode() not in z.read(name) for name in z.namelist())
    obj=Document(BytesIO(output)) if kind=='docx' else Presentation(BytesIO(output))
    assert 'SECRET' in (obj.paragraphs[0].text if kind=='docx' else obj.slides[0].shapes[0].text)
