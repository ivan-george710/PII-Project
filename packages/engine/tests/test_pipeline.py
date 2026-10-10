from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile

import pytest
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches as WordInches
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.utils import ImageReader

from pii_guard.pipeline import PrivacyEngine
from pii_guard.models import SanitizeRequest, ArtifactInput, PolicySnapshot
from pii_guard.rules.builtins import DEFAULT_BUILTIN_RULES
from pii_guard.errors import ProcessingError


@pytest.fixture(scope="module")
def real_engine():
    return PrivacyEngine()


def portrait_scan():
    with Image.open(Path(__file__).parent / "fixtures/astronaut.png") as portrait:
        image = Image.new("RGB", (1000, 650), "white")
        image.paste(portrait, (0, 0))
    font = ImageFont.truetype(
        "Arial.ttf" if Path("/System/Library/Fonts").exists() else "DejaVuSans.ttf", 32
    )
    ImageDraw.Draw(image).text(
        (30, 550), "Email: dana@example.com", font=font, fill="black"
    )
    output = BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


def make_document(kind, image):
    out = BytesIO()
    if kind == "png":
        return image, "image/png"
    if kind == "jpeg":
        with Image.open(BytesIO(image)) as im:
            im.save(out, format="JPEG", quality=95)
        return out.getvalue(), "image/jpeg"
    if kind == "pdf":
        canvas = Canvas(out, pagesize=(600, 600))
        canvas.drawImage(ImageReader(BytesIO(image)), 0, 0, width=600, height=390)
        canvas.drawString(30, 500, "Email: native@example.com")
        canvas.save()
        return out.getvalue(), "application/pdf"
    if kind == "docx":
        doc = Document()
        doc.add_paragraph("Email: native@example.com")
        doc.add_picture(BytesIO(image), width=WordInches(5))
        doc.save(out)
        return (
            out.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    ppt = Presentation()
    slide = ppt.slides.add_slide(ppt.slide_layouts[6])
    slide.shapes.add_picture(BytesIO(image), Inches(1), Inches(1), width=Inches(5))
    slide.shapes.add_textbox(
        Inches(1), Inches(0.2), Inches(6), Inches(0.5)
    ).text = "Email: native@example.com"
    ppt.save(out)
    return (
        out.getvalue(),
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )


@pytest.mark.parametrize("kind", ["png", "jpeg", "pdf", "docx", "pptx"])
def test_real_faces_scanned_text_and_verified_output(real_engine, tmp_path, kind):
    raw, mime = make_document(kind, portrait_scan())
    source = tmp_path / f"input.{kind}"
    source.write_bytes(raw)
    aid = uuid4()
    result = real_engine.sanitize(
        SanitizeRequest(
            submission_id=uuid4(),
            text="Contact message@example.com",
            artifacts=[
                ArtifactInput(
                    id=aid, filename=source.name, source_path=source, declared_mime=mime
                )
            ],
            policy=PolicySnapshot(
                organization_id=uuid4(),
                version=1,
                builtins=DEFAULT_BUILTIN_RULES,
                custom_rules=[],
            ),
            output_dir=tmp_path / "out",
        )
    )
    assert result.sanitized_text == "Contact [REDACTED]"
    assert any(f.rule_id == "FACE" and f.action == "redacted" for f in result.findings)
    assert any(
        f.rule_id == "EMAIL_ADDRESS" and f.location.kind == "ocr"
        for f in result.findings
    )
    assert "dana@example.com" not in result._artifact_texts[aid]
    assert "native@example.com" not in result._artifact_texts[aid]
    safe = result.artifacts[0].output_path.read_bytes()
    assert safe != raw
    if kind in {"png", "jpeg"}:
        images = [safe]
    elif kind in {"docx", "pptx"}:
        with ZipFile(BytesIO(safe)) as archive:
            images = [archive.read(n) for n in archive.namelist() if "/media/" in n]
    else:
        import pymupdf

        with pymupdf.open(stream=safe, filetype="pdf") as pdf:
            images = [pdf[0].get_pixmap().tobytes("png")]
    # Genuine re-detection on serialized/reopened sanitized pixels.
    from pii_guard.detection.faces import same_face

    originals = [f.location.bbox for f in result.findings if f.rule_id == "FACE"]
    # Black masks/text rectangles can create fresh false positives elsewhere;
    # the originally detected face must no longer be detectable in its region.
    for data in images:
        residual = real_engine.face_detector.detect(
            data, artifact_id=aid, source_ref="image"
        )
        assert not any(
            same_face(f.location.bbox, box) for f in residual for box in originals
        )


def test_disabled_faces_are_reported_and_not_masked(real_engine, tmp_path):
    source = Path(__file__).parent / "fixtures/astronaut.png"
    aid = uuid4()
    policy = {**DEFAULT_BUILTIN_RULES, "FACE": False}
    result = real_engine.sanitize(
        SanitizeRequest(
            submission_id=uuid4(),
            text="",
            artifacts=[
                ArtifactInput(
                    id=aid,
                    filename="face.png",
                    source_path=source,
                    declared_mime="image/png",
                )
            ],
            policy=PolicySnapshot(
                organization_id=uuid4(), version=1, builtins=policy, custom_rules=[]
            ),
            output_dir=tmp_path,
        )
    )
    assert any(
        f.rule_id == "FACE" and f.action == "allowed_disabled" for f in result.findings
    )
    assert real_engine.face_detector.detect(
        result.artifacts[0].output_path.read_bytes(),
        artifact_id=aid,
        source_ref="image",
    )


def test_multiple_faces_and_atomic_failure(real_engine, tmp_path, monkeypatch):
    with Image.open(Path(__file__).parent / "fixtures/astronaut.png") as portrait:
        image = Image.new("RGB", (1024, 512))
        image.paste(portrait, (0, 0))
        image.paste(portrait, (512, 0))
        out = BytesIO()
        image.save(out, format="PNG")
    faces = real_engine.face_detector.detect(
        out.getvalue(), artifact_id=uuid4(), source_ref="image"
    )
    # Known face centers in the two NASA portraits. Conservative detection
    # can include uncertain regions, so assert actual face coverage.
    for x, y in [(0.22, 0.23), (0.72, 0.23)]:
        assert any(
            f.location.bbox[0] <= x <= f.location.bbox[0] + f.location.bbox[2]
            and f.location.bbox[1] <= y <= f.location.bbox[1] + f.location.bbox[3]
            for f in faces
        )
    a = tmp_path / "group.png"
    a.write_bytes(out.getvalue())
    b = tmp_path / "bad.png"
    b.write_bytes(b"invalid")
    request = SanitizeRequest(
        submission_id=uuid4(),
        text="",
        artifacts=[
            ArtifactInput(
                id=uuid4(), filename=p.name, source_path=p, declared_mime="image/png"
            )
            for p in [a, b]
        ],
        policy=PolicySnapshot(
            organization_id=uuid4(),
            version=1,
            builtins=DEFAULT_BUILTIN_RULES,
            custom_rules=[],
        ),
        output_dir=tmp_path / "output",
    )
    with pytest.raises(ProcessingError):
        real_engine.sanitize(request)
    assert not (tmp_path / "output").exists() or not list(
        (tmp_path / "output").iterdir()
    )
