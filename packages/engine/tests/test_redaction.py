from __future__ import annotations

import itertools
import traceback
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pytest
import pymupdf
from PIL import Image, ImageDraw, ImageChops
from pypdf import PdfReader, PdfWriter
from pypdf.generic import NameObject, DictionaryObject, RectangleObject
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader

from pii_guard import BUILTIN_RULE_IDS, CustomRule, PolicyEvaluator, PolicySnapshot
from pii_guard.detection.validation import DetectionCandidate
from pii_guard.errors import ProcessingError, ProcessingErrorCode
from pii_guard.models import ExtractedTextBlock, SourceLocation
from pii_guard.redaction import redact_block, redact_text, redact_image, redact_pdf, redact_pdf_file
from pii_guard.extraction.pdf import extract_pdf
from pii_guard.ocr.engine import OCRBlock
from pii_guard.rules.policy import RedactionTarget
from pii_guard.redaction import redact_docx, redact_docx_file, redact_pptx, redact_pptx_file
from pii_guard.extraction.docx import extract_docx
from pii_guard.extraction.pptx import extract_pptx
from docx import Document
from docx.shared import Inches as WordInches
from pptx import Presentation
from pptx.util import Inches
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree as ET


def _location(start: int, end: int, **overrides) -> SourceLocation:
    data = dict(artifact_id=None, kind="message", page_number=None, slide_number=None,
                source_ref=None, start=start, end=end, bbox=None)
    data.update(overrides)
    return SourceLocation(**data)


def _target(start: int, end: int, **overrides) -> RedactionTarget:
    return RedactionTarget(finding_id=uuid4(), location=_location(start, end, **overrides))


@pytest.mark.parametrize("text,spans,expected", [
    ("before SECRET after", [(7, 13)], "before [REDACTED] after"),
    ("alpha, beta!", [(0, 5), (7, 11)], "[REDACTED], [REDACTED]!"),
    ("abcdefghij", [(1, 5), (3, 7)], "a[REDACTED]hij"),
    ("abcdefghij", [(1, 4), (4, 7)], "a[REDACTED]hij"),
    ("abcdefghij", [(1, 8), (2, 5), (3, 4)], "a[REDACTED]ij"),
    ("abcdefghij", [(1, 3), (2, 5), (5, 8)], "a[REDACTED]ij"),
    ("SECRET", [(0, 6), (0, 6)], "[REDACTED]"),
    ("one two", [(0, 3), (4, 7)], "[REDACTED] [REDACTED]"),
    ("", [], ""),
    (" \r\nunchanged\t\U0001f600!", [], " \r\nunchanged\t\U0001f600!"),
])
def test_text_rendering(text: str, spans: list, expected: str) -> None:
    assert redact_text(text, [_target(*span) for span in spans]) == expected


def test_unicode_combining_marks_and_multiline_preservation() -> None:
    text = "\U0001f600 e\u0301, SECRET\r\nsecond\n\U0001f469\u200d\U0001f4bb fin!"
    secrets = ("e\u0301", "SECRET\r\nsecond", "\U0001f469\u200d\U0001f4bb")
    targets = [_target(text.index(secret), text.index(secret) + len(secret)) for secret in secrets]
    assert redact_text(text, targets) == "\U0001f600 [REDACTED], [REDACTED]\n[REDACTED] fin!"
    assert all(secret not in redact_text(text, targets) for secret in secrets)


def test_offsets_are_codepoints_without_grapheme_expansion() -> None:
    # A target for only the accent must not consume the unrelated base character.
    assert redact_text("e\u0301!", [_target(1, 2)]) == "e[REDACTED]!"
    assert redact_text("\U0001f600X!", [_target(1, 2)]) == "\U0001f600[REDACTED]!"


def test_order_independent_rendering_does_not_mutate_targets() -> None:
    targets = [_target(1, 4), _target(3, 6), _target(6, 8), _target(9, 10)]
    before = [target.model_dump_json() for target in targets]
    for ordered in itertools.permutations(targets):
        assert redact_text("abcdefghijk", ordered) == "a[REDACTED]i[REDACTED]k"
    assert [target.model_dump_json() for target in targets] == before


@pytest.mark.parametrize("start,end", [(-1, 2), (0, 0), (3, 2), (0, 7), (6, 7),
                                         (True, 3), (0, False), ("0", 3), (0, 3.5), (None, 3)])
def test_malformed_or_out_of_bounds_offsets_fail_closed(start, end) -> None:
    # model_copy simulates a corrupted private target; the renderer revalidates it.
    target = _target(0, 3)
    target = target.model_copy(update={"location": target.location.model_copy(update={"start": start, "end": end})})
    with pytest.raises(ProcessingError) as error:
        redact_text("SECRET", [_target(0, 1), target])
    assert error.value.code == ProcessingErrorCode.REDACTION_FAILED
    assert "SECRET" not in str(error.value)


def test_empty_text_cannot_have_targets() -> None:
    with pytest.raises(ProcessingError):
        redact_text("", [_target(0, 1)])


@pytest.mark.parametrize("text", [None, b"SECRET", 123])
def test_non_string_original_text_rejected(text) -> None:
    with pytest.raises(ProcessingError):
        redact_text(text)


def _block(kind: str = "native_text") -> ExtractedTextBlock:
    text = "\U0001f600 SECRET, tail\n"
    return ExtractedTextBlock(text=text, extraction_method=kind, confidence=0.92,
        location=_location(0, len(text), artifact_id=uuid4(), kind=kind, page_number=2,
            slide_number=None, source_ref="unit:paragraph:7",
            bbox=(0.1, 0.2, 0.3, 0.4) if kind == "ocr" else None))


def _block_target(block: ExtractedTextBlock) -> RedactionTarget:
    location = block.location.model_copy(update={"start": 2, "end": 8})
    return RedactionTarget(finding_id=uuid4(), location=location)


@pytest.mark.parametrize("kind", ["native_text", "ocr"])
def test_blocks_preserve_original_provenance_and_do_not_mutate_inputs(kind: str) -> None:
    block = _block(kind)
    target = _block_target(block)
    before = block.model_dump_json(), target.model_dump_json()
    result = redact_block(block, [target])
    assert result.sanitized_text == "\U0001f600 [REDACTED], tail\n"
    assert result.original_location == block.location
    assert result.original_location.end == len(block.text)
    assert result.original_location.end != len(result.sanitized_text)
    assert result.extraction_method == kind
    assert result.confidence == 0.92
    assert "SECRET" not in result.model_dump_json()
    assert (block.model_dump_json(), target.model_dump_json()) == before
    # Result provenance is copied rather than aliased to mutable input models.
    result.original_location.source_ref = "changed"
    assert block.location.source_ref == "unit:paragraph:7"


@pytest.mark.parametrize("field,value", [
    ("artifact_id", uuid4()), ("kind", "native_text"), ("page_number", 3),
    ("slide_number", 4), ("source_ref", "different-unit"), ("bbox", (0.2, 0.2, 0.3, 0.4)),
])
def test_block_source_mismatches_are_errors(field: str, value) -> None:
    block = _block("ocr")
    target = _block_target(block)
    target.location = target.location.model_copy(update={field: value})
    with pytest.raises(ProcessingError) as error:
        redact_block(block, [target], submission_id=(submission_id := uuid4()))
    assert error.value.code == ProcessingErrorCode.REDACTION_FAILED
    assert error.value.submission_id == submission_id
    assert error.value.artifact_id == block.location.artifact_id


def test_message_and_block_targets_cannot_be_interchanged() -> None:
    block = _block()
    with pytest.raises(ProcessingError):
        redact_text(block.text, [_block_target(block)])
    with pytest.raises(ProcessingError):
        redact_block(block, [_target(2, 8)])


def test_block_without_targets_retains_text_and_source() -> None:
    block = _block()
    rendered = redact_block(block)
    assert rendered.sanitized_text == block.text
    assert rendered.original_location == block.location
    assert redact_text(block.text, source=block.location) == block.text


def test_malformed_target_contracts_are_rejected() -> None:
    target = _target(0, 6)
    for malformed in (target.model_dump(), None, target.model_copy(update={"finding_id": "invalid"})):
        with pytest.raises(ProcessingError) as error:
            redact_text("SECRET", [malformed])
        assert error.value.code == ProcessingErrorCode.REDACTION_FAILED


@pytest.mark.parametrize("field,value", [("start", 1), ("end", 99), ("source_ref", None)])
def test_invalid_original_block_is_revalidated_even_without_targets(field: str, value) -> None:
    block = _block()
    block = block.model_copy(update={"location": block.location.model_copy(update={field: value})})
    with pytest.raises(ProcessingError):
        redact_block(block)


def test_original_source_must_cover_text_and_message_metadata_must_match() -> None:
    with pytest.raises(ProcessingError):
        redact_text("SECRET", source=_location(1, 6))
    with pytest.raises(ProcessingError):
        redact_text("SECRET", source=_location(0, 6, page_number=1))
    assert redact_text("SECRET", [_target(0, 6)], source=_location(0, 6)) == "[REDACTED]"


def _candidate(start: int, end: int, rule_id: str = "PERSON") -> DetectionCandidate:
    return DetectionCandidate(rule_id=rule_id, category="identity" if rule_id == "PERSON" else "contact",
        confidence=0.95, recognizer_confidence=0.95, recognizer="Synthetic",
        location=_location(start, end))


@pytest.mark.parametrize("enabled", [True, False])
def test_policy_targets_control_redaction_and_public_events_stay_unchanged(enabled: bool) -> None:
    policy = PolicySnapshot(organization_id=uuid4(), version=7,
        builtins={rule: enabled for rule in BUILTIN_RULE_IDS}, custom_rules=[])
    policy.builtins["EMAIL_ADDRESS"] = False
    evaluation = PolicyEvaluator(policy).evaluate([
        _candidate(2, 8), _candidate(5, 11, "EMAIL_ADDRESS"), _candidate(3, 5)])
    before = evaluation.model_dump_json()
    targets_before = [target.model_dump_json() for target in evaluation.redaction_targets]
    rendered = redact_text("xxSECRETzzz!", evaluation.redaction_targets)
    assert rendered == ("xx[REDACTED]zzz!" if enabled else "xxSECRETzzz!")
    assert len(evaluation.findings) == 3
    assert evaluation.findings[1].action == "allowed_disabled"
    assert evaluation.model_dump_json() == before
    assert [target.model_dump_json() for target in evaluation.redaction_targets] == targets_before
    # A Finding cannot be passed accidentally as a rendering target.
    with pytest.raises(ProcessingError):
        redact_text("xxSECRETzzz!", evaluation.findings)


@pytest.mark.parametrize("kind", ["native_text", "ocr"])
def test_custom_policy_targets_render_original_extraction_blocks(kind: str) -> None:
    block = _block(kind)
    rule = CustomRule(id=uuid4(), name="Secret", kind="exact", pattern="SECRET",
        case_sensitive=True, enabled=True)
    policy = PolicySnapshot(organization_id=uuid4(), version=1,
        builtins={rule: True for rule in BUILTIN_RULE_IDS}, custom_rules=[rule])
    evaluation = PolicyEvaluator(policy).evaluate(blocks=[block])
    before = evaluation.model_dump_json()
    assert redact_block(block, evaluation.redaction_targets).sanitized_text == "\U0001f600 [REDACTED], tail\n"
    assert evaluation.model_dump_json() == before


def test_target_stream_error_is_safe_and_no_partial_success_is_returned(caplog) -> None:
    secret = "SECRET"
    def targets():
        yield _target(0, 6)
        raise ValueError("SECRET document content")

    with pytest.raises(ProcessingError) as error:
        redact_text(secret, targets())
    assert error.value.code == ProcessingErrorCode.REDACTION_FAILED
    assert "SECRET" not in "".join(traceback.format_exception(error.value))
    assert "SECRET" not in caplog.text


def test_face_targets_are_rejected_by_text_renderer() -> None:
    target = RedactionTarget(finding_id=uuid4(), location=_location(0, 1,
        artifact_id=uuid4(), kind="face", bbox=(0.1, 0.2, 0.3, 0.4)))
    with pytest.raises(ProcessingError):
        redact_text("SECRET", [target])


@pytest.mark.integration
@pytest.mark.presidio
def test_real_presidio_policy_text_redaction() -> None:
    import spacy
    from pii_guard import TextDetector
    if not spacy.util.is_package("en_core_web_sm"):
        pytest.skip("Install en_core_web_sm for real Presidio integration")
    text = "Employee EMP-38104. Email alex@example.com.\nKeep punctuation!"
    policy = PolicySnapshot(organization_id=uuid4(), version=1,
        builtins={rule: True for rule in BUILTIN_RULE_IDS}, custom_rules=[])
    detected = TextDetector().detect(text=text)
    evaluation = PolicyEvaluator(policy).evaluate(detected.accepted)
    before = evaluation.model_dump_json()
    assert redact_text(text, evaluation.redaction_targets) == "Employee [REDACTED]. Email [REDACTED].\nKeep punctuation!"
    assert evaluation.model_dump_json() == before


# Actual image/PDF rewrites use synthetic rasters and reportlab documents.
def _raster(fmt="PNG", orientation=None) -> bytes:
    with Image.new("RGB", (200, 100), "white") as image:
        image.paste((230, 10, 170), (20, 20, 80, 50))
        ImageDraw.Draw(image).text((23, 35), "SECRET", fill=(0, 0, 0))
        image.paste((20, 150, 230), (140, 60, 180, 90))
        exif = Image.Exif()
        exif[270] = "PRIVATE image metadata"
        if orientation:
            exif[274] = orientation
        with BytesIO() as stream:
            image.save(stream, format=fmt, exif=exif, **({"quality": 100, "subsampling": 0} if fmt == "JPEG" else {}))
            return stream.getvalue()


def _region_source(identity, bbox=(0.1, 0.2, 0.3, 0.3), kind="ocr", ref=None):
    return _location(0 if kind == "ocr" else None, 6 if kind == "ocr" else None,
        artifact_id=identity, kind=kind, source_ref=ref or ("image:ocr:0" if kind == "ocr" else "image"), bbox=bbox)


@pytest.mark.parametrize("fmt", ["PNG", "JPEG"])
def test_permanent_standalone_image_redaction(fmt: str) -> None:
    data = _raster(fmt)
    identity = uuid4()
    source = _region_source(identity)
    target = RedactionTarget(finding_id=uuid4(), location=source)
    result = redact_image(data, [target], artifact_id=identity, sources=[source])
    assert result.mime_type == ("image/png" if fmt == "PNG" else "image/jpeg")
    assert not result.format_changed and result.warnings == ('Image metadata was removed.',)
    assert (result.width, result.height) == (200, 100)
    assert "data" not in result.model_dump()
    assert b"PRIVATE" not in result.data
    with Image.open(BytesIO(result.data)) as image, Image.open(BytesIO(data)) as original:
        assert not image.getexif() and not image.info.get("exif")
        assert max(image.getpixel((40, 30))[:3]) <= (0 if fmt == "PNG" else 5)
        if fmt == "PNG":
            assert image.crop((20, 20, 80, 50)).tobytes() == bytes(60 * 30 * 3)
            assert image.getpixel((160, 70)) == original.getpixel((160, 70))
            assert image.getpixel((19, 20)) == original.getpixel((19, 20))
            assert image.getpixel((80, 49)) == original.getpixel((80, 49))


def test_conservative_pixel_rounding_edges_overlap_and_target_order() -> None:
    identity = uuid4()
    sources = [_region_source(identity, bbox=box, ref=f"image:ocr:{i}") for i, box in enumerate(
        [(0.101, 0.201, 0.299, 0.299), (0.4, 0.2, 0.1, 0.3), (0.99, 0.99, 0.01, 0.01)])]
    targets = [RedactionTarget(finding_id=uuid4(), location=source) for source in sources]
    first = redact_image(_raster(), [*targets, targets[0]], artifact_id=identity, sources=sources)
    second = redact_image(_raster(), reversed(targets), artifact_id=identity, sources=sources)
    assert first.data == second.data
    with Image.open(BytesIO(first.data)) as image:
        assert image.getpixel((20, 20)) == (0, 0, 0)
        assert image.getpixel((99, 49)) == (0, 0, 0)
        assert image.getpixel((199, 99)) == (0, 0, 0)
        assert image.getpixel((100, 49)) == (255, 255, 255)


@pytest.mark.parametrize("coordinate_frame", ["raw", "display"])
def test_exif_orientation_is_baked_in_and_metadata_removed(coordinate_frame) -> None:
    identity = uuid4()
    box = (0.1, 0.2, 0.3, 0.3) if coordinate_frame == "raw" else (0.5, 0.1, 0.3, 0.3)
    source = _region_source(identity, bbox=box)
    result = redact_image(_raster(orientation=6), [RedactionTarget(finding_id=uuid4(), location=source)],
        artifact_id=identity, sources=[source], coordinate_frame=coordinate_frame)
    assert (result.width, result.height) == (100, 200)
    with Image.open(BytesIO(result.data)) as image:
        assert image.getpixel((69, 40)) == (0, 0, 0)
        assert not image.getexif()


def test_image_format_conversion_is_explicit() -> None:
    result = redact_image(_raster("JPEG"), artifact_id=uuid4(), output_format="PNG")
    assert result.format_changed and result.mime_type == "image/png" and result.warnings


def test_transparent_sensitive_pixels_are_overwritten_opaquely() -> None:
    identity = uuid4()
    source = _region_source(identity, bbox=(0, 0, 1, 1))
    with Image.new("RGBA", (2, 2), (230, 10, 170, 0)) as image, BytesIO() as stream:
        image.save(stream, format="PNG")
        result = redact_image(stream.getvalue(), [RedactionTarget(finding_id=uuid4(), location=source)],
            artifact_id=identity, sources=[source])
    with Image.open(BytesIO(result.data)) as image:
        assert {image.getpixel((x, y)) for x in range(2) for y in range(2)} == {(0, 0, 0, 255)}


@pytest.mark.parametrize("enabled", [True, False])
def test_future_face_policy_targets_redact_pixels_only_when_enabled(enabled) -> None:
    identity = uuid4()
    source = _region_source(identity, kind="face")
    candidate = DetectionCandidate(rule_id="FACE", category="biometric", confidence=0.99,
        recognizer_confidence=0.99, recognizer="Synthetic", location=source)
    policy = PolicySnapshot(organization_id=uuid4(), version=1,
        builtins={rule: enabled for rule in BUILTIN_RULE_IDS}, custom_rules=[])
    evaluation = PolicyEvaluator(policy).evaluate([candidate])
    before = evaluation.model_dump_json()
    result = redact_image(_raster(), evaluation.redaction_targets, artifact_id=identity, sources=[source])
    with Image.open(BytesIO(result.data)) as image:
        assert image.getpixel((40, 30)) == ((0, 0, 0) if enabled else (230, 10, 170))
    assert evaluation.model_dump_json() == before
    with pytest.raises(ProcessingError):
        redact_image(_raster(), evaluation.findings, artifact_id=identity, sources=[source])


@pytest.mark.parametrize("field,value", [("artifact_id", uuid4()), ("source_ref", "wrong:ocr:0"),
    ("page_number", 1), ("bbox", (-0.1, 0.2, 0.3, 0.3)), ("bbox", (0.9, 0.2, 0.3, 0.3)), ("end", 7)])
def test_image_source_mismatches_and_invalid_regions_fail_closed(field, value) -> None:
    identity = uuid4()
    source = _region_source(identity)
    target = RedactionTarget(finding_id=uuid4(), location=source).model_copy(
        update={"location": source.model_copy(update={field: value})})
    with pytest.raises(ProcessingError) as error:
        redact_image(_raster(), [target], artifact_id=identity, sources=[source])
    assert error.value.code == ProcessingErrorCode.REDACTION_FAILED


def _pdf_case(tmp_path, *, native=True, image=True, rotation=0, crop=False):
    path = tmp_path / "source.pdf"
    with BytesIO() as stream:
        drawing = canvas.Canvas(stream, pagesize=(300, 200))
        if native:
            drawing.drawString(30, 175, "SECRET keep")
            drawing.drawString(30, 155, "SECOND safe")
        if image:
            drawing.drawImage(ImageReader(BytesIO(_raster())), 50, 50, width=200, height=100)
        drawing.save()
        writer = PdfWriter(clone_from=PdfReader(BytesIO(stream.getvalue())))
        writer.pages[0].rotate(rotation)
        if crop:
            writer.pages[0].cropbox = RectangleObject((10, 10, 290, 190))
        with path.open("wb") as output:
            writer.write(output)
    identity = uuid4()
    box = None
    with pymupdf.open(path) as document:
        page = document[0]
        if image:
            placement = pymupdf.Rect(page.get_image_info(xrefs=True)[0]["bbox"])
            sensitive = pymupdf.Rect(placement.x0 + 20, placement.y0 + 20,
                                     placement.x0 + 80, placement.y0 + 50) * page.rotation_matrix
            box = (sensitive.x0 / page.rect.width, sensitive.y0 / page.rect.height,
                   sensitive.width / page.rect.width, sensitive.height / page.rect.height)

    class Provider:
        def recognize_image(self, *args, **kwargs):
            return [OCRBlock("SECRET", 0.99, box)] if image else []

    extracted = extract_pdf(path, artifact_id=identity, filename=path.name, ocr_provider=Provider())
    targets = []
    for block in extracted.blocks:
        loc = block.location
        if loc.kind == "native_text":
            for secret in ("SECRET", "SECOND"):
                start = block.text.index(secret)
                targets.append(RedactionTarget(finding_id=uuid4(),
                    location=loc.model_copy(update={"start": start, "end": start + len(secret)})))
        else:
            targets.append(RedactionTarget(finding_id=uuid4(), location=loc))
    return path, identity, extracted.blocks, targets


@pytest.mark.parametrize("native,image", [(True, False), (False, True), (True, True)])
@pytest.mark.integration
def test_actual_native_scanned_and_mixed_pdf_rewrites(tmp_path, native, image) -> None:
    path, identity, blocks, targets = _pdf_case(tmp_path, native=native, image=image)
    original = path.read_bytes()
    output = tmp_path / "safe.pdf"
    result = redact_pdf_file(path, output, [*reversed(targets), targets[0]], artifact_id=identity, blocks=blocks)
    assert not result.format_changed and result.size_bytes == output.stat().st_size
    assert path.read_bytes() == original
    reader = PdfReader(output)
    assert len(reader.pages) == 1
    text = reader.pages[0].extract_text()
    assert "SECRET" not in text and "SECOND" not in text
    if native:
        assert "keep" in text and "safe" in text
        assert all(b"SECRET" not in stream.get_object().get_data()
                   for stream in (reader.pages[0]["/Contents"] if isinstance(reader.pages[0]["/Contents"], list)
                                  else [reader.pages[0]["/Contents"]]))
    assert not reader.pages[0].get("/Annots")
    assert not list(tmp_path.glob(".pii-redaction-*"))
    with pymupdf.open(output) as document:
        # Inspect all remaining image objects, not only the visible page raster.
        image_xrefs = [xref for xref in range(1, document.xref_length()) if document.xref_is_image(xref)]
        assert len(image_xrefs) == int(image)
        for xref in image_xrefs:
            with Image.open(BytesIO(document.extract_image(xref)["image"])) as raster:
                assert raster.getpixel((40, 30))[:3] == (0, 0, 0)
                assert raster.convert("RGB").crop((20, 20, 80, 50)).tobytes() == bytes(60 * 30 * 3)
                assert raster.getpixel((160, 70))[:3] == (20, 150, 230)
        pixmap = document[0].get_pixmap()
        if image:
            assert pixmap.pixel(90, 80)[:3] == (0, 0, 0)


@pytest.mark.parametrize("rotation", [0, 90, 180, 270])
@pytest.mark.parametrize("crop", [False, True])
def test_pdf_display_boxes_rotation_and_crop_mapping(tmp_path, rotation, crop) -> None:
    path, identity, blocks, targets = _pdf_case(tmp_path, rotation=rotation, crop=crop)
    output = redact_pdf(path.read_bytes(), targets, artifact_id=identity, blocks=blocks)
    assert "SECRET" not in PdfReader(BytesIO(output)).pages[0].extract_text()
    with pymupdf.open(stream=output, filetype="pdf") as document:
        page = document[0]
        assert page.rotation == rotation
        ocr = next(block for block in blocks if block.extraction_method == "ocr")
        x, y, w, h = ocr.location.bbox
        rendered = page.get_pixmap()
        assert rendered.pixel(int((x + w / 2) * rendered.width), int((y + h / 2) * rendered.height))[:3] == (0, 0, 0)
        xref = page.get_images(full=True)[0][0]
        with Image.open(BytesIO(document.extract_image(xref)["image"])) as image:
            assert image.getpixel((40, 30))[:3] == (0, 0, 0)
            assert image.getpixel((160, 70))[:3] == (20, 150, 230)


@pytest.mark.parametrize("field,value", [("artifact_id", uuid4()), ("page_number", 2),
    ("source_ref", "pdf:page:9:native_text"), ("end", 999), ("bbox", (0, 0, 0.5, 0.5))])
def test_pdf_bad_target_source_does_not_publish_output(tmp_path, field, value) -> None:
    path, identity, blocks, targets = _pdf_case(tmp_path, image=False)
    targets[0] = targets[0].model_copy(update={"location": targets[0].location.model_copy(update={field: value})})
    output = tmp_path / "safe.pdf"
    with pytest.raises(ProcessingError) as error:
        redact_pdf_file(path, output, targets, artifact_id=identity, blocks=blocks)
    assert error.value.code == ProcessingErrorCode.REDACTION_FAILED
    assert not output.exists() and not list(tmp_path.glob(".pii-redaction-*"))


def test_pdf_unmappable_whitespace_only_span_and_stale_text_fail(tmp_path) -> None:
    path, identity, blocks, _ = _pdf_case(tmp_path, image=False)
    block = blocks[0]
    offset = block.text.index("\n")
    target = RedactionTarget(finding_id=uuid4(), location=block.location.model_copy(update={"start": offset, "end": offset + 1}))
    with pytest.raises(ProcessingError):
        redact_pdf(path.read_bytes(), [target], artifact_id=identity, blocks=blocks)
    changed = block.model_copy(update={"text": block.text.replace("SECRET", "OTHER!")})
    with pytest.raises(ProcessingError):
        redact_pdf(path.read_bytes(), artifact_id=identity, blocks=[changed])


def test_disabled_pdf_findings_do_not_remove_native_text_or_pixels(tmp_path) -> None:
    path, identity, blocks, _ = _pdf_case(tmp_path)
    rule = CustomRule(id=uuid4(), name="Secret", kind="exact", pattern="SECRET", case_sensitive=True, enabled=False)
    policy = PolicySnapshot(organization_id=uuid4(), version=1,
        builtins={rule: False for rule in BUILTIN_RULE_IDS}, custom_rules=[rule])
    evaluation = PolicyEvaluator(policy).evaluate(blocks=blocks)
    before = evaluation.model_dump_json()
    output = redact_pdf(path.read_bytes(), evaluation.redaction_targets, artifact_id=identity, blocks=blocks)
    assert "SECRET" in PdfReader(BytesIO(output)).pages[0].extract_text()
    with pymupdf.open(stream=output, filetype="pdf") as document:
        assert document[0].get_pixmap().pixel(90, 80)[:3] == (230, 10, 170)
    assert evaluation.model_dump_json() == before


def test_future_pdf_face_region_path_erases_source_pixels(tmp_path) -> None:
    path, identity, blocks, _ = _pdf_case(tmp_path, native=False)
    ocr = blocks[0].location
    face = ocr.model_copy(update={"kind": "face", "start": None, "end": None,
                                 "source_ref": "pdf:page:1:rendered"})
    output = redact_pdf(path.read_bytes(), [RedactionTarget(finding_id=uuid4(), location=face)],
        artifact_id=identity, blocks=blocks, region_sources=[face])
    with pymupdf.open(stream=output, filetype="pdf") as document:
        xref = document[0].get_images()[0][0]
        with Image.open(BytesIO(document.extract_image(xref)["image"])) as image:
            assert image.getpixel((40, 30))[:3] == (0, 0, 0)


@pytest.mark.parametrize("feature", ["attachment", "annotation", "optional_content", "form", "metadata"])
def test_unsupported_hidden_pdf_structures_fail_closed(tmp_path, feature) -> None:
    path, identity, blocks, targets = _pdf_case(tmp_path, image=False)
    writer = PdfWriter(clone_from=PdfReader(path))
    if feature == "attachment":
        writer.add_attachment("hidden.txt", b"SECRET")
    elif feature == "annotation":
        writer.pages[0][NameObject("/Annots")] = DictionaryObject()
    else:
        key = {"optional_content": "/OCProperties", "form": "/AcroForm", "metadata": "/Metadata"}[feature]
        writer._root_object[NameObject(key)] = DictionaryObject()
    with BytesIO() as stream:
        writer.write(stream)
        with pytest.raises(ProcessingError):
            redact_pdf(stream.getvalue(), targets, artifact_id=identity, blocks=blocks)


def test_pdf_shared_image_payload_is_rejected(tmp_path) -> None:
    with BytesIO() as stream:
        drawing = canvas.Canvas(stream, pagesize=(300, 200))
        for _ in range(2):
            drawing.drawImage(ImageReader(BytesIO(_raster())), 50, 50, width=200, height=100)
            drawing.showPage()
        drawing.save()
        with pytest.raises(ProcessingError):
            redact_pdf(stream.getvalue(), artifact_id=uuid4())


def test_publication_failure_cleans_temporary_output(tmp_path, monkeypatch) -> None:
    import pii_guard.redaction.pdf as module
    path, identity, blocks, targets = _pdf_case(tmp_path, image=False)
    original = path.read_bytes()
    output = tmp_path / "safe.pdf"
    def fail_publish(*args): raise OSError("synthetic publication failure")
    monkeypatch.setattr(module.os, "link", fail_publish)
    with pytest.raises(ProcessingError):
        redact_pdf_file(path, output, targets, artifact_id=identity, blocks=blocks)
    assert not output.exists() and not list(tmp_path.glob(".pii-redaction-*"))
    assert path.read_bytes() == original


def test_pdf_output_destination_never_overwrites_source_or_existing_file(tmp_path) -> None:
    path, identity, blocks, targets = _pdf_case(tmp_path, image=False)
    original = path.read_bytes()
    with pytest.raises(ProcessingError):
        redact_pdf_file(path, path, targets, artifact_id=identity, blocks=blocks)
    output = tmp_path / "existing.pdf"
    output.write_bytes(b"keep existing file")
    with pytest.raises(ProcessingError):
        redact_pdf_file(path, output, targets, artifact_id=identity, blocks=blocks)
    assert output.read_bytes() == b"keep existing file" and path.read_bytes() == original


@pytest.mark.parametrize("fmt", ["image", "pdf"])
@pytest.mark.integration
def test_real_ocr_region_erasure_and_rerecognition(tmp_path, fmt) -> None:
    import os
    from pathlib import Path
    from pii_guard.ocr import RapidOCREngine
    from pii_guard.extraction.image import extract_image
    if os.environ.get("PII_GUARD_REAL_OCR") != "1":
        pytest.skip("Set PII_GUARD_REAL_OCR=1 for real OCR redaction")
    fixture = Path(__file__).parent / "fixtures" / "scan.png"
    identity = uuid4()
    engine = RapidOCREngine()
    if fmt == "image":
        original = fixture
        extracted = extract_image(original, artifact_id=identity, filename="scan.png", ocr_provider=engine)
    else:
        original = tmp_path / "scan.pdf"
        drawing = canvas.Canvas(str(original), pagesize=(600, 200))
        drawing.drawImage(str(fixture), 0, 30, width=600, height=140)
        drawing.save()
        extracted = extract_pdf(original, artifact_id=identity, filename="scan.pdf", ocr_provider=engine)
    rule = CustomRule(id=uuid4(), name="Synthetic OCR value", kind="exact", pattern="HELLO",
                      case_sensitive=False, enabled=True)
    policy = PolicySnapshot(organization_id=uuid4(), version=1,
        builtins={rule: True for rule in BUILTIN_RULE_IDS}, custom_rules=[rule])
    evaluation = PolicyEvaluator(policy).evaluate(blocks=extracted.blocks)
    assert evaluation.redaction_targets
    raster_path = tmp_path / "safe.png"
    if fmt == "image":
        result = redact_image(original.read_bytes(), evaluation.redaction_targets, artifact_id=identity,
                              sources=[block.location for block in extracted.blocks])
        raster_path.write_bytes(result.data)
    else:
        result = redact_pdf(original.read_bytes(), evaluation.redaction_targets, artifact_id=identity,
                            blocks=extracted.blocks)
        with pymupdf.open(stream=result, filetype="pdf") as document:
            document[0].get_pixmap(matrix=pymupdf.Matrix(2, 2)).save(raster_path)
    recognized = engine.recognize_image(raster_path)
    assert all("HELLO" not in block.text.upper() for block in recognized)


def test_pdf_standard_and_custom_scalar_info_are_removed(tmp_path) -> None:
    path, identity, blocks, targets = _pdf_case(tmp_path, image=False)
    writer = PdfWriter(clone_from=PdfReader(path))
    writer.add_metadata({"/Author": "PRIVATE metadata"})
    with BytesIO() as stream:
        writer.write(stream)
        output = redact_pdf(stream.getvalue(), targets, artifact_id=identity, blocks=blocks)
        assert b"PRIVATE" not in output and not PdfReader(BytesIO(output)).metadata
    writer.add_metadata({"/PrivatePayload": "PRIVATE metadata"})
    with BytesIO() as stream:
        writer.write(stream)
        output = redact_pdf(stream.getvalue(), targets, artifact_id=identity, blocks=blocks)
        assert b'PRIVATE' not in output and not PdfReader(BytesIO(output)).metadata


def test_pdf_invisible_and_offpage_text_is_rejected() -> None:
    for invisible in (True, False):
        with BytesIO() as stream:
            drawing = canvas.Canvas(stream, pagesize=(300, 200))
            text = drawing.beginText(30 if invisible else 400, 150)
            if invisible:
                text.setTextRenderMode(3)
            text.textOut("SECRET")
            drawing.drawText(text)
            drawing.save()
            with pytest.raises(ProcessingError):
                redact_pdf(stream.getvalue(), artifact_id=uuid4())


def test_pdf_target_page_identity_preserves_other_pages(tmp_path) -> None:
    path = tmp_path / "two-pages.pdf"
    drawing = canvas.Canvas(str(path), pagesize=(300, 200))
    for _ in range(2):
        drawing.drawString(30, 150, "SECRET keep")
        drawing.showPage()
    drawing.save()
    identity = uuid4()
    class NoOCR:
        def recognize_image(self, *args, **kwargs): pytest.fail("Native pages must not need OCR")
    extracted = extract_pdf(path, artifact_id=identity, filename=path.name, ocr_provider=NoOCR())
    location = extracted.blocks[1].location.model_copy(update={"end": 6})
    output = redact_pdf(path.read_bytes(), [RedactionTarget(finding_id=uuid4(), location=location)],
                        artifact_id=identity, blocks=extracted.blocks)
    pages = PdfReader(BytesIO(output)).pages
    assert len(pages) == 2 and "SECRET" in pages[0].extract_text()
    assert "SECRET" not in pages[1].extract_text() and "keep" in pages[1].extract_text()


@pytest.mark.parametrize("rotation", [0, 90, 180, 270])
def test_cropped_pdf_changes_pixels_only_in_requested_region(tmp_path, rotation) -> None:
    path, identity, blocks, targets = _pdf_case(tmp_path, native=False, rotation=rotation, crop=True)
    output = redact_pdf(path.read_bytes(), targets, artifact_id=identity, blocks=blocks)
    with pymupdf.open(path) as original, pymupdf.open(stream=output, filetype="pdf") as sanitized:
        before = original[0].get_pixmap()
        after = sanitized[0].get_pixmap()
        with Image.frombytes("RGB", (before.width, before.height), before.samples) as original_image, \
             Image.frombytes("RGB", (after.width, after.height), after.samples) as sanitized_image:
            changed = ImageChops.difference(original_image, sanitized_image).getbbox()
        assert changed is not None
        x, y, w, h = blocks[0].location.bbox
        assert changed[0] >= x * after.width - 1 and changed[1] >= y * after.height - 1
        assert changed[2] <= (x + w) * after.width + 1 and changed[3] <= (y + h) * after.height + 1

# Office integration tests exercise real packages, XML runs and decoded media.
def _office_case(tmp_path, kind, *, pictures=True, shared=False):
    aid = uuid4()
    path = tmp_path / f'original.{kind}'
    if kind == 'docx':
        obj = Document()
        p = obj.add_paragraph()
        p.add_run('😀 e\u0301 before ').bold = True
        p.add_run('SEC').italic = True
        p.add_run('RET').underline = True
        p.add_run(' after!\nkeep').bold = True
        obj.add_table(rows=1, cols=2).cell(0, 0).text = 'SECRET table'
        obj.tables[0].cell(0, 1).text = 'unrelated cell'
        if pictures:
            obj.add_picture(BytesIO(_raster()), width=WordInches(2))
            if shared:
                obj.add_picture(BytesIO(_raster()), width=WordInches(1))
    else:
        obj = Presentation()
        slide = obj.slides.add_slide(obj.slide_layouts[6])
        frame = slide.shapes.add_textbox(Inches(.2), Inches(.2), Inches(4), Inches(1)).text_frame
        p = frame.paragraphs[0]
        for text, bold in [('😀 e\u0301 before ', True), ('SEC', False), ('RET', True), (' after!', False)]:
            r = p.add_run(); r.text = text; r.font.bold = bold
        frame.add_paragraph().text = 'SECOND safe'
        slide.notes_slide.notes_text_frame.text = 'SECRET notes'
        if pictures:
            slide.shapes.add_picture(BytesIO(_raster()), Inches(1), Inches(2), width=Inches(2), height=Inches(1))
            if shared:
                slide.shapes.add_picture(BytesIO(_raster()), Inches(4), Inches(3), width=Inches(3), height=Inches(1.5))
                other = obj.slides.add_slide(obj.slide_layouts[6])
                other.shapes.add_picture(BytesIO(_raster()), Inches(1), Inches(1), width=Inches(2), height=Inches(1))
    obj.save(path)
    class OCR:
        def recognize_image(self, path, **kwargs):
            return [OCRBlock(text='SECRET', bbox=(.1, .2, .3, .3), confidence=.99)]
    extract = extract_docx if kind == 'docx' else extract_pptx
    extracted = extract(path, artifact_id=aid, filename=path.name, ocr_provider=OCR())
    targets = []
    for block in extracted.blocks:
        if 'SECRET' in block.text:
            start = block.text.index('SECRET')
            targets.append(RedactionTarget(finding_id=uuid4(), location=block.location.model_copy(update={'start': start, 'end': start+6})))
    return path, aid, extracted.blocks, targets


def _office_renderer(kind, file=False):
    return ({'docx': redact_docx_file, 'pptx': redact_pptx_file} if file else {'docx': redact_docx, 'pptx': redact_pptx})[kind]


def _zip_change(data, change):
    with ZipFile(BytesIO(data)) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    change(parts)
    b = BytesIO()
    with ZipFile(b, 'w', ZIP_DEFLATED) as z:
        for name, content in parts.items():
            z.writestr(name, content)
    return b.getvalue()


@pytest.mark.parametrize('kind', ['docx', 'pptx'])
def test_office_actual_rewrite_reopen_runs_tables_notes_media_and_leaks(tmp_path, kind):
    path, aid, blocks, targets = _office_case(tmp_path, kind)
    original = path.read_bytes()
    before = [t.model_dump_json() for t in targets]
    output = tmp_path / f'sanitized.{kind}'
    result = _office_renderer(kind, True)(path, output, reversed(targets), artifact_id=aid, blocks=blocks)
    assert path.read_bytes() == original
    assert result.size_bytes == output.stat().st_size and not result.format_changed
    with ZipFile(output) as z:
        removed = {'docProps/core.xml', 'docProps/app.xml', 'docProps/thumbnail.jpeg'}
        if kind == 'pptx': removed.add('ppt/printerSettings/printerSettings1.bin')
        assert set(z.namelist()) == set(ZipFile(BytesIO(original)).namelist()) - removed
        for name in z.namelist():
            if name.endswith('.xml'):
                root = ET.fromstring(z.read(name))
                assert 'SECRET' not in ''.join(root.itertext())
            if '/media/' in name:
                with Image.open(BytesIO(z.read(name))) as image:
                    assert image.convert('RGB').getpixel((40, 30)) == (0, 0, 0)
                    assert image.convert('RGB').getpixel((150, 70)) == (20, 150, 230)
                    assert not image.getexif()
    if kind == 'docx':
        obj = Document(output)
        assert obj.paragraphs[0].text == '😀 e\u0301 before [REDACTED] after!\nkeep'
        assert obj.paragraphs[0].runs[0].bold is True
        assert obj.paragraphs[0].runs[1].italic is True
        assert obj.paragraphs[0].runs[2].underline is True
        assert obj.tables[0].cell(0, 0).text == '[REDACTED] table'
        assert obj.tables[0].cell(0, 1).text == 'unrelated cell'
        assert len(obj.inline_shapes) == 1
    else:
        obj = Presentation(output)
        frame = obj.slides[0].shapes[0].text_frame
        assert frame.text == '😀 e\u0301 before [REDACTED] after!\nSECOND safe'
        assert frame.paragraphs[0].runs[0].font.bold is True
        assert obj.slides[0].notes_slide.notes_text_frame.text == '[REDACTED] notes'
        assert len(obj.slides[0].shapes) == 2
    assert [t.model_dump_json() for t in targets] == before
    assert not list(tmp_path.glob('.pii-redaction-*'))


@pytest.mark.parametrize('kind', ['docx', 'pptx'])
def test_office_overlaps_adjacency_containment_order_and_newline_runs(tmp_path, kind):
    path, aid, blocks, _ = _office_case(tmp_path, kind, pictures=False)
    block = next(b for b in blocks if b.text.startswith('😀'))
    start = block.text.index('SECRET')
    targets = [RedactionTarget(finding_id=uuid4(), location=block.location.model_copy(update={'start': lo, 'end': hi}))
               for lo, hi in [(start, start+3), (start+2, start+6), (start+6, start+7), (start+1, start+2)]]
    outputs = [_office_renderer(kind)(path.read_bytes(), order, artifact_id=aid, blocks=blocks)
               for order in [targets, list(reversed(targets)), targets[2:]+targets[:2]]]
    assert outputs[0] == outputs[1] == outputs[2]
    obj = Document(BytesIO(outputs[0])) if kind == 'docx' else Presentation(BytesIO(outputs[0]))
    text = obj.paragraphs[0].text if kind == 'docx' else obj.slides[0].shapes[0].text
    assert text.startswith('😀 e\u0301 before [REDACTED]after!')


@pytest.mark.parametrize('kind', ['docx', 'pptx'])
def test_office_shared_image_union_replaces_all_uses_and_keeps_placement(tmp_path, kind):
    path, aid, blocks, targets = _office_case(tmp_path, kind, shared=True)
    # Only one occurrence is targeted. Every relationship to that same media blob
    # must see the redacted raster; no duplicate unredacted payload may survive.
    target = next(t for t in targets if t.location.kind == 'ocr')
    output = _office_renderer(kind)(path.read_bytes(), [target], artifact_id=aid, blocks=blocks)
    with ZipFile(BytesIO(output)) as z:
        media = [n for n in z.namelist() if '/media/' in n]
        assert len(media) == 1
        with Image.open(BytesIO(z.read(media[0]))) as image:
            assert image.getpixel((40, 30)) == (0, 0, 0)
    if kind == 'docx':
        assert len(Document(BytesIO(output)).inline_shapes) == 2
    else:
        obj = Presentation(BytesIO(output))
        assert len(obj.slides) == 2
        assert obj.slides[0].shapes[1].left == Inches(1)
        assert obj.slides[0].shapes[2].left == Inches(4)
        assert obj.slides[0].shapes[2].width == Inches(3)
        assert obj.slides[1].shapes[0].image.blob == obj.slides[0].shapes[1].image.blob


@pytest.mark.parametrize('kind', ['docx', 'pptx'])
def test_office_disabled_findings_unchanged_and_enabled_precedence(tmp_path, kind):
    path, aid, blocks, _ = _office_case(tmp_path, kind, pictures=False)
    policy = PolicySnapshot(organization_id=uuid4(), version=7,
        builtins={rule: False for rule in BUILTIN_RULE_IDS}, custom_rules=[
            CustomRule(id=uuid4(), name='enabled', kind='exact', pattern='SECRET', case_sensitive=True, enabled=True),
            CustomRule(id=uuid4(), name='disabled', kind='exact', pattern='SECRET table' if kind=='docx' else 'SECRET notes', case_sensitive=True, enabled=False)])
    evaluated = PolicyEvaluator(policy).evaluate(blocks=blocks)
    before = [f.model_dump_json() for f in evaluated.findings]
    output = _office_renderer(kind)(path.read_bytes(), evaluated.redaction_targets, artifact_id=aid, blocks=blocks)
    assert any(f.action == 'allowed_disabled' for f in evaluated.findings)
    assert [f.model_dump_json() for f in evaluated.findings] == before
    with ZipFile(BytesIO(output)) as z:
        assert all(b'SECRET' not in z.read(n) for n in z.namelist() if n.endswith('.xml'))
    for rule in policy.custom_rules:
        rule.enabled = False
    evaluated = PolicyEvaluator(policy).evaluate(blocks=blocks)
    output = _office_renderer(kind)(path.read_bytes(), evaluated.redaction_targets, artifact_id=aid, blocks=blocks)
    obj = Document(BytesIO(output)) if kind == 'docx' else Presentation(BytesIO(output))
    assert 'SECRET' in (obj.paragraphs[0].text if kind=='docx' else obj.slides[0].shapes[0].text)


@pytest.mark.parametrize('kind', ['docx', 'pptx'])
@pytest.mark.parametrize('error', ['artifact', 'ref', 'slide', 'span', 'bbox', 'text', 'missing_blocks', 'duplicate_blocks'])
def test_office_original_source_integrity_fails_closed(tmp_path, kind, error):
    path, aid, blocks, targets = _office_case(tmp_path, kind)
    t = targets[0]
    if error == 'artifact':
        t = t.model_copy(update={'location': t.location.model_copy(update={'artifact_id': uuid4()})})
    elif error == 'ref':
        t = t.model_copy(update={'location': t.location.model_copy(update={'source_ref': 'SECRET wrong'})})
    elif error == 'slide':
        t = t.model_copy(update={'location': t.location.model_copy(update={'slide_number': 99})})
    elif error == 'span':
        t = t.model_copy(update={'location': t.location.model_copy(update={'end': 9999})})
    elif error == 'bbox':
        t = t.model_copy(update={'location': t.location.model_copy(update={'bbox': (.1, .1, .2, .2)})})
    elif error == 'text':
        b = next(b for b in blocks if b.location.source_ref == t.location.source_ref)
        blocks = [x if x is not b else b.model_copy(update={'text': 'X'*len(b.text)}) for x in blocks]
    elif error == 'missing_blocks':
        blocks = []
    else:
        blocks = blocks + [blocks[0]]
    with pytest.raises(ProcessingError) as caught:
        _office_renderer(kind)(path.read_bytes(), [t], artifact_id=aid, blocks=blocks)
    assert caught.value.code == ProcessingErrorCode.REDACTION_FAILED
    assert 'SECRET' not in ''.join(traceback.format_exception(caught.value))
    assert caught.value.artifact_id == aid


@pytest.mark.parametrize('kind', ['docx', 'pptx'])
@pytest.mark.parametrize('bad', ['extra_part', 'metadata', 'thumbnail', 'external', 'hidden', 'ole', 'comment', 'dtd'])
def test_office_unsupported_packages_even_without_targets(tmp_path, kind, bad):
    path, aid, _, _ = _office_case(tmp_path, kind, pictures=False)
    main = 'word/document.xml' if kind == 'docx' else 'ppt/slides/slide1.xml'
    def mutate(parts):
        if bad == 'extra_part':
            parts['customXml/SECRET.xml'] = b'<secret>SECRET</secret>'
        elif bad == 'metadata':
            parts['docProps/core.xml'] = parts['docProps/core.xml'].replace(b'<dc:title/>', b'<dc:title>SECRET</dc:title>')
        elif bad == 'thumbnail':
            parts['docProps/thumbnail.jpeg'] = _raster('JPEG')
        elif bad == 'external':
            rel = 'word/_rels/document.xml.rels' if kind=='docx' else 'ppt/slides/_rels/slide1.xml.rels'
            parts[rel] = parts[rel].replace(b'<Relationship ', b'<Relationship TargetMode="External" ', 1)
        elif bad == 'hidden':
            parts[main] = parts[main].replace(b'<w:r>', b'<w:r><w:rPr><w:vanish/></w:rPr>', 1) if kind=='docx' else parts[main].replace(b'<p:sld ', b'<p:sld show="0" ', 1)
        elif bad == 'ole':
            parts[main] = parts[main].replace(b'</w:body>', b'<w:object/></w:body>') if kind=='docx' else parts[main].replace(b'</p:spTree>', b'<p:graphicFrame/></p:spTree>')
        elif bad == 'comment':
            parts[main] = parts[main].replace(b'</w:body>', b'<!-- SECRET --></w:body>') if kind=='docx' else parts[main].replace(b'</p:spTree>', b'<!-- SECRET --></p:spTree>')
        else:
            parts[main] = b'<!DOCTYPE x [<!ENTITY secret "SECRET">]>' + parts[main]
    data = _zip_change(path.read_bytes(), mutate)
    if bad in {'metadata', 'thumbnail'}:
        output = _office_renderer(kind)(data, artifact_id=aid)
        with ZipFile(BytesIO(output)) as z:
            assert not any(name.startswith('docProps/') for name in z.namelist())
        return
    with pytest.raises(ProcessingError):
        _office_renderer(kind)(data, artifact_id=aid)


@pytest.mark.parametrize('kind', ['docx', 'pptx'])
@pytest.mark.parametrize('data', [b'', b'PK corrupt SECRET', b'not an office document'])
def test_office_corrupt_safe_errors(kind, data):
    with pytest.raises(ProcessingError) as error:
        _office_renderer(kind)(data, artifact_id=uuid4())
    assert error.value.code == ProcessingErrorCode.REDACTION_FAILED
    assert 'SECRET' not in str(error.value)


@pytest.mark.parametrize('kind', ['docx', 'pptx'])
@pytest.mark.parametrize('failure', ['publish', 'existing', 'original', 'rewrite'])
def test_office_file_cleanup_and_no_overwrite(tmp_path, monkeypatch, kind, failure):
    path, aid, blocks, targets = _office_case(tmp_path, kind)
    original = path.read_bytes()
    output = path if failure=='original' else tmp_path / f'output.{kind}'
    if failure == 'publish':
        def reject(*args):
            raise OSError('SECRET private path')
        monkeypatch.setattr('pii_guard.redaction.office.os.link', reject)
    elif failure == 'existing':
        output.write_bytes(b'keep existing')
    elif failure == 'rewrite':
        blocks = []
    with pytest.raises(ProcessingError) as error:
        _office_renderer(kind, True)(path, output, targets, artifact_id=aid, blocks=blocks)
    assert 'SECRET' not in str(error.value)
    assert path.read_bytes() == original
    assert not list(tmp_path.glob('.pii-redaction-*'))
    if failure == 'existing':
        assert output.read_bytes() == b'keep existing'
    elif failure != 'original':
        assert not output.exists()


@pytest.mark.parametrize('transform', ['rotation', 'crop', 'flip', 'offslide', 'group'])
def test_pptx_unsupported_picture_transforms_fail(tmp_path, transform):
    path, aid, blocks, targets = _office_case(tmp_path, 'pptx')
    obj = Presentation(path)
    shape = obj.slides[0].shapes[1]
    if transform == 'rotation':
        shape.rotation = 90
    elif transform == 'crop':
        shape.crop_left = .1
    elif transform == 'flip':
        shape.element.xpath('.//a:xfrm')[0].set('flipH', '1')
    elif transform == 'offslide':
        shape.left = -100
    else:
        group = obj.slides[0].shapes.add_group_shape()
        group.shapes.add_picture(BytesIO(_raster()), 0, 0, 10000, 10000)
    b = BytesIO(); obj.save(b)
    with pytest.raises(ProcessingError):
        redact_pptx(b.getvalue(), targets, artifact_id=aid, blocks=blocks)

@pytest.mark.parametrize('kind', ['docx', 'pptx'])
def test_office_linebreak_start_and_multiple_paragraph_targets(tmp_path, kind):
    aid = uuid4(); path = tmp_path / f'lines.{kind}'
    if kind == 'docx':
        obj = Document(); obj.add_paragraph('keep\nSECRET\tend'); obj.add_paragraph('SECOND')
    else:
        obj = Presentation(); slide = obj.slides.add_slide(obj.slide_layouts[6])
        frame = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2)).text_frame
        frame.text = 'keep\vSECRET end\nSECOND'
    obj.save(path)
    class NoOCR:
        def recognize_image(self, *args, **kwargs): return []
    extract = extract_docx if kind=='docx' else extract_pptx
    blocks = extract(path, artifact_id=aid, filename=path.name, ocr_provider=NoOCR()).blocks
    selected = [b for b in blocks if 'SECRET' in b.text or b.text=='SECOND']
    targets = [RedactionTarget(finding_id=uuid4(), location=b.location.model_copy(update={
        'start': b.text.index('\n') if 'SECRET' in b.text else 0,
        'end': b.text.index('SECRET')+6 if 'SECRET' in b.text else len(b.text)})) for b in selected]
    data = _office_renderer(kind)(path.read_bytes(), targets, artifact_id=aid, blocks=blocks)
    if kind=='docx':
        rewritten = Document(BytesIO(data)); assert rewritten.paragraphs[0].text=='keep[REDACTED]\tend'
        assert rewritten.paragraphs[1].text=='[REDACTED]'
    else:
        rewritten = Presentation(BytesIO(data)); assert rewritten.slides[0].shapes[0].text=='keep[REDACTED] end\n[REDACTED]'


@pytest.mark.parametrize('kind', ['docx', 'pptx'])
def test_office_empty_content_and_no_matches(tmp_path, kind):
    obj = Document() if kind=='docx' else Presentation()
    if kind=='pptx': obj.slides.add_slide(obj.slide_layouts[6])
    b=BytesIO();obj.save(b)
    data=_office_renderer(kind)(b.getvalue(), artifact_id=uuid4())
    if kind=='docx': assert not Document(BytesIO(data)).paragraphs
    else: assert len(Presentation(BytesIO(data)).slides[0].shapes)==0


@pytest.mark.parametrize('kind', ['docx', 'pptx'])
@pytest.mark.parametrize('bad', ['duplicate_zip', 'trailing', 'unplaced_media', 'alttext', 'unmapped_text', 'zero_font', 'namespace', 'media_type'])
def test_office_ambiguous_or_hidden_payloads_rejected(tmp_path, kind, bad):
    path,aid,blocks,targets=_office_case(tmp_path,kind)
    data=path.read_bytes()
    main='word/document.xml' if kind=='docx' else 'ppt/slides/slide1.xml'
    if bad=='duplicate_zip':
        with pytest.warns(UserWarning):
            stream=BytesIO(data)
            with ZipFile(stream,'a') as z: z.writestr(main,z.read(main))
        data=stream.getvalue()
    elif bad=='trailing': data+=b'SECRET payload'
    else:
        def mutate(parts):
            if bad=='unplaced_media': parts[('word' if kind=='docx' else 'ppt')+'/media/image99.png']=_raster()
            elif bad=='alttext': parts[main]=parts[main].replace(b'name="Picture 1"',b'name="Picture 1" descr="SECRET"') if kind=='docx' else parts[main].replace(b'descr="image.png"',b'descr="SECRET"')
            elif bad=='unmapped_text': parts[main]=parts[main].replace(b'</w:body>',b'<w:t>SECRET</w:t></w:body>') if kind=='docx' else parts[main].replace(b'</p:spTree>',b'<a:t>SECRET</a:t></p:spTree>')
            elif bad=='namespace': parts[main]=parts[main].replace(b'<w:document ',b'<w:document xmlns:SECRET="urn:SECRET" ',1) if kind=='docx' else parts[main].replace(b'<p:sld ',b'<p:sld xmlns:SECRET="urn:SECRET" ',1)
            elif bad=='media_type': parts['[Content_Types].xml']=parts['[Content_Types].xml'].replace(b'Extension="png" ContentType="image/png"',b'Extension="png" ContentType="image/jpeg"')
            else: parts[main]=parts[main].replace(b'<w:r>',b'<w:r><w:rPr><w:sz w:val="0"/></w:rPr>',1) if kind=='docx' else parts[main].replace(b'<a:r>',b'<a:r><a:rPr sz="0"/>',1)
        data=_zip_change(data,mutate)
    if bad == 'alttext':
        output = _office_renderer(kind)(data,targets,artifact_id=aid,blocks=blocks)
        with ZipFile(BytesIO(output)) as z:
            assert b'descr="SECRET"' not in z.read(main)
        return
    with pytest.raises(ProcessingError): _office_renderer(kind)(data,targets,artifact_id=aid,blocks=blocks)


@pytest.mark.parametrize('kind', ['docx', 'pptx'])
def test_office_whole_sensitive_raster_region_removed(tmp_path, kind):
    path,aid,blocks,targets=_office_case(tmp_path,kind)
    data=_office_renderer(kind)(path.read_bytes(),targets,artifact_id=aid,blocks=blocks)
    with ZipFile(BytesIO(data)) as z:
        media=next(n for n in z.namelist() if '/media/' in n)
        with Image.open(BytesIO(z.read(media))) as image:
            assert image.crop((20,20,80,50)).convert('RGB').getextrema()==((0,0),(0,0),(0,0))


@pytest.mark.integration
@pytest.mark.parametrize('kind', ['docx', 'pptx'])
def test_real_ocr_office_rewrite_removes_recognized_pixels(tmp_path, kind):
    import os
    if os.environ.get('PII_GUARD_REAL_OCR')!='1': pytest.skip('Set PII_GUARD_REAL_OCR=1 for real OCR')
    from pii_guard.ocr.engine import RapidOCREngine
    aid=uuid4();path=tmp_path/f'real.{kind}'
    scan=Path(__file__).parent/'fixtures'/'scan.png'
    if kind=='docx':
        obj=Document();obj.add_picture(BytesIO(scan.read_bytes()),width=WordInches(3))
    else:
        obj=Presentation();slide=obj.slides.add_slide(obj.slide_layouts[6])
        slide.shapes.add_picture(BytesIO(scan.read_bytes()),Inches(1),Inches(1),width=Inches(3))
    obj.save(path);provider=RapidOCREngine()
    extract=extract_docx if kind=='docx' else extract_pptx
    blocks=extract(path,artifact_id=aid,filename=path.name,ocr_provider=provider).blocks
    rule=CustomRule(id=uuid4(),name='fixture',kind='exact',pattern='HELLO',case_sensitive=False,enabled=True)
    evaluation=PolicyEvaluator(PolicySnapshot(organization_id=uuid4(),version=1,
        builtins={r:False for r in BUILTIN_RULE_IDS},custom_rules=[rule])).evaluate(blocks=blocks)
    assert evaluation.redaction_targets
    data=_office_renderer(kind)(path.read_bytes(),evaluation.redaction_targets,artifact_id=aid,blocks=blocks)
    with ZipFile(BytesIO(data)) as z:
        media=next(n for n in z.namelist() if '/media/' in n)
        image=tmp_path/'rewritten.png';image.write_bytes(z.read(media))
    recognized=provider.recognize_image(image,artifact_id=aid,source_ref='image')
    assert 'HELLO' not in ' '.join(b.text for b in recognized).upper()

@pytest.mark.parametrize('kind', ['docx', 'pptx'])
def test_office_shared_media_unions_distinct_enabled_regions(tmp_path, kind):
    path,aid,blocks,targets=_office_case(tmp_path,kind,shared=True)
    originals=[b for b in blocks if b.location.kind=='ocr']
    first=originals[0]
    if kind=='docx':
        # Word extraction enumerates the media relationship once, with OCR blocks.
        second=first.model_copy(update={'location': first.location.model_copy(update={
            'source_ref': first.location.source_ref.rpartition(':ocr:')[0]+':ocr:1',
            'bbox': (.7,.6,.2,.3)})})
        blocks=blocks+[second]
    else:
        second=originals[1]
        obj=Presentation(path);shape=obj.slides[0].shapes[2]
        box=(shape.left/obj.slide_width+.7*shape.width/obj.slide_width,
             shape.top/obj.slide_height+.6*shape.height/obj.slide_height,
             .2*shape.width/obj.slide_width,.3*shape.height/obj.slide_height)
        replacement=second.model_copy(update={'location':second.location.model_copy(update={'bbox':box})})
        blocks=[replacement if b is second else b for b in blocks];second=replacement
    selected=[RedactionTarget(finding_id=uuid4(),location=b.location) for b in (first,second)]
    output=_office_renderer(kind)(path.read_bytes(),selected,artifact_id=aid,blocks=blocks)
    with ZipFile(BytesIO(output)) as z:
        media=[n for n in z.namelist() if '/media/' in n];assert len(media)==1
        with Image.open(BytesIO(z.read(media[0]))) as image:
            assert image.getpixel((40,30))==(0,0,0)
            assert image.getpixel((150,70))==(0,0,0)
            assert image.getpixel((100,80))==(255,255,255)


@pytest.mark.parametrize('kind', ['docx', 'pptx'])
def test_office_wrong_container_format_fails(kind):
    other=Presentation() if kind=='docx' else Document()
    if kind=='docx': other.slides.add_slide(other.slide_layouts[6])
    stream=BytesIO();other.save(stream)
    with pytest.raises(ProcessingError): _office_renderer(kind)(stream.getvalue(),artifact_id=uuid4())
