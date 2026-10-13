"""Create fictional local examples. Never copy private policy files into fixtures."""
from pathlib import Path
from shutil import copyfile

ROOT = Path(__file__).resolve().parents[1]
LINES = [
    'Fictional demonstration data — not a real employee',
    'Name: Dana Mercer',
    'Employee ID: EMP-02048',
    'Email: dana@example.com',
    'Phone: +1 (202) 555-0147',
    'Address: 42 Demo Lane, Boston, MA 02110',
    'Password: DemoOnly_47!',
]
SAMPLE_NAMES = (
    'demo-sensitive-text.pdf', 'demo-sensitive-image.png', 'demo-sensitive-image.jpg',
    'demo-scanned-paper.pdf', 'demo-mixed-content.pdf',
    'demo-text-document.docx', 'demo-image-document.docx',
    'demo-text-slides.pptx', 'demo-image-slides.pptx',
    'demo-face.png', 'demo-face.pdf',
)


def ensure_demo_samples(directory=None):
    directory = Path(directory or ROOT / 'demo' / 'samples')
    directory.mkdir(parents=True, exist_ok=True)
    if all((directory / name).is_file() for name in SAMPLE_NAMES):
        return directory
    from PIL import Image, ImageDraw, ImageFont
    import pymupdf as fitz
    from docx import Document
    from docx.shared import Inches as WordInches
    from pptx import Presentation
    from pptx.util import Inches

    font = ImageFont.load_default(size=24)
    scan = Image.new('RGB', (1100, 620), 'white')
    draw = ImageDraw.Draw(scan)
    for i, line in enumerate(LINES):
        draw.text((40, 40 + i * 70), line, font=font, fill='black')
    for extension in ('png', 'jpg'):
        path = directory / f'demo-sensitive-image.{extension}'
        if not path.exists():
            scan.save(path)
    scan.close()
    image = directory / 'demo-sensitive-image.png'
    face = directory / 'demo-face.png'
    if not face.exists():
        copyfile(ROOT / 'tests' / 'fixtures' / 'visual' / 'synthetic-face.png', face)

    for name in ('demo-sensitive-text.pdf', 'demo-scanned-paper.pdf', 'demo-mixed-content.pdf', 'demo-face.pdf'):
        path = directory / name
        if path.exists():
            continue
        with fitz.open() as pdf:
            page = pdf.new_page(width=612, height=792)
            if name in ('demo-sensitive-text.pdf', 'demo-mixed-content.pdf'):
                page.insert_text((40, 45), '\n'.join(LINES), fontsize=12)
            if name in ('demo-scanned-paper.pdf', 'demo-mixed-content.pdf'):
                top = 250 if name == 'demo-mixed-content.pdf' else 40
                page.insert_image(fitz.Rect(40, top, 572, top + 300), filename=str(image))
            if name in ('demo-face.pdf', 'demo-mixed-content.pdf'):
                top = 565 if name == 'demo-mixed-content.pdf' else 70
                page.insert_image(fitz.Rect(40, top, 220, top + 180), filename=str(face))
            pdf.save(path)

    for embedded in (False, True):
        path = directory / ('demo-image-document.docx' if embedded else 'demo-text-document.docx')
        if not path.exists():
            doc = Document()
            for line in LINES:
                doc.add_paragraph(line)
            if embedded:
                doc.add_picture(str(image), width=WordInches(6))
                doc.add_picture(str(face), width=WordInches(2))
            doc.save(path)
        path = directory / ('demo-image-slides.pptx' if embedded else 'demo-text-slides.pptx')
        if not path.exists():
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(Inches(.5), Inches(.5), Inches(9), Inches(5)).text = '\n'.join(LINES)
            if embedded:
                slide = deck.slides.add_slide(deck.slide_layouts[6])
                slide.shapes.add_picture(str(image), Inches(.4), Inches(.3), width=Inches(8))
                slide.shapes.add_picture(str(face), Inches(8.5), Inches(4.2), width=Inches(1.2))
            deck.save(path)
    return directory


if __name__ == '__main__':
    print(f'Local fictional samples: {ensure_demo_samples()}')
