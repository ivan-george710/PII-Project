"""Fictional fixture coverage, real file structures, and sample API downloads."""
from fastapi.testclient import TestClient
from scripts.create_demo_samples import ensure_demo_samples, SAMPLE_NAMES
from app.web.server import create_app


def test_samples_cover_native_scanned_mixed_and_embedded_content(tmp_path):
    import pymupdf as fitz
    from docx import Document
    from pptx import Presentation
    folder = ensure_demo_samples(tmp_path / 'samples')
    assert set(p.name for p in folder.iterdir()) == set(SAMPLE_NAMES)
    with fitz.open(folder / 'demo-sensitive-text.pdf') as pdf:
        assert 'dana@example.com' in pdf[0].get_text()
    with fitz.open(folder / 'demo-scanned-paper.pdf') as pdf:
        assert not pdf[0].get_text().strip()
        assert pdf[0].get_images()
    with fitz.open(folder / 'demo-mixed-content.pdf') as pdf:
        assert 'dana@example.com' in pdf[0].get_text()
        assert len(pdf[0].get_images()) == 2
    doc = Document(folder / 'demo-image-document.docx')
    assert len(doc.inline_shapes) == 2
    deck = Presentation(folder / 'demo-image-slides.pptx')
    assert len(deck.slides) == 2
    assert sum(shape.shape_type == 13 for shape in deck.slides[1].shapes) == 2
    # Reopening the menu must not overwrite a local user's sample.
    extra = folder / 'user-file.pdf'; extra.write_bytes(b'local file')
    ensure_demo_samples(folder)
    assert extra.read_bytes() == b'local file'


def test_sample_api_lists_and_downloads_all_formats_with_correct_mime(tmp_path, monkeypatch):
    folder = ensure_demo_samples(tmp_path / 'samples')
    monkeypatch.setenv('DEMO_SAMPLE_DIR', str(folder))
    client = TestClient(create_app(work_dir=tmp_path / 'work', pipeline=object(), allow_test_host=True))
    assert {item['filename'] for item in client.get('/api/samples').json()['items']} == set(SAMPLE_NAMES)
    for name in SAMPLE_NAMES:
        response = client.get('/api/samples/' + name)
        assert response.status_code == 200
        assert response.content == (folder / name).read_bytes()
        if name.endswith('.jpg'):
            assert response.headers['content-type'] == 'image/jpeg'
    assert client.get('/api/samples/not-a-sample.pdf').status_code == 404
