"""Real OCR/artifact regression tests, with fictional sensitive markers only."""
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile
import pytest
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from pptx import Presentation
from pptx.util import Inches
import pymupdf as fitz
from app.pipeline import DocumentPipeline
from app.engines.extraction.ocr import OCREngine
from app.engines.pii_detection import PIIDetectionEngine
from app.engines.extraction.models import ExtractionResult,ExtractedPage
from app.engines.redaction import RedactionEngine


def marker_image(path):
    image=Image.new('RGB',(1000,250),'white')
    draw=ImageDraw.Draw(image)
    font=ImageFont.load_default(size=38)
    draw.text((25,30),'Password: FictionalSecret2048!',font=font,fill='black')
    draw.text((25,110),'Email: dana@example.test',font=font,fill='black')
    image.save(path)


@pytest.mark.parametrize('kind',['pdf','docx','pptx','png','jpg'])
def test_real_image_text_removed_from_serialized_artifact(tmp_path,kind):
    image=tmp_path/'source.png';marker_image(image)
    path=tmp_path/('mixed.'+kind)
    if kind=='pdf':
        with fitz.open() as doc:
            page=doc.new_page(width=600,height=400)
            page.insert_text((30,30),'A native heading')
            page.insert_image(fitz.Rect(20,80,580,220),filename=str(image))
            doc.save(path)
    elif kind=='docx':
        doc=Document();doc.add_paragraph('A native heading');doc.add_picture(str(image));doc.save(path)
    elif kind=='pptx':
        deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6]);slide.shapes.add_picture(str(image), Inches(.5),Inches(1),width=Inches(8));deck.save(path)
    else:
        with Image.open(image) as source:source.save(path)
    pipeline=DocumentPipeline(artifact_output_dir=tmp_path/'private-artifacts')
    pipeline.set_custom_rules([{'name':'fixture','pattern':'FictionalSecret2048!','entity_type':'CREDENTIAL'}])
    result=pipeline.process(path)
    assert result.failure is None
    assert result.artifact_result.reopened and result.artifact_result.supported_content_verified
    assert result.artifact_result.mapping_failures==result.artifact_result.verification_failures==0
    assert 'FictionalSecret2048!' in result.extraction_result.full_text
    assert 'FictionalSecret2048!' not in result.redaction_result.full_text
    rendered=[]
    if kind=='pdf':
        with fitz.open(result.artifact_result.path) as doc:
            rendered=[doc[0].get_pixmap(dpi=200).tobytes('png')]
    elif kind in {'docx','pptx'}:
        with ZipFile(result.artifact_result.path) as archive:
            rendered=[archive.read(n) for n in archive.namelist() if '/media/' in n]
    else:rendered=[result.artifact_result.path.read_bytes()]
    ocr=OCREngine()
    for index,data in enumerate(rendered):
        p=tmp_path/f'output-{index}.png';p.write_bytes(data)
        text='\n'.join(b.text for b in ocr.extract_image(p))
        assert 'FictionalSecret2048' not in text
        assert 'dana@example.test' not in text


@pytest.mark.parametrize('text,value',[('hi name is ds','ds'),('hii im da','da'),
    ('My employee ID is EMP-2048','EMP-2048'),('Email dana.example@example.test','dana.example@example.test'),
    ('Password: FictionalSecret2048!','FictionalSecret2048!')])
def test_original_prompt_regressions(text,value):
    source=ExtractionResult(Path('message.docx'),'message.docx','docx',[ExtractedPage(1,text,'docx')])
    detected=PIIDetectionEngine().detect(source)
    output=RedactionEngine().redact(source,detected)
    assert value not in output.full_text
    assert output.redactions
    if text=='hii im da':assert output.full_text.startswith('hii im ')


@pytest.mark.parametrize('text,value',[
    ('Address: 482 Juniper Lane, Testford, CA 90210','482 Juniper Lane, Testford, CA 90210'),
    ('DOB: 1994-07-19','1994-07-19'),
    ('MAC: 02:00:5e:10:00:01','02:00:5e:10:00:01'),
    ('Bank account: 123456789012','123456789012'),
    ('Driver license: A1234567','A1234567'),
    ('Passport: 123456789','123456789'),
    ('IP: 192.0.2.42','192.0.2.42'),
])
def test_additional_requested_pii_categories(text,value):
    source=ExtractionResult(Path('message.docx'),'message.docx','docx',[ExtractedPage(1,text,'docx')])
    result=RedactionEngine().redact(source,PIIDetectionEngine().detect(source))
    assert value not in result.full_text
    assert result.redactions


def test_crossing_rules_mask_entire_union_and_keep_audit():
    from app.engines.audit import AuditEngine
    source=ExtractionResult(Path('message.docx'),'message.docx','docx',[ExtractedPage(1,'alpha bravo charlie','docx')])
    detection=PIIDetectionEngine(custom_rules=[{'pattern':'alpha bravo','entity_type':'CUSTOM_RULE'},
                                              {'pattern':'bravo charlie','entity_type':'CUSTOM_RULE'}]).detect(source)
    result=RedactionEngine().redact(source,detection)
    assert result.full_text=='[PII_SPAN_001]'
    audit=AuditEngine().create_audit(source,detection,result)
    assert len(audit.records)==detection.raw_candidate_count
    assert all(value not in result.full_text for value in ('alpha','bravo','charlie'))


def test_password_and_account_list_from_screenshot_do_not_need_custom_rules():
    text='Username\nPassword\nLogin\nAccepted usernames are:\nfictional_user\nfictional_admin\nPassword for all users:\nFictionalSecret2048!'
    source=ExtractionResult(Path('image.png'),'image.png','image',[ExtractedPage(1,text,'ocr')])
    result=RedactionEngine().redact(source,PIIDetectionEngine().detect(source))
    assert all(value not in result.full_text for value in ('fictional_user','fictional_admin','FictionalSecret2048!'))
    assert 'Login' in result.full_text


def test_saved_pdf_leak_is_repaired_with_original_crop_evidence(tmp_path,monkeypatch):
    from app.engines.extraction.office import source_hash
    image=tmp_path/'source.png';marker_image(image)
    path=tmp_path/'scan.pdf'
    with fitz.open() as doc:
        page=doc.new_page(width=600,height=300)
        page.insert_image(fitz.Rect(20,40,580,180),filename=str(image));doc.save(path)
    pipeline=DocumentPipeline(artifact_output_dir=tmp_path/'private-artifacts')
    result=pipeline.process(path)
    assert result.failure is None
    original_hash=source_hash(path)
    # Simulate a saved output retaining the original raster after generation.
    result.artifact_result.path.write_bytes(path.read_bytes())
    original_ocr=pipeline.extraction.ocr_engine.extract_image
    def crop_unreadable(path):
        with Image.open(path) as image:
            if image.height<500:return []
        return original_ocr(path)
    monkeypatch.setattr(pipeline.extraction.ocr_engine,'extract_image',crop_unreadable)
    pipeline._repair_pdf_residuals(result)
    assert result.failure is None
    assert any(b.mask_only for page in result.extraction_result.pages for b in page.blocks)
    assert any(d.recognizer=='ResidualPrivacyRecognizer' for d in result.detection_result.detections)
    assert source_hash(path)==original_hash
    assert result.artifact_result.supported_content_verified
    with fitz.open(result.artifact_result.path) as returned:
        image=tmp_path/'returned.png';returned[0].get_pixmap(dpi=200).save(image)
    text='\n'.join(b.text for b in OCREngine().extract_image(image))
    assert 'FictionalSecret2048' not in text and 'dana@example.test' not in text


def test_residual_verifier_failure_is_a_processing_failure(tmp_path,monkeypatch):
    image=tmp_path/'source.png';marker_image(image)
    path=tmp_path/'scan.pdf'
    with fitz.open() as doc:
        page=doc.new_page(width=600,height=300)
        page.insert_image(fitz.Rect(20,40,580,180),filename=str(image));doc.save(path)
    def fail(*_):raise ValueError('Unavailable verifier')
    monkeypatch.setattr('app.engines.extraction.residual.residual_blocks',fail)
    result=DocumentPipeline(artifact_output_dir=tmp_path/'private-artifacts').process(path)
    assert result.failure.stage=='residual_verification'
    assert result.processing_status.value=='PROCESSING_FAILED'
