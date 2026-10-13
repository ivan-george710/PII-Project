"""Synthetic Office extraction, original-structure editing and independent verification."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
from types import SimpleNamespace
from unittest.mock import Mock
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from pptx import Presentation
from pptx.util import Inches
from presidio_analyzer import RecognizerResult

from app.engines.extraction.docx import extract_docx
from app.engines.extraction.pptx import extract_pptx
from app.engines.extraction.engine import ExtractionEngine
from app.engines.extraction.ocr import OCREngine
from app.engines.pii_detection import PIIDetectionEngine
from app.engines.redaction import RedactionEngine
from app.engines.redaction.artifacts import OfficeArtifactEngine, ArtifactValidationError
from app.engines.audit import AuditEngine, ReviewerFeedback
from app.pipeline import DocumentPipeline, ProcessingStatus


def docx_fixture(path):
    doc = Document()
    p = doc.add_paragraph("Ordinary business policy. Owner: ")
    p.add_run("John ").bold = True
    p.add_run("Smith").italic = True
    p = doc.add_paragraph("Employee: ")
    p.add_run("EMP-").bold = True
    p.add_run("00001")
    doc.add_paragraph("Email: audit@example.com")
    doc.add_paragraph("Missing: TokenPrivate")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "EMP-00002"
    table.cell(0, 1).text = "Non-sensitive table text"
    table.cell(0, 1).add_table(rows=1, cols=1).cell(0, 0).text = "EMP-00003"
    doc.sections[0].header.paragraphs[0].text = "Header EMP-00004"
    doc.sections[0].footer.paragraphs[0].text = "Footer EMP-00005"
    doc.save(path)


def pptx_fixture(path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Ordinary business policy"
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    p = box.text_frame.paragraphs[0]
    p.add_run().text = "Owner: "
    p.add_run().text = "John "
    p.runs[-1].font.bold = True
    p.add_run().text = "Smith"
    p.runs[-1].font.italic = True
    slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1)).text = "EMP-00001 audit@example.com"
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(6), Inches(1)).table
    table.cell(0, 0).text = "EMP-00002"
    table.cell(0, 1).text = "Non-sensitive table text"
    group = slide.shapes.add_group_shape()
    subgroup = group.shapes.add_group_shape()
    subgroup.shapes.add_textbox(Inches(1), Inches(4), Inches(3), Inches(1)).text = "EMP-00003"
    slide.notes_slide.notes_text_frame.text = "Speaker notes EMP-00004"
    second = presentation.slides.add_slide(presentation.slide_layouts[6])
    second.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = "EMP-00005 Missing: TokenPrivate"
    presentation.save(path)


class SyntheticAnalyzer:
    def analyze(self, text, language):
        results = []
        for entity, pattern in [("EMPLOYEE_ID", r"EMP-\d{5}"), ("PERSON", "John Smith"),
                                ("EMAIL_ADDRESS", r"audit@example\.com")]:
            for match in re.finditer(pattern, text):
                results.append(RecognizerResult(entity, match.start(), match.end(), .9,
                    recognition_metadata={"recognizer_name": "Synthetic"}))
        return results


@pytest.fixture(params=["docx", "pptx"])
def office_case(tmp_path, request):
    kind = request.param
    path = tmp_path / ("private-source." + kind)
    (docx_fixture if kind == "docx" else pptx_fixture)(path)
    extract = extract_docx if kind == "docx" else extract_pptx
    pipeline = DocumentPipeline(extraction=SimpleNamespace(extract=extract),
        detection=PIIDetectionEngine(analyzer=SyntheticAnalyzer()), audit=AuditEngine(tmp_path / "private-audit"),
        artifact_output_dir=tmp_path / "private-artifacts")
    return kind, path, extract, pipeline


def test_exact_mappings_all_locations(office_case):
    kind, path, extract, _ = office_case
    result = extract(path)
    assert result.structure.source_sha256 == hashlib.sha256(path.read_bytes()).hexdigest()
    assert all(f"EMP-{number:05d}" in result.full_text for number in range(1, 6))
    assert "John Smith" in result.full_text and "audit@example.com" in result.full_text
    assert "Non-sensitive table text" in result.full_text
    with ZipFile(path) as archive:
        from app.engines.extraction.office import parse_xml, NS, node_text
        roots = {s.part: parse_xml(archive.read(s.part)) for s in result.structure.segments}
        for segment in result.structure.segments:
            text = "".join(node_text(roots[segment.part].xpath(n.path, namespaces=NS)[0]) for n in segment.nodes)
            page = next(p for p in result.pages if p.page_number == segment.page_number)
            assert page.text[segment.start:segment.end] == text
    if kind == "docx":
        assert result.structure.location_kind == "docx_logical_paragraph"
        assert {s.location["kind"] for s in result.structure.segments} == {"body", "header", "footer"}
        assert any(s.location["table_depth"] == 2 for s in result.structure.segments)
    else:
        assert len(result.pages) == 2
        assert result.structure.location_kind == "pptx_slide"
        assert any(len(s.location["group_ids"]) == 2 for s in result.structure.segments)
        assert any(s.location["kind"] == "notes" for s in result.structure.segments)
        assert any("column" in s.location for s in result.structure.segments)


def test_pipeline_redaction_reopening_and_original_preserved(office_case):
    kind, path, extract, pipeline = office_case
    before = path.read_bytes()
    result = pipeline.process(path)
    assert result.failure is None
    artifact = result.artifact_result
    assert artifact.redaction_completed and artifact.verification_completed and artifact.reopened
    assert artifact.supported_content_verified and artifact.mapping_failures == artifact.verification_failures == 0
    assert artifact.applied_redactions == 7
    assert artifact.residual_confirmed_values == artifact.residual_sensitive_patterns == 0
    assert artifact.path != path and artifact.path.name.startswith("review_")
    text = extract(artifact.path).full_text
    assert all(f"EMP-{number:05d}" not in text for number in range(1, 6))
    assert "John Smith" not in text and "audit@example.com" not in text
    assert "Non-sensitive table text" in text
    assert path.read_bytes() == before
    assert "document_properties" in artifact.unsupported_structures
    assert not artifact.artifact_verified and not result.safe_for_llm
    assert result.processing_status == ProcessingStatus.PROCESSED_REVIEW_REQUIRED
    assert {"artifact_redaction", "artifact_verification"} <= set(result.stage_timings)
    summary = json.dumps(result.sanitized_summary())
    for value in (str(path), "John Smith", "EMP-00001", "audit@example.com", str(artifact.path)):
        assert value not in summary
    if kind == "docx":
        original, output = Document(path), Document(artifact.path)
        assert output.paragraphs[0].runs[1].bold == original.paragraphs[0].runs[1].bold
        assert output.paragraphs[0].runs[2].italic == original.paragraphs[0].runs[2].italic
        assert len(output.tables) == len(original.tables)
    else:
        original, output = Presentation(path), Presentation(artifact.path)
        assert len(output.slides) == len(original.slides)
        assert len(output.slides[0].shapes) == len(original.slides[0].shapes)
        assert output.slides[0].shapes[1].text_frame.paragraphs[0].runs[1].font.bold


def test_feedback_regenerates_from_original_without_upstream_reruns(office_case):
    _, path, extract, pipeline = office_case
    original = pipeline.process(path)
    snapshot = deepcopy(original)
    original_bytes = path.read_bytes()
    pipeline.file_intelligence.analyze = Mock(side_effect=AssertionError("No classification rerun"))
    pipeline.extraction.extract = Mock(side_effect=AssertionError("No source extraction rerun"))
    pipeline.detection.detect = Mock(side_effect=AssertionError("No Presidio rerun"))
    person = next(r for r in original.audit_result.records if r.original_detection.entity_type == "PERSON")
    updated = pipeline.apply_feedback(original, ReviewerFeedback("FALSE_POSITIVE", "synthetic-reviewer", person.detection_id))
    assert updated.failure is None
    assert "John Smith" in extract(updated.artifact_result.path).full_text
    assert updated.artifact_result.path != original.artifact_result.path
    assert "John Smith" not in extract(original.artifact_result.path).full_text
    employee = next(r for r in updated.audit_result.records if r.original_detection.entity_type == "EMPLOYEE_ID")
    updated = pipeline.apply_feedback(updated, ReviewerFeedback("CONFIRM_PII", "synthetic-reviewer", employee.detection_id))
    page = next(p for p in original.extraction_result.pages if "TokenPrivate" in p.text)
    start = page.text.index("TokenPrivate")
    updated = pipeline.apply_feedback(updated, ReviewerFeedback("MISSED_PII", "synthetic-reviewer", page_number=page.page_number,
        original_start=start, original_end=start + len("TokenPrivate"), entity_type="PERSON"))
    assert updated.failure is None and updated.artifact_result.supported_content_verified
    assert "TokenPrivate" not in extract(updated.artifact_result.path).full_text
    assert updated.audit_result.revision == 3
    assert original == snapshot and path.read_bytes() == original_bytes
    pipeline.file_intelligence.analyze.assert_not_called()
    pipeline.extraction.extract.assert_not_called()
    pipeline.detection.detect.assert_not_called()
    assert updated.extraction_result == original.extraction_result
    assert updated.detection_result == original.detection_result
    assert not updated.safe_for_llm


def test_mapping_snapshot_persistence_and_reload(office_case):
    _, path, _, pipeline = office_case
    original = pipeline.process(path)
    saved = pipeline.save_audit(original)
    assert saved.failure is None
    pipeline.extraction.extract = Mock(side_effect=AssertionError("No extraction rerun"))
    pipeline.detection.detect = Mock(side_effect=AssertionError("No detection rerun"))
    loaded = pipeline.load_audit(original.document_id)
    assert loaded.failure is None and loaded.audit_result == original.audit_result
    assert loaded.extraction_result.structure == original.extraction_result.structure
    assert loaded.artifact_result.reopened and loaded.artifact_result.supported_content_verified


def test_changed_source_fails_closed(office_case):
    kind, path, _, pipeline = office_case
    result = pipeline.process(path)
    path.write_bytes(path.read_bytes() + b"changed-source")
    target = result.audit_result.records[0].detection_id
    updated = pipeline.apply_feedback(result, ReviewerFeedback("CONFIRM_PII", "reviewer", target))
    assert updated.failure.stage == "artifact_redaction"
    assert updated.artifact_result is None
    assert not updated.safe_for_llm


def test_output_tamper_verification_fails(office_case):
    kind, path, _, pipeline = office_case
    result = pipeline.process(path)
    output = result.artifact_result.path
    if kind == "docx":
        doc = Document(output)
        doc.paragraphs[0].add_run(" John Smith")
        doc.save(output)
    else:
        deck = Presentation(output)
        deck.slides[0].shapes[1].text_frame.paragraphs[0].add_run().text = " John Smith"
        deck.save(output)
    verified = pipeline.artifacts.verify(result.extraction_result, result.redaction_result, result.artifact_result)
    assert verified.verification_completed and verified.verification_failures > 0
    assert not verified.supported_content_verified and not verified.artifact_verified
    assert verified.residual_confirmed_values > 0


def test_docx_text_box_is_reported_not_claimed_extracted(tmp_path):
    path = tmp_path / "textbox.docx"
    doc = Document()
    paragraph = doc.add_paragraph("Supported EMP-00001")
    box = OxmlElement("w:txbxContent")
    p, r, t = OxmlElement("w:p"), OxmlElement("w:r"), OxmlElement("w:t")
    t.text = "Hidden EMP-00002"
    r.append(t); p.append(r); box.append(p); paragraph._p.append(box)
    doc.save(path)
    result = extract_docx(path)
    assert "EMP-00001" in result.full_text and "EMP-00002" not in result.full_text
    assert result.structure.limitations["text_boxes"] == 1
    assert result.structure.limitations["unmapped_text_nodes"] == 1


def test_pptx_image_and_chart_are_unverified(tmp_path):
    from PIL import Image
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    image = tmp_path / "fixture.png"
    Image.new("RGB", (20, 20), "white").save(image)
    path = tmp_path / "image.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(str(image), Inches(1), Inches(1))
    data = CategoryChartData()
    data.categories = ["Private chart name"]
    data.add_series("Private series", [1])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(3), Inches(2), data)
    deck.save(path)
    result = extract_pptx(path)
    assert result.structure.limitations["media_without_ocr"] == 1
    assert result.structure.limitations["charts_or_diagrams"] >= 1
    assert result.structure.limitations["embedded_objects"] >= 1
    assert "Private chart name" not in result.full_text


def test_cross_structure_detection_blocks_artifact_mapping(tmp_path):
    from app.engines.pii_detection.models import DetectedPII, PIIDetectionResult
    path = tmp_path / "crossing.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for i, value in enumerate(["John", "Smith"]):
        slide.shapes.add_textbox(Inches(1), Inches(i + 1), Inches(3), Inches(1)).text = value
    deck.save(path)
    extraction = extract_pptx(path)
    text = extraction.pages[0].text
    d = DetectedPII("PERSON", text, 0, len(text), .9, "pattern", "Synthetic", 1,
                    validation_decision="CONFIRM", final_decision="CONFIRM")
    detection = PIIDetectionResult(path, path.name, [d], [], [], 1, [1])
    redaction = RedactionEngine().redact(extraction, detection)
    engine = OfficeArtifactEngine(tmp_path / "private-artifacts")
    artifact = engine.redact(extraction, detection, redaction)
    assert artifact.mapping_failures == 1 and artifact.path.exists()
    assert artifact.applied_redactions == 0
    assert len(artifact.mapping_issues) == 1
    artifact = engine.verify(extraction, redaction, artifact)
    assert artifact.reopened and artifact.verification_failures > 0 and not artifact.artifact_verified


@pytest.mark.parametrize("kind", ["docx", "pptx"])
def test_real_presidio_office_pipeline(tmp_path, kind):
    path = tmp_path / ("real-detector." + kind)
    (docx_fixture if kind == "docx" else pptx_fixture)(path)
    pipeline = DocumentPipeline(artifact_output_dir=tmp_path / "private-artifacts")
    result = pipeline.process(path)
    assert result.failure is None
    types = {d.entity_type for d in result.detection_result.detections}
    assert {"EMPLOYEE_ID", "PERSON", "EMAIL_ADDRESS"} <= types
    assert result.detection_result.raw_candidate_count == result.detection_result.disposition_count
    assert result.artifact_result.reopened and result.artifact_result.supported_content_verified
    assert not result.safe_for_llm and not result.coverage_verified


@pytest.mark.parametrize("kind", ["docx", "pptx"])
def test_extraction_engine_routes_office_with_media_ocr(tmp_path, monkeypatch, kind):
    path = tmp_path / ("routed." + kind)
    (docx_fixture if kind == "docx" else pptx_fixture)(path)
    monkeypatch.setattr(OCREngine, "__init__", lambda self: None)
    monkeypatch.setattr(OCREngine, "extract_pdf", Mock(side_effect=AssertionError("No Office OCR")))
    calls = []
    monkeypatch.setattr(OCREngine, "extract_image", lambda self,p: calls.append(p) or [])
    result = ExtractionEngine().extract(path)
    assert calls and any(p.extraction_method == "media_ocr" for p in result.pages)
    assert result.file_type == kind and result.structure is not None


@pytest.mark.parametrize("kind", ["docx", "pptx"])
def test_office_classification_uses_package_contents_not_extension(tmp_path, monkeypatch, kind):
    from app.engines.file_intelligence import FileIntelligenceEngine
    from app.engines.file_intelligence import detector
    path = tmp_path / ("fixture." + kind)
    (docx_fixture if kind == "docx" else pptx_fixture)(path)
    renamed = tmp_path / "misleading.pdf"
    path.rename(renamed)
    monkeypatch.setattr(detector.magic, "from_file", lambda *a, **k: "application/octet-stream")
    assert FileIntelligenceEngine().analyze(str(renamed)).file_type == kind


def test_non_office_zip_with_docx_extension_rejected(tmp_path):
    from app.engines.file_intelligence import FileIntelligenceEngine
    path = tmp_path / "fake.docx"
    with ZipFile(path, "w") as archive:
        archive.writestr("ordinary.txt", "Not an Office document")
    result = DocumentPipeline().process(path)
    assert result.failure.stage == "file_intelligence"


def test_library_added_empty_placeholder_does_not_invent_mapping(tmp_path):
    from app.engines.extraction.office import NS, parse_xml
    from lxml import etree
    path = tmp_path / "empty-placeholder.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = ""
    deck.save(path)
    with ZipFile(path) as archive:
        entries = [(i, archive.read(i)) for i in archive.infolist()]
    with ZipFile(path, "w") as archive:
        for info, data in entries:
            if info.filename == "ppt/slides/slide1.xml":
                root = parse_xml(data)
                for node in root.xpath("//p:txBody/a:p", namespaces=NS):
                    node.getparent().remove(node)
                data = etree.tostring(root)
            archive.writestr(info, data)
    extraction = extract_pptx(path)
    assert extraction.pages[0].text == "" and extraction.structure.segments == []
    detection = PIIDetectionEngine(analyzer=SyntheticAnalyzer()).detect(extraction)
    engine = OfficeArtifactEngine(tmp_path / "private-artifacts")
    redaction = RedactionEngine().redact(extraction, detection)
    artifact = engine.verify(extraction, redaction, engine.redact(extraction, detection, redaction))
    assert artifact.reopened and artifact.supported_content_verified


def test_output_directory_guard_fails_without_public_artifact(office_case, tmp_path):
    _, path, _, pipeline = office_case
    pipeline.artifact_output_dir = tmp_path / "public"
    result = pipeline.process(path)
    assert result.failure.stage == "artifact_redaction"
    assert not (tmp_path / "public").exists()


def test_broken_mapping_blocks_before_writing(office_case):
    _, path, extract, _ = office_case
    extraction = extract(path)
    nonempty = next(s for s in extraction.structure.segments if s.nodes)
    nonempty.nodes[0].end += 1
    detection = PIIDetectionEngine(analyzer=SyntheticAnalyzer()).detect(extraction)
    engine = OfficeArtifactEngine(path.parent / "private-artifacts")
    with pytest.raises(ArtifactValidationError):
        engine.redact(extraction, detection, RedactionEngine().redact(extraction, detection))
    assert not engine.output_dir.exists()
