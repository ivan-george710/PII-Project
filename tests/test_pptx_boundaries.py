"""Structural patterns from the real deck, without Cadence-specific values."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
import pytest

from app.engines.extraction.pptx import extract_pptx
from app.engines.pii_detection.models import DetectedPII, PIIDetectionResult
from app.engines.redaction import RedactionEngine
from app.engines.redaction.artifacts import OfficeArtifactEngine


@pytest.mark.parametrize("kind", ["text_frame", "table_cell"])
@pytest.mark.parametrize("mode", ["trailing_separator", "adjacent_paragraphs"])
def test_deterministic_boundaries_and_unrelated_text(tmp_path, kind, mode):
    path=tmp_path/"source.pptx"
    deck=Presentation()
    slide=deck.slides.add_slide(deck.slide_layouts[6])
    if kind=="table_cell":
        table=slide.shapes.add_table(1,2,Inches(1),Inches(1),Inches(6),Inches(1)).table
        frame=table.cell(0,0).text_frame
        table.cell(0,1).text="Unrelated cell"
    else:
        frame=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(2)).text_frame
    if mode=="trailing_separator":
        frame.text="Ada Example"
        frame.add_paragraph().text="Unrelated paragraph"
    else:
        frame.text="Before Ada"
        frame.add_paragraph().text="Example After"
    deck.save(path)
    extraction=extract_pptx(path)
    text=extraction.pages[0].text
    value="Ada Example\n" if mode=="trailing_separator" else "Ada\nExample"
    start=text.index(value)
    finding=DetectedPII("PERSON",value,start,start+len(value),.9,"pattern","Synthetic",1,
        validation_decision="CONFIRM",final_decision="CONFIRM")
    detection=PIIDetectionResult(path,path.name,[finding],[],[],1,[1])
    redaction=RedactionEngine().redact(extraction,detection)
    engine=OfficeArtifactEngine(tmp_path/"private-artifacts")
    artifact=engine.verify(extraction,redaction,engine.redact(extraction,detection,redaction))
    assert artifact.mapping_failures==artifact.verification_failures==0
    assert artifact.supported_content_verified and artifact.reopened
    output=extract_pptx(artifact.output_path)
    assert "[PERSON_001]" in output.full_text and value not in output.full_text
    if mode=="trailing_separator": assert "Unrelated paragraph" in output.full_text
    else: assert "Before" in output.full_text and "After" in output.full_text
    if kind=="table_cell": assert "Unrelated cell" in output.full_text
    assert len(output.pages)==1 and not artifact.safe_for_llm


def test_cross_shape_keeps_ambiguous_span(tmp_path):
    path=tmp_path/"source.pptx"
    deck=Presentation()
    slide=deck.slides.add_slide(deck.slide_layouts[6])
    for i,text in enumerate(["Ada", "Example"]):
        slide.shapes.add_textbox(Inches(1),Inches(i+1),Inches(3),Inches(1)).text=text
    deck.save(path)
    extraction=extract_pptx(path)
    text=extraction.pages[0].text
    d=DetectedPII("PERSON",text,0,len(text),.9,"pattern","Synthetic",1,validation_decision="CONFIRM",final_decision="CONFIRM")
    detection=PIIDetectionResult(path,path.name,[d],[],[],1)
    engine=OfficeArtifactEngine(tmp_path/"private-artifacts")
    redaction=RedactionEngine().redact(extraction,detection)
    artifact=engine.verify(extraction,redaction,engine.redact(extraction,detection,redaction))
    assert artifact.mapping_issue_count==1 and artifact.verification_failure_count>0
    assert not artifact.supported_content_verified
