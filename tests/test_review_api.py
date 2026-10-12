"""Local API boundaries with real small-format extraction/artifacts and deterministic analyzers."""
from io import BytesIO
from pathlib import Path
import re
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from fastapi.testclient import TestClient
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas
from presidio_analyzer import RecognizerResult

from app.pipeline import DocumentPipeline
from app.engines.audit import AuditEngine
from app.engines.extraction.pdf import extract_text_pdf
from app.engines.extraction.docx import extract_docx
from app.engines.extraction.pptx import extract_pptx
from app.engines.pii_detection import PIIDetectionEngine
from app.web.server import create_app

from manager.server import create_app as create_manager_app
from app.web.manager_client import ManagerClient

HEADERS={"X-Cadence-Local":"1"}
LINES=["Owner: John Smith", "Employee: EMP-00001", "Contact: Brightwell", "Missed: SecretToken", "Non-sensitive business policy"]


class Analyzer:
    def __init__(self): self.calls=0
    def analyze(self,text,language):
        self.calls+=1
        result=[]
        for entity,pattern in [("EMPLOYEE_ID",r"EMP-\d{5}"),("PERSON",r"John Smith|Brightwell")]:
            for m in re.finditer(pattern,text):
                result.append(RecognizerResult(entity,m.start(),m.end(),.9,recognition_metadata={"recognizer_name":"Synthetic"}))
        return result


def document_bytes(kind,lines=LINES):
    stream=BytesIO()
    if kind=="pdf":
        doc=canvas.Canvas(stream)
        for i,line in enumerate(lines): doc.drawString(50,740-i*25,line)
        doc.save()
    elif kind=="docx":
        doc=Document()
        for line in lines: doc.add_paragraph(line)
        doc.save(stream)
    else:
        doc=Presentation()
        slide=doc.slides.add_slide(doc.slide_layouts[6])
        for i,line in enumerate(lines): slide.shapes.add_textbox(Inches(1),Inches(i+1),Inches(6),Inches(.7)).text=line
        doc.save(stream)
    return stream.getvalue()


def setup(tmp_path,maximum=30*1024*1024):
    root=tmp_path/"private-app"
    analyzer=Analyzer()
    def extract(path): return {".pdf":extract_text_pdf,".docx":extract_docx,".pptx":extract_pptx}[Path(path).suffix](path)
    extraction=Mock(side_effect=extract)
    pipeline=DocumentPipeline(extraction=SimpleNamespace(extract=extraction),detection=PIIDetectionEngine(analyzer=analyzer),
        audit=AuditEngine(root/"audits"),artifact_output_dir=root/"artifacts")
    manager_app=create_manager_app(path=root/"manager.sqlite")
    manager_tc=TestClient(manager_app)
    mgr_client=ManagerClient(client=manager_tc)
    app=create_app(work_dir=root,pipeline=pipeline,max_upload_bytes=maximum,allow_test_host=True,manager_client=mgr_client)
    app.state.manager_test_client=manager_tc
    return TestClient(app),app,analyzer,extraction


@pytest.fixture
def api(tmp_path): return setup(tmp_path)


def upload(client,kind="docx",content=None,name=None):
    return client.post("/api/documents",headers=HEADERS,files={"file":(name or "private-name."+kind,content or document_bytes(kind))})


@pytest.mark.parametrize("kind",["pdf","docx","pptx"])
def test_upload_get_findings_context_artifact_and_no_path_leaks(api,kind):
    client,app,_,_=api
    mgr_tc=app.state.manager_test_client
    response=upload(client,kind)
    assert response.status_code==200,response.text
    state=response.json();sid=state["document_session_id"]
    assert re.fullmatch(r"[0-9a-f]{32}",sid)
    assert state["format"]==kind and state["summary"]["artifact"]["supported_content_verified"]
    assert state["summary"]["audit"]["counts"]["raw_detection_count"]==3
    assert not state["summary"]["safe_for_llm"] and not state["summary"]["coverage_verified"]
    assert client.get("/api/documents/"+sid).json()==state
    assert client.get("/api/documents").json()[0]==state
    findings=client.get(f"/api/documents/{sid}/findings").json()
    assert len(findings["items"])==3
    reviews=client.get(f"/api/documents/{sid}/findings?decision=REVIEW").json()
    assert reviews["total"]==1 and reviews["items"][0]["effective_decision"]=="REVIEW"
    detail=client.get(f"/api/documents/{sid}/findings/{reviews['items'][0]['finding_id']}").json()
    assert detail["candidate"]=="Brightwell"
    assert sum(len(detail[k]) for k in ["before","candidate","after"])<=480

    # Initial redacted artifact is available immediately without awaiting review/approval
    assert client.get(f"/api/documents/{sid}/artifact").status_code==200

    # Manager report is pending and unready; manager cannot approve unready engine results
    rep=mgr_tc.get(f"/api/manager/reports/{sid}").json()
    assert rep["status"]=="pending" and not rep["ready"]
    bad_approve=mgr_tc.post(f"/api/manager/reports/{sid}/decision",headers=HEADERS,json={"revision":0,"artifact_sha256":None,"action":"approved"})
    assert bad_approve.status_code==409

    # Resolve review item -> human review complete
    f_id=reviews["items"][0]["finding_id"]
    fb=client.post(f"/api/documents/{sid}/feedback",headers=HEADERS,json={"expected_revision":0,"finding_id":f_id,"action":"CONFIRM_PII"})
    assert fb.status_code==200
    rev1=fb.json()["revision"]
    assert fb.json()["summary"]["human_review_complete"]

    # Sync manager report with revised result
    m_status=client.get(f"/api/documents/{sid}/manager").json()
    assert m_status["ready"] is True and m_status["status"]=="pending"
    sha=m_status["artifact_sha256"]
    assert sha is not None

    # Redacted artifact is available immediately without awaiting manager approval
    download=client.get(f"/api/documents/{sid}/artifact")
    assert download.status_code==200 and download.content
    assert ("redacted-private-name."+kind) in download.headers["content-disposition"]

    # Manager approves revision 1
    approve=mgr_tc.post(f"/api/manager/reports/{sid}/decision",headers=HEADERS,json={"revision":rev1,"artifact_sha256":sha,"action":"approved"})
    assert approve.status_code==200
    assert client.get(f"/api/documents/{sid}/artifact").status_code==200

    # Manager denies -> download blocked (403)
    deny=mgr_tc.post(f"/api/manager/reports/{sid}/decision",headers=HEADERS,json={"revision":rev1,"artifact_sha256":sha,"action":"denied"})
    assert deny.status_code==200
    assert client.get(f"/api/documents/{sid}/artifact").status_code==403

    # Re-approve -> download unblocked (200)
    mgr_tc.post(f"/api/manager/reports/{sid}/decision",headers=HEADERS,json={"revision":rev1,"artifact_sha256":sha,"action":"approved"})
    assert client.get(f"/api/documents/{sid}/artifact").status_code==200

    for secret in [str(app.state.work_dir),"private-name", "John Smith", "Brightwell", "EMP-00001"]:
        assert secret not in response.text
        assert secret not in str(findings)
    assert response.headers["cache-control"]=="no-store"
    assert client.get(f"/api/documents/{sid}/review").json()["display_name"]=="private-name."+kind


def test_filters_pagination_and_search_metadata_only(api):
    client,*_=api
    sid=upload(client).json()["document_session_id"]
    base=f"/api/documents/{sid}/findings"
    assert client.get(base+"?decision=CONFIRM").json()["total"]==2
    assert client.get(base+"?entity=EMPLOYEE_ID").json()["total"]==1
    assert client.get(base+"?location=2").json()["total"]==1
    assert client.get(base+"?q=John").json()["total"]==0
    data=client.get(base+"?offset=1&limit=1").json()
    assert data["total"]==3 and len(data["items"])==1


def test_all_feedback_actions_regenerate_and_preserve_inputs(api):
    client,app,analyzer,extraction=api
    state=upload(client).json();sid=state["document_session_id"]
    initial=app.state.sessions[sid].result
    original_path=initial.source_path
    original_bytes=original_path.read_bytes()
    before_calls=analyzer.calls
    records=client.get(f"/api/documents/{sid}/findings").json()["items"]
    review=next(r for r in records if r["effective_decision"]=="REVIEW")
    body={"expected_revision":0,"finding_id":review["finding_id"],"action":"CONFIRM_PII","reviewer_id":"reviewer","reason":"fixture confirmation"}
    confirm=client.post(f"/api/documents/{sid}/feedback",headers=HEADERS,json=body)
    assert confirm.status_code==200 and confirm.json()["revision"]==1
    confirmed_artifact=app.state.sessions[sid].result.artifact_result.path
    assert confirmed_artifact!=initial.artifact_result.path
    body.update(expected_revision=1,action="FALSE_POSITIVE")
    rejected=client.post(f"/api/documents/{sid}/feedback",headers=HEADERS,json=body)
    assert rejected.status_code==200 and rejected.json()["revision"]==2
    current=app.state.sessions[sid].result
    assert "Brightwell" in extract_docx(current.artifact_result.path).full_text
    window=client.get(f"/api/documents/{sid}/source?location=4").json()
    start=window["text"].index("SecretToken")
    missed=client.post(f"/api/documents/{sid}/missed-pii",headers=HEADERS,json={"expected_revision":2,"location":4,
        "start":start,"end":start+11,"entity_type":"PERSON","reviewer_id":"reviewer"})
    assert missed.status_code==200 and missed.json()["revision"]==3
    assert "SecretToken" not in extract_docx(app.state.sessions[sid].result.artifact_result.path).full_text
    assert analyzer.calls==before_calls and extraction.call_count==1
    assert original_path.read_bytes()==original_bytes and initial.audit_result.revision==0
    audit=client.get(f"/api/documents/{sid}/audit").json()
    assert [e["action"] for e in audit["events"]]==["CONFIRM_PII","FALSE_POSITIVE","MISSED_PII"]
    assert "fixture confirmation" not in str(missed.json())


@pytest.mark.parametrize("name,content,status",[("bad.txt",b"not a supported document",415),
    ("bad.pdf",b"%PDF-1.7\ncorrupt",400),("bad.pdf",document_bytes("docx"),400),
    ("../bad.pdf",document_bytes("pdf"),400)],
    ids=["unsupported","corrupt","mismatched-content","traversal"])
def test_reject_uploads_without_processing_or_path_echo(api,name,content,status):
    client,app,analyzer,_=api
    response=client.post("/api/documents",headers=HEADERS,files={"file":(name,content)})
    assert response.status_code==status
    assert analyzer.calls==0 and not app.state.sessions
    assert not [p for p in (app.state.work_dir/"uploads").iterdir() if p.name != ".gitignore"]
    assert str(app.state.work_dir) not in response.text and name not in response.text


def test_oversized_stream_and_header_limit(tmp_path):
    client,app,_,_=setup(tmp_path,maximum=1024)
    response=upload(client,"pdf")
    assert response.status_code==413
    assert not [p for p in (app.state.work_dir/"uploads").iterdir() if p.name != ".gitignore"]
    response=client.post("/api/documents",headers={**HEADERS,"Content-Length":"999999"},content=b"x")
    assert response.status_code==413


def test_multipart_windows_filename_is_sanitized_before_private_storage(api):
    client,app,*_=api
    response=upload(client,"pdf",name="C:\\private\\file.pdf")
    assert response.status_code==200
    sid=response.json()["document_session_id"]
    identity=client.get(f"/api/documents/{sid}/review").json()
    assert identity["display_name"]=="file.pdf"  # python-multipart strips legacy browser fake paths.
    assert app.state.sessions[sid].result.source_path.parent==app.state.work_dir/"uploads"
    assert "C:\\private" not in response.text


def test_invalid_session_finding_and_traversal(api):
    client,*_=api
    assert client.get("/api/documents/"+"f"*32).status_code==404
    assert client.get("/api/documents/not-a-session").status_code==404
    assert client.get("/api/documents/%2e%2e%2fartifact").status_code==404
    sid=upload(client).json()["document_session_id"]
    assert client.get(f"/api/documents/{sid}/findings/D999999").status_code==404
    response=client.post(f"/api/documents/{sid}/feedback",headers=HEADERS,json={"expected_revision":0,"finding_id":"D999999","action":"CONFIRM_PII"})
    assert response.status_code==404


def test_stale_revision_conflict_and_invalid_manual_selection(api):
    client,app,*_=api
    sid=upload(client).json()["document_session_id"]
    finding=client.get(f"/api/documents/{sid}/findings?decision=REVIEW").json()["items"][0]
    body={"expected_revision":0,"finding_id":finding["finding_id"],"action":"CONFIRM_PII"}
    assert client.post(f"/api/documents/{sid}/feedback",headers=HEADERS,json=body).status_code==200
    assert client.post(f"/api/documents/{sid}/feedback",headers=HEADERS,json=body).status_code==409
    invalid={"expected_revision":1,"location":999,"start":0,"end":10,"entity_type":"PERSON"}
    assert client.post(f"/api/documents/{sid}/missed-pii",headers=HEADERS,json=invalid).status_code==422
    invalid.update(location=1,start=99999,end=100000)
    assert client.post(f"/api/documents/{sid}/missed-pii",headers=HEADERS,json=invalid).status_code==422
    invalid.update(start=0,end=3,replacement_text="arbitrary")
    assert client.post(f"/api/documents/{sid}/missed-pii",headers=HEADERS,json=invalid).status_code==422
    assert app.state.sessions[sid].result.audit_result.revision==1


def test_source_window_bounded_unicode_offsets(api):
    client,app,*_=api
    sid=upload(client,content=document_bytes("docx",["Prefix 😀 SecretToken suffix"+"x"*3000])).json()["document_session_id"]
    window=client.get(f"/api/documents/{sid}/source?location=1&length=20").json()
    assert len(window["text"])==20 and window["total_characters"]>3000
    start=window["text"].index("SecretToken")
    response=client.post(f"/api/documents/{sid}/missed-pii",headers=HEADERS,json={"expected_revision":0,"location":1,
        "start":start,"end":start+11,"entity_type":"PERSON"})
    assert response.status_code==200
    assert app.state.sessions[sid].result.audit_result.records[0].original_detection.text=="SecretToken"
    assert client.get(f"/api/documents/{sid}/source?length=2001").status_code==422


def test_artifact_verification_failure_is_visible(api):
    client,app,*_=api
    backend=app.state.pipeline
    from app.engines.redaction.document_artifacts import DocumentArtifactEngine
    engine=DocumentArtifactEngine(app.state.work_dir/"artifacts")
    real_verify=engine.verify
    def failed(e,r,a):
        result=real_verify(e,r,a)
        result.verification_failures+=1
        result.supported_content_verified=False
        return result
    engine.verify=failed
    backend.artifacts=engine
    state=upload(client).json()
    assert state["summary"]["artifact"]["verification_failures"]==1
    assert next(s for s in state["stages"] if s["name"]=="artifact_verification")["state"]=="Failed"
    assert state["artifact_available"] and not state["summary"]["safe_for_llm"]


def test_save_reload_uses_existing_audit_and_opaque_reference(api):
    client,app,analyzer,_=api
    sid=upload(client).json()["document_session_id"]
    response=client.post(f"/api/documents/{sid}/save",headers=HEADERS,json={"expected_revision":0})
    assert response.status_code==200 and response.json()["saved_revision"]==0
    calls=analyzer.calls
    saved=client.get("/api/saved-audits").json()
    assert len(saved)==1 and re.fullmatch(r"[0-9a-f]{32}",saved[0]["audit_reference"])
    response=client.post("/api/load-audit",headers=HEADERS,json={"audit_reference":saved[0]["audit_reference"]})
    assert response.status_code==200 and response.json()["document_session_id"]!=sid
    assert response.json()["summary"]["artifact"]["supported_content_verified"]
    assert analyzer.calls==calls
    assert str(app.state.work_dir) not in response.text


def test_security_headers_origin_host_and_local_mutation_guard(api):
    client,*_=api
    assert client.post("/api/documents",files={"file":("x.docx",document_bytes("docx"))}).status_code==403
    assert client.get("/api/config",headers={"Origin":"https://evil.invalid"}).status_code==403
    assert client.get("/api/config",headers={"Host":"evil.invalid"}).status_code==403
    assert client.get("/").status_code==200
    assert "script-src 'self'" in client.get("/").headers["content-security-policy"]
    assert client.get("/assets/../../../.env").status_code==404


def test_source_changed_failed_feedback_never_downloads_stale_artifact(api):
    client,app,*_=api
    sid=upload(client).json()["document_session_id"]
    current=app.state.sessions[sid].result
    current.source_path.write_bytes(current.source_path.read_bytes()+b"changed")
    finding=current.audit_result.records[0].detection_id
    response=client.post(f"/api/documents/{sid}/feedback",headers=HEADERS,json={"expected_revision":0,"finding_id":finding,"action":"CONFIRM_PII"})
    assert response.status_code==200
    assert response.json()["summary"]["failure"]["stage"]=="artifact_redaction"
    assert "original upload changed" in response.json()["issue_message"]
    assert not response.json()["artifact_available"]
    assert client.get(f"/api/documents/{sid}/artifact").status_code==409


def test_unexpected_error_no_exception_or_console_leak(api,capsys):
    client,app,*_=api
    app.state.pipeline.process=Mock(side_effect=RuntimeError("SECRET source@example.invalid /private/path"))
    response=upload(client)
    assert response.status_code==500
    assert "SECRET" not in response.text and "source@example" not in response.text
    assert "SECRET" not in capsys.readouterr().err


@pytest.mark.parametrize("kind", ["pdf", "docx", "pptx"])
def test_comparison_is_actual_bounded_redaction_without_private_paths(api, kind):
    client, app, *_ = api
    sid = upload(client, kind).json()["document_session_id"]
    current = app.state.sessions[sid].result
    record = next(r for r in current.audit_result.records if r.redaction_status == "APPLIED")
    response = client.get(f"/api/documents/{sid}/findings/{record.detection_id}/comparison")
    assert response.status_code == 200
    value = response.json()
    page = next(p for p in current.redaction_result.pages if p.page_number == record.effective_detection.page_number)
    assert value["replacement"] == page.text[record.redacted_start:record.redacted_end]
    assert value["replacement"] in value["redacted_context"]
    assert len(value["redacted_context"]) <= 480
    assert str(app.state.work_dir) not in response.text
    assert "source_sha256" not in response.text


def test_comparison_does_not_invent_a_replacement_for_unresolved_finding(api):
    client, app, *_ = api
    sid = upload(client).json()["document_session_id"]
    record = next(r for r in app.state.sessions[sid].result.audit_result.records if r.effective_decision == "REVIEW")
    response = client.get(f"/api/documents/{sid}/findings/{record.detection_id}/comparison")
    assert response.status_code == 200
    assert response.json()["redacted_context"] is None
    assert response.json()["replacement"] is None
    assert client.get(f"/api/documents/{sid}/findings/unknown/comparison").status_code == 404


def test_saved_feedback_timeline_survives_reload_without_detection(api):
    client, _, analyzer, _ = api
    sid = upload(client).json()["document_session_id"]
    finding = client.get(f"/api/documents/{sid}/findings?decision=REVIEW").json()["items"][0]
    feedback = client.post(f"/api/documents/{sid}/feedback", headers=HEADERS,
                          json={"expected_revision": 0, "finding_id": finding["finding_id"], "action": "CONFIRM_PII"})
    assert feedback.status_code == 200
    assert client.post(f"/api/documents/{sid}/save", headers=HEADERS, json={"expected_revision": 1}).status_code == 200
    reference = client.get("/api/saved-audits").json()[0]["audit_reference"]
    calls = analyzer.calls
    restored = client.post("/api/load-audit", headers=HEADERS, json={"audit_reference": reference}).json()["document_session_id"]
    timeline = client.get(f"/api/documents/{restored}/audit")
    assert timeline.status_code == 200
    assert timeline.json()["events"][0]["action"] == "CONFIRM_PII"
    assert analyzer.calls == calls


def test_manager_unavailable_does_not_block_local_download(tmp_path):
    root = tmp_path / "private-app"
    analyzer = Analyzer()
    pipeline = DocumentPipeline(extraction=SimpleNamespace(extract=lambda p: extract_docx(p)), detection=PIIDetectionEngine(analyzer=analyzer),
        audit=AuditEngine(root / "audits"), artifact_output_dir=root / "artifacts")
    unreachable_client = ManagerClient(url="http://127.0.0.1:59999")
    app = create_app(work_dir=root, pipeline=pipeline, allow_test_host=True, manager_client=unreachable_client)
    client = TestClient(app)
    res = upload(client, "docx")
    assert res.status_code == 200
    sid = res.json()["document_session_id"]
    finding = client.get(f"/api/documents/{sid}/findings?decision=REVIEW").json()["items"][0]
    client.post(f"/api/documents/{sid}/feedback", headers=HEADERS, json={"expected_revision": 0, "finding_id": finding["finding_id"], "action": "CONFIRM_PII"})
    down = client.get(f"/api/documents/{sid}/artifact")
    assert down.status_code == 200


def test_revision_reset_resets_approval_and_blocks_stale_download(api):
    client, app, _, _ = api
    mgr_tc = app.state.manager_test_client
    sid = upload(client, "docx").json()["document_session_id"]
    finding = client.get(f"/api/documents/{sid}/findings?decision=REVIEW").json()["items"][0]
    client.post(f"/api/documents/{sid}/feedback", headers=HEADERS, json={"expected_revision": 0, "finding_id": finding["finding_id"], "action": "CONFIRM_PII"})
    m_status = client.get(f"/api/documents/{sid}/manager").json()
    assert m_status["ready"] is True
    sha1 = m_status["artifact_sha256"]
    approve1 = mgr_tc.post(f"/api/manager/reports/{sid}/decision", headers=HEADERS, json={"revision": 1, "artifact_sha256": sha1, "action": "approved"})
    assert approve1.status_code == 200
    assert client.get(f"/api/documents/{sid}/artifact").status_code == 200

    rec = client.get(f"/api/documents/{sid}/findings?decision=CONFIRM").json()["items"][0]
    client.post(f"/api/documents/{sid}/feedback", headers=HEADERS, json={"expected_revision": 1, "finding_id": rec["finding_id"], "action": "FALSE_POSITIVE"})
    assert client.get(f"/api/documents/{sid}").json()["revision"] == 2

    # Deny blocks download
    m_status2 = client.get(f"/api/documents/{sid}/manager").json()
    deny = mgr_tc.post(f"/api/manager/reports/{sid}/decision", headers=HEADERS, json={"revision": 2, "artifact_sha256": m_status2["artifact_sha256"], "action": "denied"})
    assert deny.status_code == 200
    assert client.get(f"/api/documents/{sid}/artifact").status_code == 403

    stale_app = mgr_tc.post(f"/api/manager/reports/{sid}/decision", headers=HEADERS, json={"revision": 1, "artifact_sha256": sha1, "action": "approved"})
    assert stale_app.status_code == 409

    approve2 = mgr_tc.post(f"/api/manager/reports/{sid}/decision", headers=HEADERS, json={"revision": 2, "artifact_sha256": m_status2["artifact_sha256"], "action": "approved"})
    assert approve2.status_code == 200
    assert client.get(f"/api/documents/{sid}/artifact").status_code == 200


def test_text_message_manager_approval_and_sanitized_txt_download(api):
    client, app, _, _ = api
    mgr_tc = app.state.manager_test_client
    res = client.post("/api/messages", headers=HEADERS, json={"text": "Contact Brightwell about EMP-00001 please."})
    assert res.status_code == 200
    sid = res.json()["document_session_id"]
    assert res.json()["input_text"] == "Contact Brightwell about EMP-00001 please."

    finding = client.get(f"/api/documents/{sid}/findings?decision=REVIEW").json()["items"][0]
    client.post(f"/api/documents/{sid}/feedback", headers=HEADERS, json={"expected_revision": 0, "finding_id": finding["finding_id"], "action": "CONFIRM_PII"})
    m_status = client.get(f"/api/documents/{sid}/manager").json()
    assert m_status["ready"] is True
    mgr_tc.post(f"/api/manager/reports/{sid}/decision", headers=HEADERS, json={"revision": 1, "artifact_sha256": m_status["artifact_sha256"], "action": "approved"})
    down = client.get(f"/api/documents/{sid}/artifact")
    assert down.status_code == 200
    assert down.headers["content-type"].startswith("text/plain")
    assert "redacted-message.txt" in down.headers["content-disposition"]
    text = down.text
    assert "Brightwell" not in text
    assert "EMP-00001" not in text
    assert "please." in text
