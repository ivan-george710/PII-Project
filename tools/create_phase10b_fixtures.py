"""Synthetic browser acceptance documents; never alter detector evaluation data."""
from pathlib import Path
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen.canvas import Canvas

def main():
    root = Path('.private_review_app/phase10b-fixtures')
    root.mkdir(parents=True, exist_ok=True)
    lines = ['Policy owner: John Smith.', 'Bank account number: 73000000',
             'Employee: EMP-00001', 'Contact: ada@example.com', 'Missed: SecretToken']
    pdf = Canvas(str(root/'synthetic.pdf'))
    for i, line in enumerate(lines):
        pdf.drawString(50, 750-i*30, line)
    pdf.save()
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    doc.save(root/'synthetic.docx')
    ppt = Presentation()
    slide = ppt.slides.add_slide(ppt.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(4)).text = '\n'.join(lines)
    ppt.save(root/'synthetic.pptx')
    print('Created three fictional acceptance fixtures in the private review area.')

if __name__ == '__main__':
    main()
