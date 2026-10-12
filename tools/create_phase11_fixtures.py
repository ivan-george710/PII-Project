"""Local fictional face documents. Never copy real Cadence photos into fixtures."""
from pathlib import Path
import sys
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.shared import Inches as WordInches
from pptx import Presentation
from pptx.util import Inches
import pymupdf as fitz
from app.engines.visual_detection.images import encode


def create(directory):
    directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
    face=ROOT/'tests/fixtures/visual/synthetic-face.png'
    doc=Document();doc.add_paragraph('Synthetic biometric fixture. Employee: EMP-00001. Contact: ada@example.com.')
    doc.add_picture(str(face),width=WordInches(3));doc.add_paragraph('Surrounding process text is preserved.')
    doc.save(directory/'synthetic-face.docx')
    ppt=Presentation()
    for n in range(2):
        slide=ppt.slides.add_slide(ppt.slide_layouts[6])
        slide.shapes.add_textbox(Inches(.5),Inches(.2),Inches(8),Inches(.5)).text='Synthetic image privacy. Employee: EMP-00001; ada@example.com'
        if n==0:slide.shapes.add_picture(str(face),Inches(1),Inches(1),width=Inches(3))
        else:slide.shapes.add_group_shape().shapes.add_picture(str(face),Inches(1),Inches(1),width=Inches(3))
    ppt.save(directory/'synthetic-face.pptx')
    with fitz.open() as pdf:
        page=pdf.new_page(width=600,height=800)
        page.insert_image(fitz.Rect(50,50,450,450),filename=str(face))
        page.insert_text((50,650),'Synthetic face privacy fixture. Employee EMP-00001. Contact ada@example.com',fontsize=12)
        pdf.save(directory/'synthetic-face-native.pdf')
    with Image.open(face) as photo:
        scan=Image.new('RGB',(512,620),'white');scan.paste(photo,(0,0))
        ImageDraw.Draw(scan).text((20,540),'Employee EMP-00001',font=ImageFont.load_default(size=24),fill='black')
        data=encode(scan);scan.close()
    with fitz.open() as pdf:
        page=pdf.new_page(width=512,height=620);page.insert_image(page.rect,stream=data)
        pdf.save(directory/'synthetic-face-scan.pdf')
    return directory


if __name__=='__main__':
    from app.engines.audit.storage import AuditStore
    store=AuditStore(ROOT/'.private_phase11/fixtures');store._prepare();create(store.directory)
    print('Created local synthetic face PDF/DOCX/PPTX fixtures in private storage.')
